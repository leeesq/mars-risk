"""核对公开快照证据；显式提供 Home Credit 原文件时进行完整复算。"""

from __future__ import annotations

import argparse
import ast
import hashlib
import html
import json
import math
import os
import re
import struct
import subprocess
import sys
import tempfile
import zipfile
from datetime import date
from html.parser import HTMLParser
from pathlib import Path, PurePosixPath
from typing import Any
from xml.etree import ElementTree

ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "docs/assets/cases"
GENERATOR = ROOT / "docs/snippets/task_cases.py"
CASE_PAGES = (
    "data-quality", "binning-stability", "selection-correlation", "score-cross",
    "rule-evidence", "saved-reports", "report-delivery",
)
VOLATILE_FIELDS = frozenset({
    "report_id", "created_at", "generated_at", "created_at_utc", "generated_at_utc",
    "source_commit", "source_revision", "elapsed_seconds",
})


def _reject_constant(value: str) -> None:
    """拒绝 JSON 非有限常量，避免机器证据依赖非标准解析行为。"""
    raise ValueError(f"Non-finite JSON constant: {value}")


def _read_json(path: Path) -> Any:
    """读取标准有限 JSON；错误保留具体文件路径。"""
    try:
        return json.loads(path.read_text(encoding="utf-8"), parse_constant=_reject_constant)
    except (ValueError, OSError) as error:
        raise ValueError(f"Invalid case JSON {path.name}: {error}") from error


def _safe_path(directory: Path, name: str) -> Path:
    """仅允许产物目录内的相对 POSIX 路径。"""
    relative = PurePosixPath(name)
    if relative.is_absolute() or ".." in relative.parts or "\\" in name or ":" in name:
        raise ValueError(f"Unsafe artifact path: {name!r}")
    candidate = directory.joinpath(*relative.parts)
    if not candidate.is_file():
        raise ValueError(f"Missing artifact: {name}")
    return candidate


def _semantic(value: Any) -> Any:
    """排除本来随运行变化的身份和时间，同时保留业务字段、状态和数值。"""
    if isinstance(value, dict):
        return {
            key: _semantic(item) for key, item in value.items()
            if key not in VOLATILE_FIELDS and not key.endswith("_report_id")
        }
    if isinstance(value, list):
        return [_semantic(item) for item in value]
    return value


def _assert_equal(expected: Any, actual: Any, location: str = "summary") -> None:
    """逐字段比较语义，允许浮点最后几位的跨平台差异并报告首个陈旧路径。"""
    if isinstance(expected, dict) and isinstance(actual, dict):
        if expected.keys() != actual.keys():
            raise ValueError(f"Stale {location}: keys {sorted(expected)} != {sorted(actual)}")
        for key in expected:
            _assert_equal(expected[key], actual[key], f"{location}.{key}")
        return
    if isinstance(expected, list) and isinstance(actual, list):
        if len(expected) != len(actual):
            raise ValueError(f"Stale {location}: row count {len(expected)} != {len(actual)}")
        for index, (left, right) in enumerate(zip(expected, actual)):
            _assert_equal(left, right, f"{location}[{index}]")
        return
    if (
        isinstance(expected, (int, float)) and not isinstance(expected, bool)
        and isinstance(actual, (int, float)) and not isinstance(actual, bool)
    ):
        if math.isclose(expected, actual, rel_tol=1e-7, abs_tol=1e-10):
            return
    elif type(expected) is type(actual) and expected == actual:
        return
    raise ValueError(f"Stale {location}: published={expected!r}, generated={actual!r}")


def _table_rows(frame: Any) -> list[dict[str, Any]]:
    """仅物化已经分页的小证据表，支持公共接口的两个后端。"""
    from mars.reporting._serialization import json_safe

    rows: list[dict[str, Any]] = json_safe(
        frame.to_dicts() if hasattr(frame, "to_dicts") else frame.to_dict("records")
    )
    return rows


def _code_digest(path: Path) -> str:
    """只比较可执行 AST，说明文字、注释和排版不会强迫重新分析。"""
    tree = ast.parse(path.read_text(encoding="utf-8"))
    for node in ast.walk(tree):
        if isinstance(node, (ast.Module, ast.FunctionDef, ast.AsyncFunctionDef, ast.ClassDef)):
            if node.body and isinstance(node.body[0], ast.Expr):
                value = node.body[0].value
                if isinstance(value, ast.Constant) and isinstance(value.value, str):
                    node.body.pop(0)
    return hashlib.sha256(ast.dump(tree, include_attributes=False).encode("utf-8")).hexdigest()


def _validate_preparation(directory: Path, manifest: dict[str, Any]) -> dict[str, Any]:
    """无 CSV 模式只核对聚合输入证据与分母，不宣称重新验证真实成员。"""
    preparation: dict[str, Any] = _read_json(directory / "preparation-summary.json")
    _assert_equal(manifest["data_config"]["preparation"], preparation, "manifest.preparation")
    if preparation["preparation_version"] != "home-credit-2018-v1":
        raise ValueError("Unsupported Home Credit preparation version")
    files = {entry["role"]: entry for entry in preparation["input_files"]}
    if len(files) != 3 or set(files) != {"train", "test", "description"}:
        raise ValueError("Preparation must identify exactly the three required input files")
    for role, entry in files.items():
        if not re.fullmatch(r"[0-9a-f]{64}", entry["sha256"]) or entry["bytes"] <= 0:
            raise ValueError(f"Invalid source SHA256 or byte count: {role}")
        if Path(entry["name"]).name != entry["name"] or re.search(r"[\\/:]", entry["name"]):
            raise ValueError(f"Source fingerprint leaks a private path: {role}")
    total = files["train"]["rows"] + files["test"]["rows"]
    width = files["train"]["columns"]
    _assert_equal([total, width], preparation["merged_shape_before_roles"], "preparation.merged_shape")
    _assert_equal([total, width + 2], preparation["final_shape"], "preparation.final_shape")
    if files["test"]["columns"] != width - 1:
        raise ValueError("Original test must have one fewer column because TARGET is unavailable")
    if preparation["full_dataset"]:
        for role, shape in (("train", [307511, 122]), ("test", [48744, 121]), ("description", [219, 5])):
            _assert_equal(shape, [files[role]["rows"], files[role]["columns"]], f"input.{role}.shape")
    for key in ("seed", "rows"):
        _assert_equal(preparation[key], manifest["data_config"][key], f"manifest.data_config.{key}")
    if preparation["hash_config"]["algorithm"] != "SHA256" or preparation["hash_config"]["uses_binary_target_value"] is not False:
        raise ValueError("Preparation must use stable SHA256 independent of binary target values")
    if preparation["hash_config"]["encoding"] != "UTF-8":
        raise ValueError("Stable ID hash encoding differs from the recorded UTF-8 contract")
    roles = {"SK_ID_CURR", "TARGET", "APPLY_DATE", "dataset_flag", "AMT_CREDIT", "__home_credit_source"}
    features = preparation["fcols"]
    if len(features) != len(set(features)) or roles.intersection(features):
        raise ValueError("Candidate features contain duplicate or role/source columns")
    if len(features) != preparation["candidate_feature_count"] or len(features) != width - 3:
        raise ValueError("Candidate count does not exclude ID, TARGET and the separate amount role")
    if not set(preparation["display_features"]).issubset(features):
        raise ValueError("Display features must be a declared subset of the complete candidate fields")
    groups = preparation["partition_stats"]
    labels = [entry["dataset_flag"] for entry in groups]
    if labels != ["0train", "1val", "2test", "3oot", "4_oot_undue"]:
        raise ValueError("Preparation must retain the five exact ordered partition labels")
    if sum(entry["total_count"] for entry in groups) != total:
        raise ValueError("Partition aggregates do not cover the merged row count")
    previous_end: date | None = None
    for entry in groups:
        group = entry["dataset_flag"]
        observed, unobserved = entry["observed_count"], entry["null_count"]
        count = entry["total_count"]
        start, end = date.fromisoformat(entry["date_min"]), date.fromisoformat(entry["date_max"])
        if count <= 0 or start > end or (previous_end is not None and previous_end >= start):
            raise ValueError(f"Partitions are empty, overlap, or violate date order: {group}")
        previous_end = end
        if observed + unobserved != count or entry["target_0_count"] + entry["target_1_count"] != observed:
            raise ValueError(f"Observed/null and binary-label denominators differ: {group}")
        _assert_equal(observed / count, entry["label_coverage"], f"preparation.{group}.coverage")
        if group == "4_oot_undue":
            if start < date(2026, 7, 1) or end > date(2026, 7, 31) or count != files["test"]["rows"]:
                raise ValueError("All original application_test rows must remain in the July source group")
            if observed != 0 or entry["observed_bad_rate"] is not None or entry["label_status"] != "unobserved":
                raise ValueError("Original test labels cannot be invented or treated as good samples")
        else:
            if start < date(2026, 1, 1) or end > date(2026, 6, 30) or unobserved or observed != count:
                raise ValueError(f"Labeled partition must preserve observed labels and Jan–Jun dates: {group}")
            _assert_equal(entry["target_1_count"] / observed, entry["observed_bad_rate"], f"preparation.{group}.bad_rate")
    if sum(entry["observed_count"] for entry in groups) != files["train"]["rows"]:
        raise ValueError("Observed labels do not match the original training row count")
    summary = _read_json(directory / "summary.json")
    _assert_equal(total, summary["rows"], "summary.rows")
    _assert_equal({entry["dataset_flag"]: entry["total_count"] for entry in groups}, summary["splits"], "summary.splits")
    metadata = _read_json(directory / "feature-metadata.json")
    for field, entry in metadata.items():
        if set(entry) - {"display_name", "description", "data_source", "unit"}:
            raise ValueError(f"Public feature metadata violates the MARS four-field contract: {field}")
    if not set(features).issubset(metadata):
        raise ValueError("Candidate features lack delivered business metadata")
    return preparation


def _validate_saved_denominators(preparation: dict[str, Any], snapshots: dict[str, Any]) -> None:
    """将五组人数与保存的全量分箱/交叉对齐，监督量仅使用观测标签。"""
    groups = {entry["dataset_flag"]: entry for entry in preparation["partition_stats"]}
    all_groups = set(groups)
    expected_total = {
        "total_count": preparation["rows"],
        "observed_count": sum(entry["observed_count"] for entry in groups.values()),
        "target_1_count": sum(entry["target_1_count"] for entry in groups.values()),
    }
    checked_binning = False
    checked_cross = False
    for report in snapshots.values():
        description = report.describe()
        context = description["business_context"]
        if context.get("dataset_id") != "home-credit-default-risk-2018":
            raise ValueError("Saved report does not belong to the declared Home Credit batch")
        if context["date"]["simulated"] is not True or context["real_features"]["simulated"] is not False:
            raise ValueError("Saved context must distinguish real features from simulated dates")
        for table in description["tables"]:
            columns = report.get_table(table, limit=0).columns
            if "SK_ID_CURR" in columns:
                raise ValueError(f"Saved statistics must not distribute actual sample IDs: {table}")
        if description["report_type"] == "MarsBinningReport":
            details = _table_rows(report.get_table("detail"))
            actual_groups = {row["mars_group"] for row in details}
            if not all_groups.issubset(actual_groups):
                continue  # 月趋势与单字段特殊值视图只检验适用维度。
            checked_binning = True
            keys: dict[tuple[str, str, str], list[dict[str, Any]]] = {}
            for row in details:
                keys.setdefault((row["feature"], row["y"], row["mars_group"]), []).append(row)
            for (feature, target, group), rows in keys.items():
                if target != "TARGET" or group not in all_groups | {"Total"}:
                    raise ValueError("Binning uses an unexpected target or partition")
                aggregate = [row for row in rows if row["bin_type"] == "汇总组"]
                if len(aggregate) != 1:
                    raise ValueError(f"Binning feature/group requires one aggregate row: {feature}/{group}")
                total_row = aggregate[0]
                expected = groups.get(group, expected_total)
                for field, source in (("count", "total_count"), ("observed_count", "observed_count"), ("bad", "target_1_count")):
                    _assert_equal(expected[source], total_row[field], f"binning.{feature}.{group}.{field}")
                    _assert_equal(total_row[field], sum(row[field] for row in rows if row["bin_type"] != "汇总组"), f"binning.{feature}.{group}.{field}.bins")
                if group == "4_oot_undue" and total_row["bad_rate"] is not None:
                    raise ValueError("Binning must not create an unobserved bad rate")
                if group == "4_oot_undue":
                    for row in rows:
                        for metric in ("bad_rate", "lift", "cum_bad_rate", "ks_bin", "auc_bin", "iv_bin", "amt_bad_rate", "lift_amt"):
                            if metric in row and row[metric] is not None:
                                raise ValueError(f"Unobserved bin including aggregate cannot have supervised {metric}: {feature}")
            states = _table_rows(report.get_table("calculation_status", filters={"group": "4_oot_undue"}))
            if not states or any(row["status"] != "unobserved" for row in states if row["metric"] == "risk_metrics"):
                raise ValueError("Binning unobserved target requires its real calculation status")
            supervised_trends: list[str] = ["iv", "ks", "auc"]
            # 新报告在 detail 保留总体与逐箱坏率；历史快照若保存旧趋势表仍逐值核查。
            if "trend.bad_rate" in description["tables"]:
                supervised_trends.append("bad_rate")
            for metric in supervised_trends:
                for row in _table_rows(report.get_table(f"trend.{metric}")):
                    if row["4_oot_undue"] is not None:
                        raise ValueError(f"Unobserved group cannot have supervised {metric}")
            if "4_oot_undue" not in report.get_table("trend.psi", limit=0).columns:
                raise ValueError("Unlabeled distribution/PSI group was discarded")
        elif description["report_type"] == "score_cross":
            overall = _table_rows(report.get_table("overall"))
            if {row["group"] for row in overall} != all_groups:
                raise ValueError("Cross report must preserve all five source partitions")
            checked_cross = True
            cells = _table_rows(report.get_table("cells"))
            for row in overall:
                group = row["group"]
                expected = groups[group]
                subset = [cell for cell in cells if cell["group"] == group and cell["target"] == row["target"] and cell["period"] == row["period"]]
                for field, source in (("sample_count", "total_count"), ("observed_sample_count", "observed_count"), ("bad_sample_count", "target_1_count")):
                    _assert_equal(expected[source], row[field], f"cross.{group}.{field}")
                    _assert_equal(row[field], sum(cell[field] for cell in subset), f"cross.{group}.{field}.cells")
                _assert_equal(expected["observed_bad_rate"], row["bad_rate"], f"cross.{group}.bad_rate")
                if group == "4_oot_undue" and (row["status"] != "unobserved" or any(cell["bad_rate"] is not None for cell in subset)):
                    raise ValueError("Cross unlabeled state or observed denominator was fabricated")
            fit = description["parameters"]["binning_parameters"]
            for axis in ("x", "y"):
                _assert_equal(groups["0train"]["total_count"], fit[axis]["reference_row_count"], f"cross.fit.{axis}.reference_rows")
    if not checked_binning or not checked_cross:
        raise ValueError("Batch lacks complete five-partition binning or score-cross denominator evidence")


def _json_metadata_rows(rows: list[dict[str, Any]], label: str) -> dict[str, Any]:
    """按 key/part 重组 HTML 与 Excel 的 JSON 片段，检测截断、缺片和重复片段。"""
    chunks: dict[str, dict[int, str]] = {}
    counts: dict[str, int] = {}
    for row in rows:
        key = row["key"]
        part, parts = int(row.get("part", 1)), int(row.get("parts", 1))
        if key is None or not isinstance(row["value"], str) or not 1 <= part <= parts:
            raise ValueError(f"Invalid Excel metadata chunk: {label}/{key}")
        if part in chunks.setdefault(key, {}) or (key in counts and counts[key] != parts):
            raise ValueError(f"Duplicate or inconsistent Excel metadata chunk: {label}/{key}")
        chunks[key][part], counts[key] = row["value"], parts
    result: dict[str, Any] = {}
    for key, indexed_parts in chunks.items():
        if set(indexed_parts) != set(range(1, counts[key] + 1)):
            raise ValueError(f"Missing Excel metadata chunk: {label}/{key}")
        result[key] = json.loads("".join(indexed_parts[index] for index in range(1, counts[key] + 1)), parse_constant=_reject_constant)
    return result


def _excel_json_mapping(sheet: Any, label: str) -> dict[str, Any]:
    """读取实际工作表的 key/value 及有顺序长 JSON 片段。"""
    rows = sheet.iter_rows(values_only=True)
    header = list(next(rows))
    if not {"key", "value"}.issubset(header):
        raise ValueError(f"Excel metadata has no key/value header: {label}")
    return _json_metadata_rows([dict(zip(header, values)) for values in rows], label)


class _HtmlEvidenceParser(HTMLParser):
    """只读取静态 HTML 表格和原生分箱数值属性，不执行页面脚本。"""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.tables: list[list[list[str]]] = []
        self.metrics: list[tuple[str, dict[str, float]]] = []
        self._table: list[list[str]] | None = None
        self._row: list[str] | None = None
        self._cell: list[str] | None = None

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag == "table":
            self._table = []
        elif tag == "tr" and self._table is not None:
            self._row = []
            properties = dict(attrs)
            payload = properties.get("data-metrics")
            if payload:
                self.metrics.append((properties.get("data-feature") or "", json.loads(payload)))
        elif tag in {"td", "th"} and self._row is not None:
            self._cell = []

    def handle_data(self, data: str) -> None:
        if self._cell is not None:
            self._cell.append(data)

    def handle_endtag(self, tag: str) -> None:
        if tag in {"td", "th"} and self._cell is not None and self._row is not None:
            self._row.append("".join(self._cell))
            self._cell = None
        elif tag == "tr" and self._row is not None and self._table is not None:
            self._table.append(self._row)
            self._row = None
        elif tag == "table" and self._table is not None:
            self.tables.append(self._table)
            self._table = None


def _validate_html_exports(directory: Path) -> None:
    """比对 HTML 的实际目录、业务元数据、交叉表及原生数值属性。"""
    from mars.reporting import load_report

    for path in directory.glob("*.marsreport"):
        html_path = path.with_suffix(".html")
        if not html_path.exists():
            continue
        report = load_report(path)
        expected = report.describe()
        content = html_path.read_text(encoding="utf-8")
        if expected["report_type"] == "score_cross":
            payload = re.search(r'<script type="application/json" id="data">(.*?)</script>', content, re.S)
            if payload is None:
                raise ValueError("Cross HTML lacks its actual saved payload")
            actual = json.loads(payload[1], parse_constant=_reject_constant)
            _assert_equal(expected, actual["description"], f"HTML.{path.stem}.description")
            for name, values in actual["tables"].items():
                _assert_equal(_table_rows(report.get_table(name)), values, f"HTML.{path.stem}.{name}")
            continue
        if expected["report_type"] == "rule":
            descriptions = [json.loads(html.unescape(value), parse_constant=_reject_constant)
                            for value in re.findall(r"<pre>(.*?)</pre>", content, re.S)
                            if html.unescape(value).lstrip().startswith("{")]
            if descriptions != [expected]:
                raise ValueError(f"Rule HTML lacks the complete saved description: {path.stem}")
            continue
        parser = _HtmlEvidenceParser()
        parser.feed(content)
        contexts: list[dict[str, Any]] = []
        catalogs: list[dict[str, Any]] = []
        for table in parser.tables:
            if not table:
                continue
            header = table[0]
            if {"key", "value"}.issubset(header):
                rows = [dict(zip(header, values)) for values in table[1:]]
                keys = {row["key"] for row in rows}
                if "report_id" in keys or "dataset_id" in keys:
                    decoded = _json_metadata_rows(rows, f"HTML.{path.stem}.metadata")
                    (catalogs if "report_id" in keys else contexts).append(decoded)
        semantics = {key: value for key, value in expected.items() if key not in {"feature_metadata", "business_context"}}
        if contexts != [expected["business_context"]] or catalogs != [semantics]:
            raise ValueError(f"HTML metadata/parameters differ from its snapshot: {path.stem}")
        candidates: dict[str, list[dict[str, Any]]] = {}
        for table in expected["tables"]:
            for row in _table_rows(report.get_table(table)):
                if "feature" in row:
                    candidates.setdefault(row["feature"], []).append(row)
        for feature, values in parser.metrics:
            if not any(all(key in row and isinstance(row[key], (int, float))
                           and math.isclose(value, row[key], rel_tol=1e-7, abs_tol=1e-10)
                           for key, value in values.items()) for row in candidates.get(feature, [])):
                raise ValueError(f"Native HTML statistical values differ from saved evidence: {path.stem}/{feature}")


def _validate_excel_tables(workbook: Any, report: Any, label: str) -> None:
    """核对完整元数据和每张适用统计工作表，而非只抽交叉首行。"""
    description = report.describe()
    for name, expected in (
        ("BusinessContext", description["business_context"]),
        ("ReportSemantics", {key: value for key, value in description.items() if key not in {"business_context", "feature_metadata"}}),
    ):
        if name not in workbook.sheetnames:
            raise ValueError(f"Excel delivery lacks {name}: {label}")
        _assert_equal(expected, _excel_json_mapping(workbook[name], f"{label}.{name}"), f"Excel.{label}.{name}")
    metadata_rows = workbook["FeatureMetadata"].iter_rows(values_only=True)
    metadata_header = list(next(metadata_rows))
    expected_header: list[str] = [
        "feature", "display_name", "dtype", "data_source", "description", "unit",
    ]
    if metadata_header != expected_header:
        raise ValueError(f"Excel feature identity header differs: {label}")
    metadata = {row[0]: dict(zip(metadata_header[1:], row[1:])) for row in metadata_rows}
    source_types: dict[str, str] = description["parameters"].get("feature_dtypes", {})
    expected_metadata: dict[str, dict[str, Any]] = {
        feature: {**entry, "dtype": source_types.get(feature)}
        for feature, entry in description["feature_metadata"].items()
    }
    _assert_equal(expected_metadata, metadata, f"Excel.{label}.FeatureMetadata")
    for index, table in enumerate(description["tables"]):
        source = _table_rows(report.get_table(table))
        name = f"{index:03d}_{table}"[:31]
        if name not in workbook.sheetnames:
            name = table if table in workbook.sheetnames else {
                "overview": "Overview", "calculation_status": "CalculationStatus",
                "dq.missing": "DQ_Missing", "dq.zeros": "DQ_Zeros",
                "stats.mean": "Trend_Mean", "stats.psi": "Trend_Psi",
            }.get(table, "")
        if not name or name not in workbook.sheetnames:
            raise ValueError(f"Excel lacks a saved statistical table: {label}/{table}")
        rows = workbook[name].iter_rows(values_only=True)
        header = list(next(rows))
        fields = list(report.get_table(table, limit=0).columns)
        if not set(fields).issubset(header):
            raise ValueError(f"Excel statistical header differs: {label}/{table}")
        actual = [{field: values[header.index(field)] for field in fields} for values in rows]
        if table in {"overview", "dq.missing", "dq.zeros", "stats.mean", "stats.psi"}:
            source.sort(key=lambda row: row["feature"])
            actual.sort(key=lambda row: row["feature"])
        for expected_row, actual_row in zip(source, actual):
            for field, expected in expected_row.items():
                value = actual_row[field]
                if isinstance(expected, (list, dict)) and isinstance(value, str):
                    actual_row[field] = ast.literal_eval(value)
                elif expected == "" and value is None:
                    actual_row[field] = ""
        _assert_equal(source, actual, f"Excel.{label}.{table}")


def _validate_evidence(directory: Path) -> dict[str, Any]:
    """只加载已保存快照，重放七例中的实际公共查询与身份关联。"""
    from mars.reporting import load_report

    snapshots: dict[str, Any] = {
        report.report_id: report
        for path in directory.glob("*.marsreport")
        for report in [load_report(path)]
    }
    if not snapshots:
        raise ValueError("Cases contain no loadable .marsreport files")
    manifest: dict[str, Any] = _read_json(directory / "manifest.json")
    summary: dict[str, Any] = _read_json(directory / "summary.json")
    preparation = _validate_preparation(directory, manifest)
    for number in range(1, 8):
        case: dict[str, Any] = _read_json(directory / f"case-{number}.json")
        if str(case["case"]) != str(number) or case["report_id"] not in snapshots:
            raise ValueError(f"case-{number}: report identity is not saved in this batch")
        if not case["queries"] or not case["findings"] or not case["unavailable"]:
            raise ValueError(f"case-{number}: missing evidence, findings, or explicit limitation")
        for key, identity in case["findings"].items():
            if key.endswith("_report_id") and identity not in snapshots:
                raise ValueError(f"case-{number}: related {key} is not saved in this batch")
        for index, query in enumerate(case["queries"]):
            reference = query["reference"]
            if reference["report_id"] not in snapshots:
                raise ValueError(f"case-{number}.queries[{index}]: foreign report identity")
            report = snapshots[reference["report_id"]]
            replay = report.query_page(reference["table"], **reference["query"])
            _assert_equal(query["rows"], _table_rows(replay["data"]), f"case-{number}.query-{index}.rows")
            for field in ("total_rows", "omitted_rows", "next_offset"):
                if query[field] != replay[field]:
                    raise ValueError(f"case-{number}.query-{index}: {field} does not replay")
            if "returned_rows" in query and query["returned_rows"] != len(query["rows"]):
                raise ValueError(f"case-{number}.query-{index}: returned_rows differs from evidence")
        description = case["description"]
        if description["report_id"] != case["report_id"]:
            raise ValueError(f"case-{number}: human description and Agent evidence differ in identity")
        _assert_equal(description, snapshots[case["report_id"]].describe(), f"case-{number}.description")
        if (
            manifest["reports"][str(number)] != case["report_id"]
            or summary["cases"][str(number)]["report_id"] != case["report_id"]
        ):
            raise ValueError(f"case-{number}: manifest, human summary and query identities differ")
        _assert_equal(
            {"report_id": case["report_id"], **case["findings"]},
            summary["cases"][str(number)], f"summary.case-{number}",
        )
        compact: dict[str, Any] = _read_json(directory / f"query-{number}.json")
        _assert_equal(compact["query"], case["queries"][0], f"query-{number}.query")
        for field in ("case", "report_id", "question", "unavailable", "consumer"):
            _assert_equal(compact[field], case[field], f"query-{number}.{field}")
    context_path = directory / "agent-context.json"
    context_text = context_path.read_text(encoding="utf-8").rstrip("\n")
    bounded: dict[str, Any] = _read_json(context_path)
    restore: dict[str, Any] = _read_json(directory / "case-6.json")
    report = snapshots[restore["report_id"]]
    if len(context_text) > restore["findings"]["max_chars"]:
        raise ValueError("Bounded Agent JSON exceeds its declared Unicode character budget")
    if len(context_text) != restore["findings"]["context_chars"]:
        raise ValueError("Bounded Agent JSON character count differs from the saved finding")
    if bounded["description"]["report_id"] != report.report_id:
        raise ValueError("Bounded Agent JSON has a foreign report identity")
    for index, entry in enumerate(bounded["evidence"]):
        if entry["report_id"] != report.report_id:
            raise ValueError(f"Bounded evidence {index} has a foreign report identity")
        rows = _table_rows(report.get_table(entry["reference"], **entry["query"]))
        _assert_equal(entry["rows"], rows, f"bounded.evidence[{index}].rows")
    _validate_saved_denominators(preparation, snapshots)
    _validate_html_exports(directory)
    # 每张报告工作簿重组长元数据，再逐表对照同源快照的真实统计。
    import openpyxl

    for path in directory.glob("*.xlsx"):
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            if not workbook.sheetnames or not any(sheet.max_row for sheet in workbook):
                raise ValueError(f"Empty Excel delivery: {path.name}")
            snapshot_path = path.with_suffix(".marsreport")
            if snapshot_path.exists():
                _validate_excel_tables(workbook, load_report(snapshot_path), path.stem)
            elif path.name == "selection.xlsx":
                excel_rows = workbook.active.iter_rows(values_only=True)
                header = list(next(excel_rows))
                actual = [dict(zip(header, values)) for values in excel_rows]
                selected = _read_json(directory / "selection.json")
                for expected_row, actual_row in zip(selected, actual):
                    for field, value in expected_row.items():
                        if value == "" and actual_row[field] is None:
                            actual_row[field] = ""
                _assert_equal(selected, actual, "Excel.selection")
            elif path.name == "correlation.xlsx":
                from mars.reporting import show_correlation_matrix
                from mars.reporting._serialization import json_safe

                matrix = show_correlation_matrix(load_report(directory / "raw-correlation.marsreport")).data
                excel_rows = workbook.active.iter_rows(values_only=True)
                _assert_equal([None, *matrix.columns], list(next(excel_rows)), "Excel.correlation.header")
                source_rows = [[feature, *values.to_list()] for feature, values in matrix.iterrows()]
                _assert_equal(json_safe(source_rows), [list(values) for values in excel_rows], "Excel.correlation.matrix")
        finally:
            workbook.close()
    preview_path = directory / "preview-evidence.json"
    if preview_path.exists():
        preview: dict[str, Any] = _read_json(preview_path)
        reference = preview["reference"]
        report = snapshots.get(reference["report_id"])
        if report is None:
            raise ValueError("README preview references a foreign report batch")
        rows = _table_rows(report.query_page(reference["table"], **reference["query"])["data"])
        if len(rows) != 1:
            raise ValueError("README preview must cite exactly one queried cell")
        _assert_equal(preview["row"], rows[0], "preview.row")
        excerpt = preview["excerpt"]
        if excerpt["report_id"] != report.report_id or excerpt["table"] != reference["table"]:
            raise ValueError("README preview excerpt identity differs from the actual query")
        for field, value in excerpt.items():
            if field not in {"report_id", "table", "scope"}:
                _assert_equal(value, rows[0][field], f"preview.excerpt.{field}")
        for field, value in excerpt["scope"].items():
            _assert_equal(value, rows[0][field], f"preview.scope.{field}")
    _validate_native_binning_preview(directory, snapshots)
    automatic: dict[str, Any] = _read_json(directory / "rule-generation.json")
    report = snapshots.get(automatic["report_id"])
    if report is None:
        raise ValueError("Automatic rule evidence references an unsaved report")
    _assert_equal(automatic["description"], report.describe(), "automatic.description")
    _replay_queries(automatic["queries"], report, "automatic")
    if automatic["raw_input_rows"] != automatic["eligible_rows"] + automatic["excluded_rows"]:
        raise ValueError("Automatic rule input exclusions do not preserve the raw denominator")
    return snapshots


def _replay_queries(queries: list[dict[str, Any]], report: Any, label: str) -> None:
    """重放附加图表和生成器证据的实际分页，不用示意数字替代公共结果。"""
    if not queries:
        raise ValueError(f"{label}: no saved query evidence")
    for index, query in enumerate(queries):
        reference = query["reference"]
        if reference["report_id"] != report.report_id:
            raise ValueError(f"{label}.query-{index}: foreign report identity")
        replay = report.query_page(reference["table"], **reference["query"])
        _assert_equal(query["rows"], _table_rows(replay["data"]), f"{label}.query-{index}.rows")
        for field in ("total_rows", "returned_rows", "omitted_rows", "next_offset"):
            if query[field] != replay[field]:
                raise ValueError(f"{label}.query-{index}: {field} does not replay")


def _validate_native_binning_preview(directory: Path, snapshots: dict[str, Any]) -> None:
    """README 使用原生分箱图，公开图文件及其业务查询须来自同一快照。"""
    evidence: dict[str, Any] = _read_json(directory / "binning-native-evidence.json")
    report = snapshots.get(evidence["report_id"])
    binning_case: dict[str, Any] = _read_json(directory / "case-2.json")
    if report is None or evidence["report_id"] != binning_case["report_id"]:
        raise ValueError("Native binning image has no matching binning snapshot")
    metadata = report.describe()["feature_metadata"]
    if evidence["feature"] not in metadata:
        raise ValueError("Native binning image feature has no business metadata")
    _replay_queries(evidence["queries"], report, "native-binning")
    for query in evidence["queries"]:
        if any(row.get("feature") not in metadata for row in query["rows"]):
            raise ValueError("Native binning comparative evidence contains an unknown feature")
        if not any(row.get("feature") == evidence["feature"] for row in query["rows"]):
            raise ValueError("Native binning query does not cite the actual image feature")
    png = _safe_path(directory, evidence["chart_png"]).read_bytes()
    if not png.startswith(b"\x89PNG\r\n\x1a\n") or len(png) < 24:
        raise ValueError("Native binning PNG is not a readable PNG")
    width, height = struct.unpack(">II", png[16:24])
    if width < 500 or height < 300:
        raise ValueError(f"Native binning PNG is too small for its panels: {width} x {height}")
    offset = 8
    density: tuple[int, int, int] | None = None
    while offset + 12 <= len(png):
        length = struct.unpack(">I", png[offset:offset + 4])[0]
        kind = png[offset + 4:offset + 8]
        if kind == b"pHYs" and length == 9:
            density = struct.unpack(">IIB", png[offset + 8:offset + 17])
            break
        offset += length + 12
    if density is None or density[2] != 1 or min(density[:2]) * 0.0254 < 299.99:
        raise ValueError("Native PNG must preserve the actual at-least-300-dpi export")
    if evidence["parameters"]["dpi"] < 300:
        raise ValueError("Native image parameters do not record the required export DPI")
    svg = ElementTree.parse(_safe_path(directory, evidence["chart_svg"])).getroot()
    if svg.tag != "{http://www.w3.org/2000/svg}svg" or not list(svg.iter("{http://www.w3.org/2000/svg}path")):
        raise ValueError("Native binning SVG has no rendered chart paths")


def _validate_rule_memberships(
    directory: Path, config: dict[str, Any], data: Any = None,
) -> None:
    """显式原数据复算正常格规则成员；公开模式只验证保存统计。"""
    import polars as pl

    from mars.reporting import load_report
    from mars.rule import MarsRule, MarsRuleSet

    if data is None:
        raise ValueError("Rule membership verification requires explicit prepared inputs")
    cross = load_report(directory / "score-cross.marsreport")
    rules = load_report(directory / "rules.marsreport")
    description = rules.describe()
    origins = description["business_context"]["candidate_origins"]
    candidates = rules.get_table("candidates").to_dicts()
    cells = cross.get_table("cells").to_dicts()
    checked = 0
    for candidate in candidates:
        origin = origins.get(candidate["rule_id"], {})
        if origin.get("kind") != "development_cell":
            continue
        filters = origin["reference"]["query"]["filters"]
        rule = MarsRule(candidate["expression"])
        applied = MarsRuleSet(rules=[rule]).transform(data)
        hit_column = f"rule__{rule.rule_id}"
        for group in data["dataset_flag"].unique().to_list():
            hits = applied.filter((pl.col("dataset_flag") == group) & (pl.col(hit_column) == 1))
            normal = [row for row in cells if row["group"] == group
                      and row["target"] == filters["target"]
                      and row["x_bin"] == filters["x_bin"] and row["y_bin"] == filters["y_bin"]]
            for field, actual in (
                ("sample_count", hits.height),
                ("observed_sample_count", hits[filters["target"]].count()),
                ("bad_sample_count", int(hits[filters["target"]].sum() or 0)),
            ):
                expected = sum(row[field] for row in normal)
                if actual != expected:
                    raise ValueError(
                        f"Normal cell / DSL membership differs: {rule.rule_id}, {group}, "
                        f"{field}: cell={expected}, rule={actual}"
                    )
        checked += 1
    if not checked:
        raise ValueError("Rule case contains no normal-cell candidates to validate")


def _validate_manifest(directory: Path, site_dir: Path | None) -> dict[str, Any]:
    """核对 manifest 的实际文件、内容 hash、编码、打包成员和构建资源。"""
    attributes = (ROOT / ".gitattributes").read_text(encoding="utf-8")
    if not re.search(r"^docs/assets/cases/\*\*\s+-text\s*$", attributes, re.M):
        raise ValueError("Missing .gitattributes byte preservation for docs/assets/cases/** -text")
    manifest: dict[str, Any] = _read_json(directory / "manifest.json")
    if manifest["schema_version"] != 2 or not manifest["artifacts"]:
        raise ValueError("Case manifest has unsupported schema or no artifacts")
    for name in ("task_cases.py", "external_agent_rule_case.py", "home_credit_data.py"):
        delivered = _safe_path(directory, name)
        current = ROOT / "docs/snippets" / name
        if _code_digest(delivered) != _code_digest(current):
            raise ValueError(f"Downloadable source is stale: {name}; run --phase finalize")
    seen: set[str] = set()
    for artifact in manifest["artifacts"]:
        name = artifact["path"]
        if Path(name).suffix.lower() in {".csv", ".parquet", ".pickle", ".pkl"} or "human-writing" in name:
            raise ValueError(f"Raw data or personal Skill is forbidden in public downloads: {name}")
        if name in seen:
            raise ValueError(f"Duplicate artifact in manifest: {name}")
        seen.add(name)
        path = _safe_path(directory, name)
        data = path.read_bytes()
        if len(data) != artifact["size_bytes"]:
            raise ValueError(f"Artifact size differs from manifest: {name}")
        if hashlib.sha256(data).hexdigest() != artifact["sha256"]:
            raise ValueError(f"Artifact hash differs from manifest: {name}; run --phase finalize")
        if path.suffix in {".json", ".txt", ".md", ".py", ".html", ".svg"}:
            text = data.decode("utf-8")
            if re.search(r"[A-Za-z]:[\\/](?:Users|Desktop|my_download_program)[\\/]", text):
                raise ValueError(f"Private absolute path leaked in public artifact: {name}")
            if path.suffix == ".json":
                _read_json(path)
        if site_dir is not None:
            built = site_dir / "assets/cases" / name
            if not built.is_file() or built.read_bytes() != data:
                raise ValueError(f"Built site does not contain the real case download: {name}")
    archive_path = directory / "cases.zip"
    if not archive_path.is_file():
        raise ValueError("Missing shared download cases.zip")
    actual_files = {path.name for path in directory.iterdir() if path.is_file()}
    if actual_files != seen | {"manifest.json", "cases.zip"}:
        raise ValueError(f"Case files are absent from manifest: {sorted(actual_files - seen - {'manifest.json', 'cases.zip'})}")
    with zipfile.ZipFile(archive_path) as archive:
        names = archive.namelist()
        if len(names) != len(set(names)):
            raise ValueError("cases.zip contains duplicate members")
        for name in names:
            source = _safe_path(directory, name)
            if archive.read(name) != source.read_bytes():
                raise ValueError(f"cases.zip member differs from its source: {name}")
        required = (seen - {"cases.zip"}) | {"manifest.json"}
        if required != set(names):
            raise ValueError(f"cases.zip members differ from manifest: {sorted(required.symmetric_difference(names))}")
        if set(manifest["zip_allowlist"]) != required or len(manifest["zip_allowlist"]) != len(required):
            raise ValueError("ZIP members differ from the explicit, duplicate-free publication allowlist")
    if site_dir is not None:
        for name in ("manifest.json", "cases.zip"):
            built = site_dir / "assets/cases" / name
            if not built.is_file() or built.read_bytes() != (directory / name).read_bytes():
                raise ValueError(f"Built site download is missing or stale: {name}")
        for page in CASE_PAGES:
            if not (site_dir / "demos" / page / "index.html").is_file():
                raise ValueError(f"Built case page is missing: demos/{page}/")
    return manifest


def _recompute(
    directory: Path, manifest: dict[str, Any], input_arguments: list[str],
) -> None:
    """从显式三个原文件完整重算，比较参数、依赖、统计和有限查询。"""
    if not input_arguments:
        raise ValueError("Recompute mode requires --data-dir or three explicit CSV paths")
    config = manifest["data_config"]
    sys.path.insert(0, str(GENERATOR.parent))
    try:
        from home_credit_data import HomeCreditConfig, prepare_home_credit_data

        values = dict(zip(input_arguments[::2], input_arguments[1::2]))
        inputs = prepare_home_credit_data(
            values.get("--data-dir", "examples/home-credit-2018/data/raw_data"),
            train_csv=values.get("--train-csv"), test_csv=values.get("--test-csv"),
            columns_description_csv=values.get("--columns-description-csv"),
            config=HomeCreditConfig(seed=config["seed"], target_group_rows=config["target_group_rows"]),
            require_full_dataset=config["preparation"]["full_dataset"],
        )
    finally:
        sys.path.pop(0)
    _assert_equal(config["preparation"], inputs.summary, "recompute.preparation")
    with tempfile.TemporaryDirectory(prefix="mars-case-check-", dir=ROOT) as temporary:
        fresh = Path(temporary)
        execution: subprocess.CompletedProcess[bytes] = subprocess.run(
            [
                sys.executable, str(GENERATOR), "--case", "all", "--phase", "generate",
                "--seed", str(config["seed"]), "--output-dir", str(fresh), *input_arguments,
                "--target-group-rows", str(config["target_group_rows"]),
                *([] if config["preparation"]["full_dataset"] else ["--fixture"]),
            ],
            check=False, capture_output=True,
            env={**os.environ, "PYTHONIOENCODING": "utf-8"},
        )
        # Windows 文本 reader thread 不传播解码异常；在主线程严格解码保留失败证据。
        stdout = execution.stdout.decode("utf-8", errors="strict")
        stderr = execution.stderr.decode("utf-8", errors="strict")
        if execution.returncode:
            raise ValueError(f"Public case generation failed (exit {execution.returncode}):\n{stdout}\n{stderr}")
        names = ["summary.json", "agent-context.json", "binning-native-evidence.json", "rule-generation.json",
                 "preparation-summary.json", "feature-metadata.json", "selection.json"]
        names.extend(f"{prefix}-{number}.json" for prefix in ("case", "query") for number in range(1, 8))
        for name in names:
            _assert_equal(_semantic(_read_json(directory / name)), _semantic(_read_json(fresh / name)), name)
        _validate_rule_memberships(fresh, config, inputs.df)
        for path in directory.glob("*.marsreport"):
            from mars.reporting import load_report

            published = load_report(path)
            generated = load_report(fresh / path.name)
            _assert_equal(_semantic(published.describe()), _semantic(generated.describe()), path.name + ".catalog")
            for table in published.describe()["tables"]:
                _assert_equal(
                    _table_rows(published.get_table(table)), _table_rows(generated.get_table(table)),
                    path.name + "." + table,
                )
        _assert_equal(
            _semantic(_read_json(directory / "generation-environment.json")),
            _semantic(_read_json(fresh / "generation-environment.json")), "generation-environment",
        )
        _assert_equal(
            (directory / "previews.txt").read_text(encoding="utf-8"),
            (fresh / "previews.txt").read_text(encoding="utf-8"), "previews.txt",
        )
    print(f"Full CSV recomputation matches published statistics: rows={config['rows']}, seed={config['seed']}")


def main() -> None:
    """执行检查，任何陈旧或无效证据均以清晰错误退出非零。"""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--assets-dir", type=Path, default=ASSETS)
    parser.add_argument("--site-dir", type=Path, default=None)
    parser.add_argument("--mode", choices=["public", "recompute"], default="public")
    parser.add_argument("--data-dir", type=Path)
    parser.add_argument("--train-csv", type=Path)
    parser.add_argument("--test-csv", type=Path)
    parser.add_argument("--columns-description-csv", type=Path)
    args = parser.parse_args()
    try:
        manifest = _validate_manifest(args.assets_dir, args.site_dir)
        for name in ("previews.txt", "binning-native-home-credit.png", "binning-native-home-credit.svg"):
            _safe_path(args.assets_dir, name)
        snapshots = _validate_evidence(args.assets_dir)
        input_arguments = [
            argument for name in ("data_dir", "train_csv", "test_csv", "columns_description_csv")
            if getattr(args, name) is not None
            for argument in ("--" + name.replace("_", "-"), str(getattr(args, name)))
        ]
        if args.mode == "recompute":
            _recompute(args.assets_dir, manifest, input_arguments)
        elif input_arguments:
            raise ValueError("CSV arguments require --mode recompute; public mode never reads raw files")
        print(f"Case assets passed ({args.mode}): {len(manifest['artifacts'])} files, {len(snapshots)} snapshots, seven cases")
    except (ValueError, OSError, KeyError, subprocess.CalledProcessError, zipfile.BadZipFile) as error:
        parser.exit(1, f"Case asset check failed: {error}\n")


if __name__ == "__main__":
    main()
