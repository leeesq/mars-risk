"""报告数据区域的 Excel 字面文字写入，保留数值和模板公式。"""

from __future__ import annotations

from datetime import date, datetime, timezone
from typing import Any

import pandas as pd


def _xlwings_value(value: Any, *, com_dates: bool = True) -> Any:
    """转义 Excel 会消费的首引号；COM 日期副本用 UTC 标记保留墙钟时间。"""
    # Excel 将首个单引号当作文字输入标记；传输副本加一个，单元格仍保存原文。
    if isinstance(value, str) and value.startswith("'"):
        return "'" + value
    if not com_dates:
        return value
    if isinstance(value, pd.Timestamp):
        value = value.to_pydatetime()
    if isinstance(value, datetime):
        return value.replace(tzinfo=timezone.utc)
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day, tzinfo=timezone.utc)
    return value


def _write_openpyxl_value(cell: Any, value: Any) -> None:
    """显式清空缺失单元格，并让等号开头的业务字符串仍保存为原文。"""
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"


def _preserve_openpyxl_strings(worksheet: Any, frame: pd.DataFrame) -> None:
    """只校正刚由 Pandas 写入的表头及数据字符串，不扫描模板区域。"""
    for column_index, column in enumerate(frame.columns, start=1):
        _write_openpyxl_value(worksheet.cell(1, column_index), column)
    for row_index, row in enumerate(frame.itertuples(index=False, name=None), start=2):
        for column_index, value in enumerate(row, start=1):
            if isinstance(value, str):
                _write_openpyxl_value(worksheet.cell(row_index, column_index), value)


def _prepare_xlwings_strings(
    worksheet: Any, rows: list[list[Any]], *, start_row: int, first_col: int,
) -> None:
    """在写值前只给字符串连续区间设置文字格式，防止 Excel 自动识别。"""
    from openpyxl.utils import get_column_letter

    if not rows:
        return
    for column_offset in range(len(rows[0])):
        letter = get_column_letter(first_col + column_offset)
        run_start: int | None = None
        # 哨兵结束末段，连续文字只需一次原生调用；数值与日期格式不受影响。
        for row_offset in range(len(rows) + 1):
            is_text = row_offset < len(rows) and isinstance(rows[row_offset][column_offset], str)
            if is_text and run_start is None:
                run_start = start_row + row_offset
            elif not is_text and run_start is not None:
                worksheet.range(f"{letter}{run_start}:{letter}{start_row + row_offset - 1}").number_format = "@"
                run_start = None
