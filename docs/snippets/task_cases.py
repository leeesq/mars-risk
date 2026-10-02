"""七个任务案例的共享确定性生成入口；仅公开合成统计，不调用 LLM。"""

from __future__ import annotations

import argparse
import ast
import hashlib
import importlib.metadata
import json
import math
import re
import subprocess
import sys
import zipfile
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
import polars as pl
from external_agent_rule_case import DIRECTIONS, METADATA, _sample, produce
from numpy.typing import NDArray

import mars
from mars.analysis import (
    cross_scores,
    evaluate_score_policy,
    get_score_bin_definitions,
    profile_risk,
    profile_stats,
    write_score_cross_html,
)
from mars.feature import MarsLinearSelector, MarsStatsSelector
from mars.reporting import Report, load_report, show_correlation_matrix
from mars.reporting._serialization import json_safe

DEFAULT_SEED = 20261001
DEFAULT_ROWS = 18000
FEATURES = ["main_score", "aux_score", "income", "channel"]
DICTIONARY: dict[str, dict[str, Any]] = {
    **METADATA,
    "income": {
        "display_name": "申报月收入", "data_source": "application", "unit": "CNY/month",
        "description": "模拟申报月收入；验证、观察期施加可复算分布变化",
    },
    "income_copy": {
        "display_name": "收入冗余副本", "data_source": "application", "unit": "CNY/month",
        "description": "income 的两倍；用于 raw 相关冗余审计",
    },
    "score_inverse": {
        "display_name": "主分负向副本", "data_source": "champion", "unit": "score_points",
        "description": "main_score 的相反数；raw 空间相关系数为 -1",
    },
    "channel": {
        "display_name": "申请渠道", "data_source": "application", "unit": "category",
        "description": "观察期出现仅在该期可见的 branch 渠道",
    },
    "constant": {"display_name": "常量字段", "data_source": "application"},
    "sparse": {"display_name": "高缺失字段", "data_source": "application"},
    "aux_probability": {
        "display_name": "辅助概率检查字段", "data_source": "challenger", "unit": "fraction",
        "description": "由辅助分转换，含显式非法概率和特殊值；仅用于概率域审计",
    },
}
BASELINE: dict[str, Any] = {
    "type": "x_only", "x_max_risk_rank": 2, "missing_score": "reject",
}
POLICY: dict[str, Any] = {
    "type": "expression", "expression": "X <= X2 AND Y <= Y2", "missing_score": "reject",
}


def _write_json(path: Path, value: Any) -> None:
    """写入可解析 UTF-8 JSON，非有限数交给报告公共编码契约。"""
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False) + "\n",
                    encoding="utf-8")


def _rows(frame: Any) -> list[dict[str, Any]]:
    """仅物化已经分页或明确受限的证据，复用报告公开 JSON 的序列化结果。"""
    rows = frame.to_dicts() if isinstance(frame, pl.DataFrame) else frame.to_dict("records")
    return json_safe(rows)


def _query(report: Report, table: str, question: str, **query: Any) -> dict[str, Any]:
    """保存实际分页结果及其可重放引用，不把省略行冒充完整证据。"""
    page: dict[str, Any] = report.query_page(table, **query)
    return {
        "question": question,
        "reference": page["reference"],
        "total_rows": page["total_rows"],
        "returned_rows": page["returned_rows"],
        "omitted_rows": page["omitted_rows"],
        "next_offset": page["next_offset"],
        "rows": _rows(page["data"]),
    }


def _case(
    output: Path, number: int, report: Report, question: str,
    queries: list[dict[str, Any]], findings: dict[str, Any], unavailable: str,
) -> dict[str, Any]:
    """人工摘要与 Agent 查询绑定同一报告；无法回答的信息始终单独声明。"""
    payload: dict[str, Any] = {
        "case": number, "report_id": report.report_id, "question": question,
        "consumer": "deterministic public API queries; no LLM",
        "description": report.describe(), "queries": queries,
        "findings": findings, "unavailable": unavailable,
    }
    _write_json(output / f"case-{number}.json", payload)
    _write_json(output / f"query-{number}.json", {
        "case": number, "consumer": payload["consumer"], "report_id": report.report_id,
        "question": question, "query": queries[0], "unavailable": unavailable,
    })
    return {"report_id": report.report_id, **findings}


def _data(rows: int, seed: int) -> pl.DataFrame:
    """扩展已有旗舰案例的同一生成器，保留九个月与独立行的三个分区。"""
    data: pl.DataFrame = _sample(rows, seed)
    rng = np.random.default_rng(seed + 1)
    partition: NDArray[np.int64] = np.arange(rows, dtype=np.int64) // (rows // 3)
    income: NDArray[np.float64] = rng.lognormal(8.4, 0.45, rows) * (1 + 0.35 * partition)
    income[::131] = -999
    income[::173] = np.nan
    channel = rng.choice(["app", "web"], rows).astype(object)
    channel[(partition == 2) & (np.arange(rows) % 7 == 0)] = "branch"
    probability = 1 / (1 + np.exp(-(data["aux_score"].to_numpy() - 500) / 100))
    probability[::101] = 1.2
    probability[::157] = -999
    # 非有限模型分进入交叉 invalid 箱；不会转成缺失或正常风险等级。
    auxiliary = data["aux_score"].to_numpy().copy()
    auxiliary[::281] = np.inf
    return data.with_columns(
        pl.Series("aux_score", auxiliary).fill_nan(None),
        pl.Series("income", income).fill_nan(None),
        pl.Series("channel", channel.tolist(), dtype=pl.String),
        pl.Series("aux_probability", probability).fill_nan(None),
        pl.Series("weight", rng.uniform(0.6, 1.4, rows)),
        pl.lit(1.0).alias("constant"),
        pl.when(pl.int_range(pl.len()) % 50 == 0).then(1.0).otherwise(None).alias("sparse"),
    ).with_columns(
        (2 * pl.col("income")).alias("income_copy"),
        (-pl.col("main_score")).alias("score_inverse"),
        pl.when(pl.col("dataset") == "observation").then(None)
        .otherwise(pl.col("late60")).alias("late60"),
    )


def _context(data: pl.DataFrame, seed: int) -> dict[str, Any]:
    """明确模拟标签、样本、单位和观察边界，避免按字段名推断业务。"""
    return {
        "dataset_id": f"synthetic-consumer-credit-{seed}", "simulated": True,
        "sample_unit": "one synthetic application; unique sample_id",
        "sample_count": data.height, "seed": seed, "currency": "CNY",
        "labels": {
            "bad30": {"definition": "模拟30日违约", "performance_window": "30天"},
            "late60": {"definition": "模拟60日逾期", "performance_window": "60天"},
        },
        "splits": {
            "discovery": "1–3月；参考分箱与候选发现",
            "validation": "4–6月；独立验证，不重新拟合分箱或选择候选",
            "observation": "7–9月；bad30 有模拟标签，late60 全未表现",
        },
        "amount_definition": "模拟申请金额 CNY；不是损失、利润或预期收益",
        "weight_definition": "模拟分析权重；与原始申请人数分别保存",
        "score_directions": DIRECTIONS,
        "binning_origin": "discovery only; frozen definitions",
    }


def _export(report: Any, output: Path, stem: str) -> None:
    """仅调用真实报告支持的导出方法，快照恢复边界见公共报告目录。"""
    report.save(output / f"{stem}.marsreport", overwrite=True)
    report.write_html(str(output / f"{stem}.html"))
    report.write_excel(str(output / f"{stem}.xlsx"))


# --8<-- [start:score_cross]
def _score_cross(
    data: pl.DataFrame, output: Path, context: dict[str, Any],
) -> Report:
    """只在发现样本拟合分段，再在固定月度范围计算联合证据和策略回放。"""
    reference: pl.DataFrame = data.filter(pl.col("dataset") == "discovery")
    discovery = cross_scores(
        reference, score_x="main_score", score_y="aux_score", targets=["bad30", "late60"],
        score_directions=DIRECTIONS, n_bins=4, special_values={"aux_score": [-999]},
        feature_metadata=DICTIONARY, business_context=context,
    )
    definitions: dict[str, Any] = get_score_bin_definitions(discovery)
    cross = cross_scores(
        data, score_x="main_score", score_y="aux_score", targets=["bad30", "late60"],
        score_directions=DIRECTIONS, bin_definitions=definitions,
        group_col="dataset", time_col="application_date", time_grain="month",
        weights_col="weight", amount_col="amount", min_observed=30,
        feature_metadata=DICTIONARY, business_context=context,
    )
    cross.save(output / "score-cross.marsreport", overwrite=True)
    policy = evaluate_score_policy(cross, POLICY, baseline=BASELINE)
    policy.save(output / "policy.marsreport", overwrite=True)
    policy.write_excel(str(output / "policy.xlsx"))
    write_score_cross_html(
        cross, output / "score-cross.html", policy_reports=[policy],
        report_name="模拟消费信贷：同一主分层内的辅助风险梯度",
    )
    cross.write_excel(str(output / "score-cross.xlsx"))
    scope: dict[str, Any] = {"group": "discovery", "period": "202601", "target": "bad30"}
    queries: list[dict[str, Any]] = [
        _query(cross, "cells", "主模型 b1 内，辅助分的风险梯度和分母是什么？",
               filters={**scope, "x_bin": "b1"}, sort_by="y_risk_rank",
               columns=["group", "period", "target", "x_bin", "y_bin", "x_risk_rank",
                        "y_risk_rank", "sample_count", "observed_sample_count", "bad_sample_count",
                        "weight_sum", "observed_weight_sum", "bad_weight_sum", "bad_rate", "row_bad_rate",
                        "delta_vs_row", "lift_vs_overall", "status"], limit=10),
        _query(cross, "row_summary", "反向查看各 X 分层的边际风险；边际包含另一轴特殊箱",
               filters=scope, sort_by="x_risk_rank", limit=10),
        _query(cross, "column_summary", "辅助分固定后，X 方向的风险梯度是什么？",
               filters=scope, sort_by="y_risk_rank", limit=10),
        _query(cross, "cells", "观察期 late60 全未表现，不能据此判断模型效果",
               filters={"group": "observation", "period": "202607", "target": "late60"},
               columns=["x_bin", "y_bin", "sample_count", "observed_sample_count", "bad_rate",
                        "status"], limit=6),
        _query(policy, "changes", "固定正常分箱表达式相对 X-only 基线的历史样本变化", limit=4),
    ]
    normal = [r for r in queries[0]["rows"] if r["y_bin"].startswith("b")]
    states = cross.get_table("cells").group_by("status").len().sort("status").to_dicts()
    findings: dict[str, Any] = {
        "scope": scope, "x_bin": "b1", "normal_cells": normal, "status_counts": states,
        "input_row_count": data.height, "bad_rate_denominator": "observed_weight_sum",
        "ratio_unit": "fraction; delta shown as pp = 100 * delta_vs_row",
        "bin_definitions": definitions, "policy_report_id": policy.report_id,
    }
    _case(output, 4, cross, "主模型同等级内，辅助分还能进一步区分风险吗？", queries, findings,
          "这些是模拟关联与历史留存回放；无法证明新模型整体更好或真实因果增量。")
    return cross
# --8<-- [end:score_cross]


# --8<-- [start:profile]
def _profile(data: pl.DataFrame, output: Path, context: dict[str, Any]) -> None:
    """画像与案例侧 schema/category 前置审计并列；不虚构自动清洗能力。"""
    reference = data.filter(pl.col("dataset") == "discovery")
    report = profile_stats(
        data, features=[*FEATURES, "sparse", "aux_probability"],
        metrics=["missing", "zeros", "mean", "psi"], benchmark_df=reference,
        categorical_features=["channel"], group_col="dataset", special_values=[-999],
        feature_metadata=DICTIONARY, business_context=context,
    )
    _export(report, output, "profile")
    labels = data.group_by("dataset", maintain_order=True).agg(
        pl.len().alias("sample_count"),
        pl.col("bad30").count().alias("bad30_observed"),
        pl.col("late60").count().alias("late60_observed"),
    ).to_dicts()
    observed = data.filter(pl.col("dataset") == "observation")
    ref_categories = set(reference["channel"].unique())
    unseen = sorted(set(observed["channel"].unique()) - ref_categories)
    # schema 差异只是显式案例前置检查，不冒充 profile 表字段。
    schema_probe = observed.drop("income_copy").with_columns(pl.col("main_score").cast(pl.String))
    schema_difference: dict[str, Any] = {
        "consumer": "case-side prerequisite check; not a MARS automatic cleaning API",
        "missing_columns": sorted(set(reference.columns) - set(schema_probe.columns)),
        "dtype_differences": {name: [str(dtype), str(schema_probe.schema[name])]
                              for name, dtype in reference.schema.items()
                              if name in schema_probe.schema and dtype != schema_probe.schema[name]},
    }
    queries: list[dict[str, Any]] = [
        _query(report, "overview", "首先检查哪些字段的数据质量？", limit=8),
        _query(report, "dq.missing", "各分区缺失率如何？", limit=8),
    ]
    _case(output, 1, report, "这份数据能直接用吗？", queries,
          {"input_row_count": data.height, "labels_by_split": labels,
           "unseen_observation_categories": unseen, "schema_probe": schema_difference,
           "invalid_probability_count": data.filter(pl.col("aux_probability") > 1).height},
          "观察期 late60 没有有效标签，画像和 PSI 无法判断这个目标的模型效果下降。")
# --8<-- [end:profile]


# --8<-- [start:binning]
def _binning(data: pl.DataFrame, output: Path, context: dict[str, Any]) -> None:
    """参考样本决定多目标共用分箱，独立验证和观察只消费固定边界。"""
    reference = data.filter(pl.col("dataset") == "discovery")
    run = profile_risk(
        data, target=["bad30", "late60"], features=FEATURES, benchmark_df=reference,
        method="quantile", n_bins=4, special_values=[-999], group_col="dataset",
        weights_col="weight", amount_col="amount", feature_metadata=DICTIONARY,
        business_context=context, n_jobs=1,
    )
    report = run.report
    # 原 Excel 模板需原生刷新；此处公共 snapshot 导出静态表，不带旧透视缓存。
    report.save(output / "binning.marsreport", overwrite=True)
    report.write_html(str(output / "binning.html"), include_charts=False)
    load_report(output / "binning.marsreport").write_excel(str(output / "binning.xlsx"))
    _binning_evidence(report, output, reference.height)


def _binning_evidence(report: Report, output: Path, reference_rows: int) -> None:
    """从同一报告的真实固定分箱证据生成分区对比和风险曲线；也支持快照重读。"""
    queries: list[dict[str, Any]] = [
        _query(report, "summary", "哪些特征同时有区分度和稳定性证据？",
               sort_by="iv", descending=True, limit=8),
        _query(report, "detail", "收入在多个分区和目标下的固定分箱风险表现如何？",
               features="income", sort_by="bin_index", limit=12),
        _query(report, "calculation_status", "未表现目标以什么状态保存？",
               filters={"target": "late60", "group": "observation"}, limit=8),
    ]
    for metric in ["iv", "ks", "psi"]:
        queries.append(_query(report, f"trend.{metric}", f"bad30 的各分区 {metric} 如何对比？",
                              features=FEATURES, limit=4))
    risk_query = _query(
        report, "detail", "辅助分正常箱在开发/验证的固定边界与加权风险如何比较？",
        features="aux_score",
        filters={"y": "bad30", "mars_group": {"op": "in", "value": ["discovery", "validation"]},
                 "bin_index": {"op": "in", "value": [0, 1, 2, 3]}},
        columns=["y", "feature", "mars_group", "bin_index", "bin_label", "count",
                 "observed_count", "bad", "bad_rate", "lift"], sort_by=["mars_group", "bin_index"],
        limit=8,
    )
    queries.append(risk_query)
    # 静态科学图逐点读取已计算 detail，不在图中拟合新边界或改变指标。
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    figure, axes = plt.subplots(figsize=(8.2, 3.9), constrained_layout=True)
    for group, color in [("discovery", "#7755b8"), ("validation", "#e29738")]:
        values = [row for row in risk_query["rows"] if row["mars_group"] == group]
        axes.plot([row["bin_index"] for row in values],
                  [100 * row["bad_rate"] for row in values], label=group,
                  color=color, marker="o", linewidth=2)
    axes.set(title="Auxiliary score: frozen discovery bins, target bad30",
             xlabel="Normal bin index (higher aux_score = higher risk)",
             ylabel="Weighted event rate (%)", xticks=[0, 1, 2, 3])
    axes.legend(loc="upper left")
    axes.grid(axis="y", alpha=0.18)
    figure.savefig(output / "binning-risk.png", dpi=150)
    plt.close(figure)
    _case(output, 2, report, "哪些特征有区分度，而且足够稳定？", queries,
          {"summary_rows": queries[0]["rows"], "status_rows": queries[2]["rows"],
           "reference_row_count": reference_rows, "targets": ["bad30", "late60"],
           "split_metrics": {metric: queries[index]["rows"]
                             for index, metric in enumerate(["iv", "ks", "psi"], start=3)},
           "auxiliary_bin_risk": risk_query["rows"],
           "units": {"ks": "percentage points on 0–100 scale", "iv": "dimensionless",
                     "psi": "dimensionless", "bad_rate": "fraction"}},
          "分箱 IV/KS/PSI 不等于最终业务价值或因果增量；观察期 late60 未表现。")
# --8<-- [end:binning]


# --8<-- [start:selection]
def _selection(data: pl.DataFrame, output: Path, context: dict[str, Any]) -> None:
    """选择与相关性复用各自真实审计；raw 与 WOE 从不混成一份矩阵。"""
    discovery: pl.DataFrame = data.filter(pl.col("dataset") == "discovery")
    candidates = ["main_score", "score_inverse", "income", "income_copy", "constant", "sparse"]
    stats = MarsStatsSelector(
        missing_thr=0.9, rough_iv_thr=-1, rough_lift_thr=0, skip_fine_scan=True,
        psi_thr=None, rc_thr=None, corr_thr=0.8,
        rough_binning_params={"method": "quantile", "n_bins": 4}, n_jobs=1,
    ).fit(discovery, target="bad30", features=candidates, feature_metadata=DICTIONARY,
          business_context=context)
    selection = stats.get_report()
    selection.write_excel(output / "selection.xlsx") if hasattr(selection, "write_excel") else (
        pd.DataFrame(_rows(selection)).to_excel(output / "selection.xlsx", index=False)
    )
    _write_json(output / "selection.json", _rows(selection))
    woe = stats.get_correlation_report()
    woe.save(output / "woe-correlation.marsreport", overwrite=True)
    numeric = ["main_score", "score_inverse", "income", "income_copy"]
    linear = MarsLinearSelector(corr_thr=0.95, n_jobs=1).fit(
        discovery.select(numeric), discovery["bad30"], features=numeric,
        feature_metadata=DICTIONARY, business_context=context,
    )
    raw = linear.get_correlation_report()
    raw.save(output / "raw-correlation.marsreport", overwrite=True)
    matrix = show_correlation_matrix(raw)
    (output / "correlation.html").write_text(matrix.to_html(), encoding="utf-8")
    matrix.to_excel(output / "correlation.xlsx")
    queries: list[dict[str, Any]] = [
        _query(raw, "correlation_decisions", "raw 空间删除了哪些冗余特征？", limit=8),
        _query(raw, "pairs", "正负相关与绝对相关阈值如何区分？", limit=8),
        _query(woe, "correlation_decisions", "目标感知 WOE 空间的决定是什么？", limit=8),
    ]
    _case(output, 3, raw, "为什么保留这个特征，删除另一个？", queries,
          {"candidate_features": candidates, "stats_selected": stats.selected_features_,
           "stats_audit": _rows(selection), "linear_selected": linear.selected_features_,
           "woe_report_id": woe.report_id, "raw_pairs": queries[1]["rows"],
           "raw_parameters": raw.describe()["parameters"],
           "woe_parameters": woe.describe()["parameters"]},
          "相关性不能说明业务因果；raw 与 WOE 有不同表示、候选集及标签口径。")
# --8<-- [end:selection]


# --8<-- [start:rules]
def _rules(data: pl.DataFrame, cross: Report, output: Path) -> None:
    """复用已有旗舰规则发现/独立验证流程，完整保留生产规格的拒绝与不足状态。"""
    produce(output, data=data, cross_report=cross)
    report = load_report(output / "rules.marsreport")
    queries: list[dict[str, Any]] = [
        _query(report, "summary", "生产规格实际保留多少候选？", limit=1),
        _query(report, "candidates", "每个候选保留或拒绝的原因是什么？", limit=10),
        _query(report, "evaluation", "独立验证主目标的覆盖和风险是什么？",
               filters={"dataset": "validation", "target": "bad30", "group": "hit"}, limit=10),
        _query(report, "rules", "可审查规则的条件和身份是什么？", limit=10),
    ]
    _case(output, 5, report, "从候选规则走到可以审查的证据", queries,
          {"summary_rows": queries[0]["rows"], "candidates": queries[1]["rows"],
           "validation": queries[2]["rows"], "cross_report_id": cross.report_id,
           "raw_data_saved": False, "rule_status": "Experimental"},
          "报告不代表生产策略批准；没有规则 observation 客群评估或高级分析，不能推断其结果。")
# --8<-- [end:rules]


# --8<-- [start:restore]
def _consume(output: Path) -> None:
    """独立新进程仅加载快照、查询与聚合回放；不调用生成器或访问原宽表。"""
    report = load_report(output / "score-cross.marsreport")
    original = json.loads((output / "case-4.json").read_text(encoding="utf-8"))
    assert report.report_id == original["report_id"]
    description: dict[str, Any] = report.describe()
    for key in ["feature_metadata", "business_context", "parameters", "tables"]:
        assert description[key] == original["description"][key], key
    for query in original["queries"]:
        reference = query["reference"]
        if reference["report_id"] == report.report_id:
            assert _rows(report.get_table(reference["table"], **reference["query"])) == query["rows"]
    scope = {"group": "discovery", "period": "202601", "target": "bad30"}
    first = _query(report, "cells", "发现目录后的第一页", filters=scope,
                   columns=["x_bin", "y_bin", "sample_count", "bad_rate", "status"], limit=3)
    second = _query(report, "cells", "使用真实 next_offset 读取下一页", filters=scope,
                    columns=["x_bin", "y_bin", "sample_count", "bad_rate", "status"],
                    offset=first["next_offset"], limit=3)
    empty = _query(report, "cells", "合法查询但没有该分区", filters={"group": "not-present"}, limit=3)
    try:
        report.get_table("cells", columns=["not_a_column"], limit=1)
    except ValueError as exc:
        invalid_request = {"exception": "ValueError", "message": str(exc)}
    else:
        raise AssertionError("无效列请求必须失败，不能返回假空表。")
    # 最终字符预算包含目录与省略信息；保留裁剪后生效查询供重放。
    context = report.to_ai_context(
        queries={"cells": {"filters": scope, "columns": ["x_bin", "y_bin", "sample_count",
                                                          "bad_rate", "status"], "limit": 40}},
        max_chars=9000,
    )
    bounded: dict[str, Any] = json.loads(context)
    assert len(context) <= 9000
    for entry in bounded["evidence"]:
        assert _rows(report.get_table(entry["reference"], **entry["query"])) == entry["rows"]
    (output / "agent-context.json").write_text(context + "\n", encoding="utf-8")
    replay = evaluate_score_policy(report, POLICY, baseline=BASELINE)
    policy = load_report(output / "policy.marsreport")
    assert _rows(replay.get_table("changes")) == _rows(policy.get_table("changes"))
    _case(output, 6, report, "分析完成后，如何保存并继续查询？", [first, second, empty],
          {"snapshot_type": type(report).__name__, "identity_preserved": True,
           "separate_process": True, "context_chars": len(context), "max_chars": 9000,
           "budget_omissions": bounded.get("omissions", bounded.get("omitted", [])),
           "invalid_request": invalid_request, "policy_replay_matches": True,
           "table_names": list(report.describe()["tables"])},
          "ReportSnapshot 不恢复原分析器、原始宽表、任意新分段或 RuleSet 部署资格。")
# --8<-- [end:restore]


# --8<-- [start:delivery]
def _delivery(output: Path) -> None:
    """同一份交叉报告交付五种真实格式，并核对静态 Excel 的关键数值。"""
    import openpyxl

    report = load_report(output / "score-cross.marsreport")
    workbook = openpyxl.load_workbook(output / "score-cross.xlsx", read_only=True, data_only=True)
    sheet = workbook["000_cells"]
    rows = sheet.iter_rows(values_only=True)
    header = list(next(rows))
    excel_sample = next(rows)
    table_sample = report.get_table("cells", limit=1).row(0, named=True)
    for name in ["sample_count", "observed_sample_count", "bad_sample_count", "bad_rate"]:
        expected = table_sample[name]
        actual = excel_sample[header.index(name)]
        assert actual == expected or (
            actual is not None and expected is not None and np.isclose(actual, expected)
        ), name
    workbook.close()
    prompt = (
        "你是外部分析消费者。此任务无需内置 Agent、LLM Key 或原始宽表。\n"
        "1. 用 load_report('score-cross.marsreport') 读取可信输入。\n"
        "2. describe() 发现实际表、单位、状态和业务上下文。\n"
        "3. query_page('cells', filters={'group':'discovery','period':'202601','target':'bad30'}, "
        "limit=3)；按 next_offset 继续，引用实际 reference。\n"
        "4. to_ai_context(max_chars=9000, queries=...) 预算单位是 Unicode 字符；"
        "只引用裁剪后 evidence 的有效 query。\n"
        "5. 区分全样本、已表现标签、权重、金额；delta 比例乘 100 才是 pp。\n"
        "6. 缺失、invalid、低样本、未表现与有效零值分别回答；未知信息明确说不能回答。\n"
        "7. 模拟关联不是因果、实际收益或部署批准。快照不恢复原分析器或原数据。\n"
    )
    (output / "external-agent-task.txt").write_text(prompt, encoding="utf-8")
    formats: list[dict[str, str]] = [
        {"format": "HTML", "path": "score-cross.html", "use": "当前矩阵、选格和正常箱表达式交互；离线请求由浏览器验收"},
        {"format": "Excel", "path": "score-cross.xlsx", "use": "静态工作表；不提供 HTML 等价交互或计算"},
        {"format": ".marsreport", "path": "score-cross.marsreport", "use": "报告目录、查询、聚合 policy 回放；不是原始数据或分析器"},
        {"format": "bounded JSON", "path": "agent-context.json", "use": "真实目录和预算内 evidence，最终 Unicode 字符计数"},
        {"format": "TXT", "path": "external-agent-task.txt", "use": "外部消费者提示材料；不声称实际 LLM 分析"},
    ]
    _case(output, 7, report, "同一次分析，如何交付给不同使用者？",
          [_query(report, "cells", "Excel 与快照的首行关键值对齐", limit=1)],
          {"formats": formats, "excel_sheets": workbook.sheetnames,
           "excel_key_values_match": True, "shared_cases": [4, 6, 7]},
          "各格式用途不同；静态 Excel 不恢复交互或计算，快照不恢复任意 Python 对象。")
# --8<-- [end:delivery]


def _fingerprint(path: Path) -> str:
    """AST 去除说明文本与位置信息，注释/排版调整不会触发全量证据失效。"""
    tree = ast.parse(path.read_text(encoding="utf-8"))
    for node in ast.walk(tree):
        if isinstance(node, (ast.Module, ast.FunctionDef, ast.AsyncFunctionDef, ast.ClassDef)):
            if node.body and isinstance(node.body[0], ast.Expr):
                value = node.body[0].value
                if isinstance(value, ast.Constant) and isinstance(value.value, str):
                    node.body.pop(0)
    return hashlib.sha256(ast.dump(tree, include_attributes=False).encode("utf-8")).hexdigest()


def _environment() -> dict[str, Any]:
    """区分实际源码版本与已安装 distribution 元数据，不公开本机解释器或目录。"""
    return {
        "python": sys.version.split()[0],
        "imported_mars_version": mars.__version__,
        "installed_distribution_version": importlib.metadata.version("mars-risk"),
        "import_source": "repository src/mars checkout",
        "dependencies": {name: importlib.metadata.version(name)
                         for name in ["numpy", "pandas", "polars", "scikit-learn",
                                      "pyarrow", "openpyxl", "xlsxwriter"]},
    }


def _previews(output: Path) -> None:
    """首屏小表仅从冻结案例 JSON 生成，正文与卡片复用同一数值来源。"""
    payloads: dict[int, dict[str, Any]] = {
        number: json.loads(path.read_text(encoding="utf-8"))
        for number in range(1, 8)
        if (path := output / f"case-{number}.json").exists()
    }
    blocks: dict[str, str] = {}

    def table(headers: list[str], values: list[list[Any]]) -> str:
        """将受限结果行渲染为 Markdown 表，保留缺失并转义自然语言分隔符。"""
        lines = ["| " + " | ".join(headers) + " |", "| " + " | ".join(["---"] * len(headers)) + " |"]
        for row in values:
            cells = [str(value).replace("|", "\\|").replace("\n", " ") for value in row]
            lines.append("| " + " | ".join(cells) + " |")
        return "\n".join(lines)

    def number(value: float | None, scale: float = 1, suffix: str = "") -> str:
        """按展示单位格式化真实有限数，未计算值保留不可用标识。"""
        return "—" if value is None else f"{value * scale:.4f}{suffix}"

    for case, payload in payloads.items():
        finding = payload["findings"]
        if case == 1:
            rows = [[row["dataset"], row["sample_count"], row["bad30_observed"], row["late60_observed"]]
                    for row in finding["labels_by_split"]]
            blocks["case1"] = table(["分区", "全样本人数", "bad30 已表现人数", "late60 已表现人数"], rows)
            blocks["card1"] = table(["总样本", "观察期 late60 已表现", "新类别"],
                [[finding["input_row_count"], finding["labels_by_split"][-1]["late60_observed"],
                  ", ".join(finding["unseen_observation_categories"])]])
        elif case == 2:
            metrics = finding["split_metrics"]
            rows = []
            for feature in FEATURES:
                evidence = {metric: next(row for row in values if row["feature"] == feature)
                            for metric, values in metrics.items()}
                rows.append([feature, number(evidence["iv"]["discovery"]),
                             number(evidence["iv"]["validation"]),
                             number(evidence["ks"]["discovery"]), number(evidence["ks"]["validation"]),
                             number(evidence["psi"]["validation"]), number(evidence["psi"]["observation"])])
            blocks["case2"] = "bad30；KS 为 0–100 点，IV/PSI 无量纲。分箱只由 discovery 参考拟合。\n\n" + table(
                ["特征", "发现 IV", "验证 IV", "发现 KS", "验证 KS", "验证 PSI", "观察 PSI"], rows)
            auxiliary = next(row for row in metrics["ks"] if row["feature"] == "aux_score")
            blocks["card2"] = table(["辅助分验证 KS", "参考行数"],
                [[number(auxiliary["validation"]), finding["reference_row_count"]]])
        elif case == 3:
            audit = [row for row in finding["stats_audit"] if row["stage"] in {"Quality", "Corr_Filter"}
                     and not (row["stage"] == "Quality" and row["status"] == "Selected")]
            blocks["case3"] = table(["特征", "状态", "阶段", "实际原因"],
                [[row["feature"], row["status"], row["stage"], row["reason"]] for row in audit])
            blocks["case3"] += "\n\n" + table(["表示", "方法", "实际相关样本人数", "阈值比较"],
                [[finding[name]["representation"], finding[name]["method"],
                  finding[name]["correlation_row_count"],
                  finding[name]["operator"] + " " + str(finding[name]["threshold"])]
                 for name in ["raw_parameters", "woe_parameters"]])
            blocks["case3"] += "\n\n" + table(["raw 特征对", "有符号相关", "绝对相关"],
                [[row["feature_a"] + " / " + row["feature_b"], number(row["correlation"]),
                  number(row["abs_correlation"])] for row in finding["raw_pairs"]
                 if row["abs_correlation"] >= 0.95])
            blocks["card3"] = table(["Stats 保留", "Linear 保留"],
                [[", ".join(finding["stats_selected"]), ", ".join(finding["linear_selected"])]])
        elif case == 4:
            rows = finding["normal_cells"]
            blocks["case4"] = "discovery / 202601 / bad30；X=b1，风险等级 X3。坏率按已表现权重计算。\n\n" + table(
                ["辅助箱 / 风险等级", "全样本人数", "已表现人数", "已表现权重和", "坏率", "Δ 行基线 pp", "整体 Lift"],
                [[row["y_bin"] + " / Y" + str(row["y_risk_rank"]), row["sample_count"],
                  row["observed_sample_count"], number(row["observed_weight_sum"]),
                  number(row["bad_rate"], 100, "%"), number(row["delta_vs_row"], 100),
                  number(row["lift_vs_overall"])] for row in rows])
            comparable = [row for row in rows if isinstance(row["bad_rate"], (int, float))
                          and math.isfinite(row["bad_rate"])]
            lower = min(comparable, key=lambda row: row["bad_rate"]) if comparable else None
            upper = max(comparable, key=lambda row: row["bad_rate"]) if comparable else None
            blocks["card4"] = table(["同一 X3 层最低加权坏率", "最高加权坏率"],
                [[number(lower["bad_rate"], 100, "%") if lower else "不可比较：无有效分母",
                  number(upper["bad_rate"], 100, "%") if upper else "不可比较：无有效分母"]])
        elif case == 5:
            blocks["case5"] = table(["规则", "审计状态", "拒绝阶段", "实际原因"],
                [[row["rule_id"], row["status"], row["rejection_stage"] or "—", row["reason"] or "—"]
                 for row in finding["candidates"]])
            blocks["case5"] += "\n\n独立 validation / bad30 / hit；sample_count 是已表现人数。\n\n" + table(
                ["规则", "已表现命中人数", "事件人数", "覆盖", "坏率", "Lift", "Lift 下界"],
                [[row["rule_id"], row["sample_count"], row["event_count"], number(row["coverage"], 100, "%"),
                  number(row["event_rate"], 100, "%"), number(row["lift"]), number(row["lift_ci_lower"])]
                 for row in finding["validation"]])
            summary = finding["summary_rows"][0]
            blocks["card5"] = table(["候选数", "保留数", "验证范围"],
                [[summary["candidate_count"], summary["selected_count"], summary["validation_status"]]])
        elif case == 6:
            first, second, empty = payload["queries"]
            blocks["case6"] = table(["实际验证", "结果"],
                [["第一页 / 下一页", str(first["returned_rows"]) + " / " + str(second["returned_rows"])],
                 ["next_offset", first["next_offset"]], ["合法空结果", empty["returned_rows"]],
                 ["无效列请求", finding["invalid_request"]["exception"]],
                 ["最终 Unicode 字符 / 上限", str(finding["context_chars"]) + " / " + str(finding["max_chars"])],
                 ["只读聚合 policy 回放与保存结果一致", finding["policy_replay_matches"]]])
            blocks["card6"] = table(["恢复对象", "预算字符上限"],
                [[finding["snapshot_type"], finding["max_chars"]]])
        else:
            blocks["case7"] = table(["格式", "共享文件", "用途与边界"],
                [[row["format"], "`" + row["path"] + "`", row["use"]] for row in finding["formats"]])
            blocks["card7"] = table(["实际格式数", "Excel 关键值与快照对齐"],
                [[len(finding["formats"]), finding["excel_key_values_match"]]])
    content = "\n\n".join(f"--8<-- [start:{name}]\n{body}\n--8<-- [end:{name}]"
                           for name, body in blocks.items()) + "\n"
    (output / "previews.txt").write_text(content, encoding="utf-8")


def _finalize(output: Path, rows: int, seed: int) -> None:
    """汇总实际已生成结果，收集文件哈希；晚到的真实浏览器截图也可再次收集。"""
    repository = Path(__file__).resolve().parents[2]
    source = subprocess.run(["git", "rev-parse", "HEAD"], cwd=repository, check=True,
                            capture_output=True, text=True).stdout.strip()
    snippets = ["docs/snippets/task_cases.py", "docs/snippets/external_agent_rule_case.py"]
    modules = ["src/mars/analysis/score_cross.py", "src/mars/analysis/score_cross_view.py",
               "src/mars/analysis/profiler.py", "src/mars/analysis/evaluator.py",
               "src/mars/analysis/_risk_profile.py", "src/mars/feature/selection/stats.py",
               "src/mars/feature/selection/linear.py", "src/mars/reporting/_query.py",
               "src/mars/reporting/_artifact.py", "src/mars/rule/workflow.py"]
    cases: dict[str, Any] = {}
    for number in range(1, 8):
        path = output / f"case-{number}.json"
        if path.exists():
            payload = json.loads(path.read_text(encoding="utf-8"))
            cases[str(number)] = {"report_id": payload["report_id"], **payload["findings"]}
    _previews(output)
    _write_json(output / "summary.json", {"seed": seed, "rows": rows,
                "splits": {name: rows // 3 for name in ["discovery", "validation", "observation"]},
                "cases": cases})
    environment_path = output / "generation-environment.json"
    if not environment_path.exists():
        _write_json(environment_path, _environment())
    environment = json.loads(environment_path.read_text(encoding="utf-8"))
    material = (
        "# 共享案例复现材料\n\n"
        f"本包结果：{rows:,} 行，seed={seed}，九个月、三个等量独立分区。\n"
        "公开站点展示规模为 18,000 行；其他规模仅用于 API/语义回归。\n"
        "运行基于本任务源码；固定基线 commit 746b8fa 提供核心 API，新生成入口随任务分支提供。\n"
        "从仓库根目录执行：\n\n"
        "python docs/snippets/task_cases.py --case all --rows 18000 --seed 20261001 "
        "--output-dir docs/assets/cases\n\n"
        "轻量 API/语义回归用 --rows 900，不能冒充公开数字。\n"
        "--case 1..7 按需运行；6/7 需要已有 score-cross 和 policy 快照，不重复计算宽表。\n"
        "--phase consume 在新进程只加载快照；--phase finalize 更新摘要、来源和 ZIP。\n"
        "此 ZIP 是预生成结果与复现材料，不含原始宽表。源码文件见同目录 task_cases.py、"
        "external_agent_rule_case.py；依赖及哈希见 manifest.json。\n"
        "复制脚本到已安装本任务源码环境的同一目录运行，或直接在 checkout 运行。\n"
        "信任边界：仅加载可信 .marsreport 文件；当前格式和专用回放能力以 describe 为准。\n"
    )
    (output / "REPRODUCE.txt").write_text(material, encoding="utf-8")
    for filename in ["task_cases.py", "external_agent_rule_case.py", "external_agent_rule_prompt.md"]:
        source_path = Path(__file__).with_name(filename)
        (output / filename).write_bytes(source_path.read_bytes())
    # 静态 HTML 只规范行尾空格，保留原编码、换行和有效内容；再计算下载字节哈希。
    for html_path in output.glob("*.html"):
        html_bytes = html_path.read_bytes()
        normalized_bytes = re.sub(rb"[ \t]+(?=\r?$)", b"", html_bytes, flags=re.MULTILINE)
        if normalized_bytes != html_bytes:
            html_path.write_bytes(normalized_bytes)
    artifacts = [{"path": path.name, "size_bytes": path.stat().st_size,
                  "sha256": hashlib.sha256(path.read_bytes()).hexdigest()}
                 for path in sorted(output.iterdir())
                 if path.is_file() and path.name not in {"manifest.json", "cases.zip"}]
    manifest: dict[str, Any] = {
        "schema_version": 1, "source_commit": source,
        "command": f"python docs/snippets/task_cases.py --case all --rows {rows} --seed {seed} --output-dir docs/assets/cases",
        "data_config": {"rows": rows, "seed": seed, "reference": "discovery", "n_score_bins": 4,
                        "scope_dimensions": ["dataset", "application_date month", "target"]},
        "dependencies": environment["dependencies"], "python": environment["python"],
        "generation_environment": environment,
        "code_fingerprints": {name: _fingerprint(repository / name) for name in [*snippets, *modules]},
        "reports": {number: payload["report_id"] for number, payload in cases.items()},
        "artifacts": artifacts,
    }
    _write_json(output / "manifest.json", manifest)
    with zipfile.ZipFile(output / "cases.zip", "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for path in sorted(output.iterdir()):
            if path.is_file() and path.name != "cases.zip":
                archive.write(path, arcname=path.name)


def main() -> None:
    """按任务运行必要计算；6/7 仅消费现有产物；所有失败保留非零退出。"""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--case", choices=["all", *map(str, range(1, 8))], default="all")
    parser.add_argument("--rows", type=int, default=DEFAULT_ROWS)
    parser.add_argument("--seed", type=int, default=DEFAULT_SEED)
    parser.add_argument("--output-dir", type=Path, default=Path("docs/assets/cases"))
    parser.add_argument("--phase", choices=["generate", "compute", "consume", "finalize"], default="generate")
    args = parser.parse_args()
    output: Path = args.output_dir
    output.mkdir(parents=True, exist_ok=True)
    if args.phase == "finalize":
        _finalize(output, args.rows, args.seed)
        return
    if args.phase == "consume":
        required = ["score-cross.marsreport", "score-cross.xlsx", "policy.marsreport", "case-4.json"]
        missing = [name for name in required if not (output / name).is_file()]
        if missing:
            parser.error(f"consume 缺少前置文件 {missing}；先在同一输出目录运行 --case 4。")
        _consume(output)
        _delivery(output)
        return
    cases = set(range(1, 8)) if args.case == "all" else {int(args.case)}
    if args.phase == "generate":
        # 计算进程先完整退出，之后新进程只能读取持久化报告；没有原宽表可复用。
        if cases.intersection({1, 2, 3, 4, 5}):
            subprocess.run([sys.executable, str(Path(__file__).resolve()), "--phase", "compute",
                            "--case", args.case, "--rows", str(args.rows), "--seed", str(args.seed),
                            "--output-dir", str(output.resolve())], check=True)
        if cases.intersection({4, 5, 6, 7}):
            subprocess.run([sys.executable, str(Path(__file__).resolve()), "--phase", "consume",
                            "--output-dir", str(output.resolve())], check=True)
        _finalize(output, args.rows, args.seed)
        print(f"已生成案例 {args.case}：rows={args.rows} seed={args.seed}；合成统计，无 LLM。")
        return
    if cases.intersection({1, 2, 3, 4, 5}):
        _write_json(output / "generation-environment.json", _environment())
        data = _data(args.rows, args.seed)
        context = _context(data, args.seed)
        cross: Report | None = None
        if cases.intersection({4, 5}):
            cross = _score_cross(data, output, context)
        if 1 in cases:
            _profile(data, output, context)
        if 2 in cases:
            _binning(data, output, context)
        if 3 in cases:
            _selection(data, output, context)
        if 5 in cases:
            assert cross is not None
            _rules(data, cross, output)


if __name__ == "__main__":
    main()
