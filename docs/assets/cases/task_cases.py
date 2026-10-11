"""Home Credit 2018 七个共享任务案例；全量计算、固定 Train 基准和无 CSV 消费。"""

from __future__ import annotations

import argparse
import ast
import hashlib
import importlib.metadata
import json
import math
import re
import shutil
import subprocess
import sys
import time
import zipfile
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
import polars as pl
from external_agent_rule_case import DIRECTIONS, consume, produce
from home_credit_data import (
    DEFAULT_DATA_DIR,
    DEFAULT_DISPLAY_FEATURES,
    HomeCreditConfig,
    HomeCreditInputs,
    prepare_home_credit_data,
)

import mars
from mars.analysis import (
    MarsBinEvaluator,
    cross_scores,
    evaluate_score_policy,
    get_score_bin_definitions,
    profile_risk,
    profile_stats,
    write_score_cross_html,
)
from mars.feature import MarsLinearSelector, MarsStatsSelector
from mars.reporting import Report, load_report, show_correlation_matrix
from mars.reporting._serialization import json_safe
from mars.rule import MarsCombinationRuleGenerator, MarsRuleMiningSpec, MarsRuleReport, mine_rules

DEFAULT_SEED = 20261001
FEATURES = ["EXT_SOURCE_2", "EXT_SOURCE_3", "AMT_INCOME_TOTAL", "DAYS_BIRTH"]
BASELINE: dict[str, Any] = {"type": "x_only", "x_max_risk_rank": 2, "missing_score": "reject"}
POLICY: dict[str, Any] = {"type": "expression", "expression": "X <= X2 AND Y <= Y2", "missing_score": "reject"}
REPORT_STEMS = ["profile", "binning", "binning-time", "days-employed-special", "score-cross", "policy",
                "woe-correlation", "raw-correlation", "rules", "rules-independent", "auto-rules"]
SCRIPT_FILES = ["task_cases.py", "external_agent_rule_case.py", "home_credit_data.py"]
TEXT_ARTIFACTS = ["preparation-summary.json", "feature-metadata.json", "generation-environment.json",
                  "generation-timings.json", "selection.json", "selection.xlsx", "correlation.html",
                  "correlation.xlsx", "rule-generation.json", "case-notes.json", "ruleset.json",
                  "external-agent-task.md", "external-agent-rule-task.txt", "external-agent-task.txt",
                  "query-trace.json", "evidence-review.json", "review.md", "rule-agent-context.json",
                  "agent-context.json", "summary.json", "previews.txt", "REPRODUCE.txt",
                  "binning-native-evidence.json", "binning-native-home-credit.png",
                  "binning-native-home-credit.svg", "binning-native-time-home-credit.png",
                  "binning-native-time-home-credit.svg"]


def _write_json(path: Path, value: Any) -> None:
    """使用报告编码合同写有限聚合统计，不保存原始样本。"""
    path.write_text(json.dumps(json_safe(value), ensure_ascii=False, indent=2,
                               allow_nan=False) + "\n", encoding="utf-8")


def _write_correlation_html(matrix: pd.io.formats.style.Styler, path: Path) -> None:
    """保留原生矩阵色带与数值，将宽表放在自身可横向滚动的完整文档中。"""
    document = (
        '<!doctype html>\n<html lang="zh-CN"><head><meta charset="utf-8">\n'
        '<meta name="viewport" content="width=device-width, initial-scale=1">\n'
        '<title>Home Credit 原始数值相关矩阵</title>\n'
        '<style>body{font-family:Arial,sans-serif;margin:24px;color:#202124}'
        'h1{font-size:24px;overflow-wrap:anywhere}'
        '.mars-correlation-scroll{max-width:100%;width:100%;overflow-x:auto}'
        '@media(max-width:600px){body{margin:16px}}</style></head>\n'
        '<body><h1>Home Credit 原始数值相关矩阵</h1>\n'
        '<div class="mars-correlation-scroll" tabindex="0" aria-label="原始数值相关矩阵，可横向滚动">\n'
        + matrix.to_html()
        + '\n</div></body></html>\n'
    )
    path.write_text(document, encoding="utf-8")


def _refresh_rule_html_styles(output: Path) -> None:
    """同步当前原生渲染器的视口与滚动样式，保留既有表值及报告身份。"""
    current = MarsRuleReport().render_html()
    style = re.search(r"<style>.*?</style>", current, re.S)
    viewport = re.search(r'<meta name="viewport"[^>]*>', current)
    # 兼容已发布 0.0.28 的只读消费；旧渲染器不覆盖下载包已经修正的样式。
    if style is None or viewport is None:
        return
    for stem in ("rules", "rules-independent", "auto-rules"):
        path = output / f"{stem}.html"
        if not path.is_file():
            continue
        content = path.read_text(encoding="utf-8")
        content = re.sub(r"<style>.*?</style>", lambda match: style.group(), content, count=1, flags=re.S)
        if re.search(r'<meta name="viewport"[^>]*>', content):
            content = re.sub(r'<meta name="viewport"[^>]*>', lambda match: viewport.group(), content, count=1)
        else:
            content = content.replace("</head>", viewport.group() + "\n</head>", 1)
        path.write_text(content, encoding="utf-8")


def _rows(frame: Any) -> list[dict[str, Any]]:
    """只物化已分页或明确受限的证据。"""
    return json_safe(frame.to_dicts() if isinstance(frame, pl.DataFrame) else frame.to_dict("records"))


def _query(report: Report, table: str, question: str, **query: Any) -> dict[str, Any]:
    """保存实际分页结果及可重放引用。"""
    page = report.query_page(table, **query)
    return {"question": question, "reference": page["reference"], "total_rows": page["total_rows"],
            "returned_rows": page["returned_rows"], "omitted_rows": page["omitted_rows"],
            "next_offset": page["next_offset"], "rows": _rows(page["data"])}


def _case(
    output: Path, number: int, report: Report, question: str,
    queries: list[dict[str, Any]], findings: dict[str, Any], unavailable: str,
) -> None:
    """人工摘要和 Agent 查询绑定同一报告身份。"""
    payload = {"case": number, "report_id": report.report_id, "question": question,
               "consumer": "deterministic public API queries; no LLM", "description": report.describe(),
               "queries": queries, "findings": findings, "unavailable": unavailable}
    _write_json(output / f"case-{number}.json", payload)
    _write_json(output / f"query-{number}.json", {
        "case": number, "report_id": report.report_id, "consumer": payload["consumer"],
        "question": question, "query": queries[0], "unavailable": unavailable,
    })


def _export(report: Any, output: Path, stem: str) -> None:
    """快照、HTML 与 Excel 由同一实际报告导出。"""
    report.save(output / f"{stem}.marsreport", overwrite=True)
    report.write_html(str(output / f"{stem}.html"))
    report.write_excel(str(output / f"{stem}.xlsx"))


# --8<-- [start:profile]
def _profile(inputs: HomeCreditInputs, output: Path) -> None:
    """全部候选字段画像，全部五组进入质量与 Train 参考 PSI。"""
    report = profile_stats(
        inputs.df, features=inputs.fcols, metrics=["missing", "zeros", "mean", "psi"],
        benchmark_df=inputs.benchmark_df, group_col="dataset_flag",
        categorical_features=[name for name in inputs.fcols if inputs.df.schema[name] == pl.String],
        feature_metadata=inputs.feature_metadata, business_context=inputs.business_context,
    )
    _export(report, output, "profile")
    queries = [
        _query(report, "overview", "真实外部来源分覆盖率及字段单位", features=list(DIRECTIONS), limit=3),
        _query(report, "dq.missing", "五分区外部来源分缺失率", features=list(DIRECTIONS), limit=3),
        _query(report, "overview", "收入与原始就业天数的质量", features=["AMT_INCOME_TOTAL", "DAYS_EMPLOYED"], limit=3),
    ]
    _case(output, 1, report, "全量数据与无标签组有哪些质量差异？", queries, {
        "input_row_count": inputs.df.height, "candidate_features": inputs.fcols,
        "analyzed_features": inputs.fcols, "labels_by_split": inputs.summary["partition_stats"],
        "special_values": inputs.summary.get("special_values", {}),
        "reference_row_count": inputs.benchmark_df.height,
        "feature_example": report.describe()["feature_metadata"]["EXT_SOURCE_2"],
    }, "July 是原始比赛测试来源；模拟日期的差异不能解释成真实时间漂移。")
# --8<-- [end:profile]


# --8<-- [start:binning]
def _binning(inputs: HomeCreditInputs, output: Path) -> None:
    """Train 决定所有固定切点与参考分布，五分区和模拟月趋势复用。"""
    display = [name for name in DEFAULT_DISPLAY_FEATURES if name in inputs.fcols]
    run = profile_risk(
        inputs.df, target="TARGET", features=inputs.fcols, benchmark_df=inputs.benchmark_df,
        method="quantile", n_bins=4, group_col="dataset_flag", time_col="APPLY_DATE",
        time_grain="month", amount_col="AMT_CREDIT",
        feature_metadata=inputs.feature_metadata, business_context=inputs.business_context,
        n_jobs=1, batch_size=30,
    )
    report = run.report
    report.save(output / "binning.marsreport", overwrite=True)
    report.write_html(str(output / "binning.html"), max_plots=len(display), include_charts=True,
                      chart_embed_mode="inline", report_name="Home Credit 固定分箱与五分区风险")
    load_report(output / "binning.marsreport").write_excel(str(output / "binning.xlsx"))
    time_report = MarsBinEvaluator().evaluate(
        inputs.df, target="TARGET", features=display, binner=run.binner,
        benchmark_df=inputs.benchmark_df, time_col="APPLY_DATE", time_grain="month",
        amount_col="AMT_CREDIT", feature_metadata=inputs.feature_metadata,
        business_context=inputs.business_context,
    ).report
    time_report.save(output / "binning-time.marsreport", overwrite=True)
    time_report.write_html(str(output / "binning-time.html"), max_plots=4, chart_embed_mode="inline")
    load_report(output / "binning-time.marsreport").write_excel(str(output / "binning-time.xlsx"))
    # 特殊码只应用于 DAYS_EMPLOYED 单字段，完整候选报告同时保留原始未改写值。
    employment = profile_risk(
        inputs.df, target="TARGET", features=["DAYS_EMPLOYED"], benchmark_df=inputs.benchmark_df,
        method="quantile", n_bins=4, special_values=[365243], group_col="dataset_flag",
        time_col="APPLY_DATE", time_grain="month",
        amount_col="AMT_CREDIT", feature_metadata=inputs.feature_metadata,
        business_context={**inputs.business_context, "special_value_processing": {"DAYS_EMPLOYED": [365243]}}, n_jobs=1,
    ).report
    employment.save(output / "days-employed-special.marsreport", overwrite=True)
    employment.write_html(str(output / "days-employed-special.html"), max_plots=1)
    load_report(output / "days-employed-special.marsreport").write_excel(str(output / "days-employed-special.xlsx"))
    actual_parameters = {"features": "EXT_SOURCE_2", "target": "TARGET", "dpi": 300, "show_risk": "both"}
    for image_format in ("png", "svg"):
        for native, stem in ((report, "binning-native-home-credit"), (time_report, "binning-native-time-home-credit")):
            generated = native.save_risk_trend_images(output, **actual_parameters,
                image_format=image_format, filename_prefix="native-export")
            shutil.copyfile(generated[0], output / f"{stem}.{image_format}")
            generated[0].unlink()
    queries = [
        _query(report, "summary", "外部来源分 2 与 3 的已观测区分度", features=list(DIRECTIONS), limit=2),
        _query(report, "detail", "主图五组固定箱人数和已观测坏率", features="EXT_SOURCE_2", limit=30),
        _query(report, "calculation_status", "无标签组监督指标真实状态", filters={"group": "4_oot_undue"}, features=list(DIRECTIONS), limit=10),
        _query(report, "detail", "收入与年龄天数原始分析入口", features=["AMT_INCOME_TOTAL", "DAYS_BIRTH"], limit=8),
    ]
    metrics: dict[str, Any] = {}
    for metric in ("iv", "ks", "auc", "psi"):
        evidence = _query(report, f"trend.{metric}", f"五分区 {metric} 固定分箱比较", features=FEATURES, limit=4)
        queries.append(evidence)
        metrics[metric] = evidence["rows"]
    _write_json(output / "binning-native-evidence.json", {
        "report_id": report.report_id, "time_report_id": time_report.report_id,
        "feature": "EXT_SOURCE_2", "target": "TARGET", "chart_png": "binning-native-home-credit.png",
        "chart_svg": "binning-native-home-credit.svg", "source_api": "MarsBinningReport.save_risk_trend_images",
        "parameters": actual_parameters, "layout": "existing native figure; no crop/redraw",
        "group": "dataset_flag", "time_col": "APPLY_DATE", "time_grain": "month",
        "scope": "five partitions; separate simulated monthly view", "queries": queries[:3],
    })
    _case(output, 2, report, "固定分箱怎样比较五分区与无标签分布？", queries, {
        "summary_rows": queries[0]["rows"], "status_rows": queries[2]["rows"], "split_metrics": metrics,
        "reference_row_count": inputs.benchmark_df.height, "targets": ["TARGET"],
        "candidate_features": inputs.fcols, "analyzed_features": inputs.fcols, "display_features": display,
        "primary_feature": "EXT_SOURCE_2", "time_report_id": time_report.report_id,
        "units": {"ks": "0-100 percentage points", "bad_rate": "fraction", "iv": "dimensionless", "psi": "dimensionless"},
    }, "模拟日期不能证明真实跨期泛化；无标签组只有分布/PSI，没有可观测坏率或 KS。")
# --8<-- [end:binning]


# --8<-- [start:selection]
def _selection(inputs: HomeCreditInputs, output: Path) -> None:
    """Stats 的 WOE Pearson 与 Linear 原始值 Spearman 均只拟合 Train。"""
    train = inputs.benchmark_df
    stats = MarsStatsSelector(
        missing_thr=0.9, rough_iv_thr=0.01, rough_lift_thr=1.2, iv_thr=0.01,
        lift_thr=1.2, psi_thr=0.25, rc_thr=0.5, corr_thr=0.8,
        rough_binning_params={"method": "quantile", "n_bins": 4},
        skip_fine_scan=True, n_jobs=1,
    ).fit(train, target="TARGET", features=inputs.fcols, benchmark_df=train,
          feature_metadata=inputs.feature_metadata, time_col="APPLY_DATE", time_grain="month",
          business_context=inputs.business_context)
    selection = stats.get_report()
    pd.DataFrame(_rows(selection)).to_excel(output / "selection.xlsx", index=False)
    _write_json(output / "selection.json", _rows(selection))
    woe = stats.get_correlation_report()
    woe.save(output / "woe-correlation.marsreport", overwrite=True)
    # Linear 对请求字段池整行删除缺失；限定真实数值子集，明确未计算全宽相关矩阵。
    numeric = ["EXT_SOURCE_1", "EXT_SOURCE_2", "EXT_SOURCE_3", "AMT_INCOME_TOTAL", "AMT_ANNUITY",
               "AMT_GOODS_PRICE", "DAYS_BIRTH", "REGION_RATING_CLIENT", "REGION_RATING_CLIENT_W_CITY",
               "OBS_30_CNT_SOCIAL_CIRCLE", "OBS_60_CNT_SOCIAL_CIRCLE"]
    numeric = [name for name in numeric if name in inputs.fcols]
    linear = MarsLinearSelector(corr_thr=0.95, corr_method="spearman", n_jobs=1).fit(
        train.select(numeric), train["TARGET"], features=numeric,
        feature_metadata=inputs.feature_metadata, business_context=inputs.business_context,
    )
    raw = linear.get_correlation_report()
    raw.save(output / "raw-correlation.marsreport", overwrite=True)
    matrix = show_correlation_matrix(raw)
    _write_correlation_html(matrix, output / "correlation.html")
    matrix.to_excel(output / "correlation.xlsx")
    queries = [
        _query(raw, "correlation_decisions", "真实数值相关性是否触发淘汰", limit=12),
        _query(raw, "pairs", "实际最强原始字段关系", sort_by="abs_correlation", descending=True, limit=6),
        _query(woe, "correlation_decisions", "目标感知 WOE 空间的实际决策", limit=8),
    ]
    _case(output, 3, raw, "为什么删除一个真实字段，保留另一个？", queries, {
        "candidate_features": inputs.fcols, "stats_selected": stats.selected_features_,
        "stats_audit": _rows(selection), "linear_candidate_features": numeric,
        "linear_selected": linear.selected_features_, "woe_report_id": woe.report_id,
        "raw_pairs": queries[1]["rows"], "raw_parameters": raw.describe()["parameters"],
        "woe_parameters": woe.describe()["parameters"], "selection_scope": "0train only; 1val/2test/3oot not used for selection",
    }, "相关性是该字段池、表示与缺失处理下的关联，不说明业务因果。")
# --8<-- [end:selection]


# --8<-- [start:score_cross]
def _score_cross(inputs: HomeCreditInputs, output: Path) -> Report:
    """Train 固定外部来源分切点，五分区保留真实人数、金额与未观测状态。"""
    context = {**inputs.business_context, "score_directions": DIRECTIONS,
               "score_direction_basis": "official normalized external score; observed Train association only",
               "binning_origin": "0train only; frozen cutpoints"}
    initial = cross_scores(inputs.benchmark_df, score_x="EXT_SOURCE_2", score_y="EXT_SOURCE_3",
                           targets=["TARGET"], score_directions=DIRECTIONS, n_bins=4,
                           feature_metadata=inputs.feature_metadata, business_context=context)
    definitions = get_score_bin_definitions(initial)
    cross = cross_scores(
        inputs.df, score_x="EXT_SOURCE_2", score_y="EXT_SOURCE_3", targets=["TARGET"],
        score_directions=DIRECTIONS, bin_definitions=definitions, group_col="dataset_flag",
        amount_col="AMT_CREDIT", min_observed=30,
        feature_metadata=inputs.feature_metadata, business_context=context,
    )
    cross.save(output / "score-cross.marsreport", overwrite=True)
    policy = evaluate_score_policy(cross, POLICY, baseline=BASELINE)
    policy.save(output / "policy.marsreport", overwrite=True)
    policy.write_excel(str(output / "policy.xlsx"))
    write_score_cross_html(cross, output / "score-cross.html", policy_reports=[policy],
                           report_name="Home Credit：外部来源分 2 × 外部来源分 3")
    cross.write_excel(str(output / "score-cross.xlsx"))
    scope = {"group": "0train", "target": "TARGET"}
    queries = [
        _query(cross, "cells", "同一来源分 2 箱内的来源分 3 风险梯度", filters={**scope, "x_bin": "b1"}, sort_by="y_risk_rank", limit=8),
        _query(cross, "row_summary", "来源分 2 边际风险包含缺失箱", filters=scope, sort_by="x_risk_rank", limit=6),
        _query(cross, "column_summary", "来源分 3 的边际风险", filters=scope, sort_by="y_risk_rank", limit=6),
        _query(cross, "cells", "无标签比赛测试的分布和未观测状态", filters={"group": "4_oot_undue", "target": "TARGET"}, limit=6),
        _query(policy, "changes", "固定表达式与单来源分基线的聚合回放", limit=5),
    ]
    _case(output, 4, cross, "两个外部来源分联合后，哪些区间有风险差异？", queries, {
        "scope": scope, "x_bin": "b1", "normal_cells": [row for row in queries[0]["rows"] if row["y_bin"].startswith("b")],
        "status_counts": cross.get_table("cells").group_by("status").len().sort("status").to_dicts(),
        "input_row_count": inputs.df.height, "bad_rate_denominator": "observed_sample_count",
        "ratio_unit": "fraction; delta pp = 100 * delta_vs_row", "bin_definitions": definitions,
        "policy_report_id": policy.report_id,
    }, "外部来源分不是本例训练模型或校准违约概率；关联与历史留存回放不证明因果或利润。")
    return cross
# --8<-- [end:score_cross]


# --8<-- [start:automatic_rules]
def _automatic_rules(inputs: HomeCreditInputs, output: Path) -> dict[str, Any]:
    """组合生成器只用 Train/Val，记录其实际必需有限数值输入人数。"""
    eligible = inputs.df.filter(pl.all_horizontal(pl.col(name).is_finite() for name in DIRECTIONS))
    train = eligible.filter(pl.col("dataset_flag") == "0train")
    validation = eligible.filter(pl.col("dataset_flag") == "1val")
    result = mine_rules(
        train, target="TARGET", validation_df=validation, features=list(DIRECTIONS),
        generators=[MarsCombinationRuleGenerator(n_bins=4, max_candidates=24, random_state=42)],
        amount_col="AMT_CREDIT", spec=MarsRuleMiningSpec.production(max_candidates=24, top_k=5),
    )
    report = result.to_report(feature_metadata=inputs.feature_metadata, business_context={
        **inputs.business_context, "candidate_origin": "MarsCombinationRuleGenerator; no hand-written seeds",
        "score_domain": "finite EXT_SOURCE_2/3; generator input prerequisite",
        "selection_roles": "0train discovery, 1val selection; no 2test/3oot access; temporal gate disabled",
    })
    _export(report, output, "auto-rules")
    evidence = {"report_id": report.report_id, "description": report.describe(), "candidate_budget": 24,
                "raw_input_rows": inputs.df.height, "eligible_rows": eligible.height,
                "excluded_rows": inputs.df.height - eligible.height,
                "generation_rows": train.height, "selection_rows": validation.height,
                "queries": [
                    _query(report, "summary", "实际生成及保留多少自动候选", limit=1),
                    _query(report, "candidates", "生成来源与拒绝原因", limit=24),
                    _query(report, "evaluation", "Val 参与筛选的证据", filters={"dataset": "validation", "target": "TARGET", "group": "hit"}, limit=8),
                ]}
    _write_json(output / "rule-generation.json", evidence)
    return evidence
# --8<-- [end:automatic_rules]


# --8<-- [start:rules]
def _rules(inputs: HomeCreditInputs, cross: Report, output: Path) -> None:
    """种子格子和自动生成器各有真实审计；独立评估不重新选择。"""
    produce(output, inputs=inputs, cross_report=cross)
    report = load_report(output / "rules.marsreport")
    fixed = load_report(output / "rules-independent.marsreport")
    queries = [
        _query(report, "summary", "种子候选最终保留多少", limit=1),
        _query(report, "candidates", "实际入选/拒绝阶段及原因", limit=8),
        _query(report, "evaluation", "Val 参与选择的风险与覆盖", filters={"dataset": "validation", "target": "TARGET", "group": "hit"}, limit=8),
        _query(fixed, "slices", "固定规则在独立 test/oot 的表现", filters={"slice": {"op": "in", "value": ["2test", "3oot"]}, "group": "hit"}, limit=8),
        _query(fixed, "slices", "无标签组的全部人数与未观测状态", filters={"slice": "4_oot_undue", "group": "hit"}, limit=8),
    ]
    automatic = _automatic_rules(inputs, output)
    _case(output, 5, report, "从 Train 发现、Val 选择，走到固定规则的独立复核", queries, {
        "summary_rows": queries[0]["rows"], "candidates": queries[1]["rows"],
        "validation": queries[2]["rows"], "independent": queries[3]["rows"], "unobserved": queries[4]["rows"],
        "cross_report_id": cross.report_id, "independent_report_id": fixed.report_id,
        "candidate_origins": report.describe()["business_context"]["candidate_origins"],
        "raw_data_saved": False, "rule_status": "Experimental", "automatic_generation": automatic,
        "auto_report_id": automatic["report_id"],
    }, "Val 被用于选择；3oot 仅为模拟时间外样本。报告不授予生产策略部署许可。")
# --8<-- [end:rules]


# --8<-- [start:restore]
def _consume(output: Path) -> None:
    """新进程只读快照，验证分页、字符预算、引用和固定聚合策略回放。"""
    report = load_report(output / "binning.marsreport")
    original = json.loads((output / "case-2.json").read_text(encoding="utf-8"))
    assert report.report_id == original["report_id"]
    for key in ("feature_metadata", "business_context", "parameters", "tables"):
        assert report.describe()[key] == original["description"][key], key
    for query in original["queries"]:
        ref = query["reference"]
        if ref["report_id"] == report.report_id:
            assert _rows(report.get_table(ref["table"], **ref["query"])) == query["rows"]
    scope = {"mars_group": "1val", "y": "TARGET"}
    columns = ["feature", "mars_group", "bin_index", "bin_label", "count", "observed_count", "bad_rate"]
    first = _query(report, "detail", "验证组固定箱实际第一页", features="EXT_SOURCE_2", filters=scope,
                   columns=columns, sort_by="bin_index", limit=3)
    second = _query(report, "detail", "按 next_offset 下一页", features="EXT_SOURCE_2", filters=scope,
                    columns=columns, sort_by="bin_index", offset=first["next_offset"], limit=3)
    empty = _query(report, "detail", "合法空结果", filters={"mars_group": "not-present"}, limit=3)
    unobserved = _query(report, "detail", "无标签组的分布人数与空坏率", features="EXT_SOURCE_2",
                        filters={"mars_group": "4_oot_undue"}, columns=columns, limit=6)
    try:
        report.get_table("detail", columns=["not_a_column"], limit=1)
    except ValueError as exc:
        invalid_request = {"exception": "ValueError", "message": str(exc)}
    else:
        raise AssertionError("无效列必须失败。")
    context = report.to_ai_context(queries={"detail": {"features": "EXT_SOURCE_2", "filters": scope,
        "columns": columns, "limit": 40}}, max_chars=9000)
    bounded = json.loads(context)
    assert len(context) <= 9000
    for entry in bounded["evidence"]:
        assert _rows(report.get_table(entry["reference"], **entry["query"])) == entry["rows"]
    (output / "agent-context.json").write_text(context + "\n", encoding="utf-8")
    replay = evaluate_score_policy(load_report(output / "score-cross.marsreport"), POLICY, baseline=BASELINE)
    policy = load_report(output / "policy.marsreport")
    assert _rows(replay.get_table("changes")) == _rows(policy.get_table("changes"))
    _case(output, 6, report, "保存分箱报告之后，如何在新进程继续查询？", [first, second, empty, unobserved], {
        "snapshot_type": type(report).__name__, "identity_preserved": True, "separate_process": True,
        "context_chars": len(context), "max_chars": 9000, "budget_omissions": bounded.get("omissions", bounded.get("omitted", [])),
        "invalid_request": invalid_request, "policy_replay_matches": True, "table_names": list(report.describe()["tables"]),
    }, "快照保存聚合结果与能力边界，不恢复原始宽表或任意分析器。")
    if (output / "rules-independent.marsreport").exists():
        consume(output)
    if (output / "raw-correlation.marsreport").is_file():
        matrix = show_correlation_matrix(load_report(output / "raw-correlation.marsreport"))
        _write_correlation_html(matrix, output / "correlation.html")
    _refresh_rule_html_styles(output)
# --8<-- [end:restore]


# --8<-- [start:delivery]
def _delivery(output: Path) -> None:
    """同批 HTML、Excel 和可查询快照分别交付，核对 Excel 实际数值。"""
    import openpyxl

    report = load_report(output / "binning.marsreport")
    workbook = openpyxl.load_workbook(output / "binning.xlsx", read_only=True, data_only=True)
    rows = workbook["001_detail"].iter_rows(values_only=True)
    header = list(next(rows))
    excel_sample = next(rows)
    table_sample = report.get_table("detail", limit=1).row(0, named=True)
    for name in ("count", "observed_count", "bad", "bad_rate"):
        expected, actual = table_sample[name], excel_sample[header.index(name)]
        assert actual == expected or (actual is not None and expected is not None and np.isclose(actual, expected)), name
    sheet_names = workbook.sheetnames
    workbook.close()
    prompt = (
        "用 load_report('binning.marsreport') 读取可信快照，describe 查看目录、单位与来源。\n"
        "query_page('detail', features='EXT_SOURCE_2', filters={'mars_group':'1val'}, limit=3)，按 next_offset 继续。\n"
        "引用 reference，区分 count 与 observed_count；坏率是 observed 分母下的比例。\n"
        "4_oot_undue 标签未公开；APPLY_DATE 是模拟日期；外部来源分不是校准概率。\n"
        "to_ai_context 预算是 Unicode 字符。无需内置 mars.agent、LLM 或原始 CSV。\n"
    )
    (output / "external-agent-task.txt").write_text(prompt, encoding="utf-8")
    formats = [
        {"format": "HTML", "path": "binning.html", "use": "原生分箱图、指标与业务元数据交互"},
        {"format": "Excel", "path": "binning.xlsx", "use": "同批分箱静态工作表"},
        {"format": ".marsreport", "path": "binning.marsreport", "use": "目录、查询、保存统计的导出；附加交叉聚合回放另用 score-cross"},
        {"format": "bounded JSON", "path": "agent-context.json", "use": "Unicode 字符预算内的目录与证据"},
        {"format": "TXT", "path": "external-agent-task.txt", "use": "外部消费者提示材料；未执行 LLM"},
    ]
    _case(output, 7, report, "同一次分箱分析怎样交付给不同使用者？", [_query(report, "detail", "Excel 首行实际值与快照对齐", limit=1)], {
        "formats": formats, "excel_sheets": sheet_names, "excel_key_values_match": True, "shared_cases": [2, 6, 7],
    }, "Excel 不恢复交互计算；快照不恢复任意 Python 对象。")
# --8<-- [end:delivery]


def _fingerprint(path: Path) -> str:
    """AST 排除说明与位置信息，样式改变不触发重算。"""
    tree = ast.parse(path.read_text(encoding="utf-8-sig"))
    for node in ast.walk(tree):
        if isinstance(node, (ast.Module, ast.FunctionDef, ast.AsyncFunctionDef, ast.ClassDef)) and node.body:
            first = node.body[0]
            if isinstance(first, ast.Expr) and isinstance(first.value, ast.Constant) and isinstance(first.value.value, str):
                node.body.pop(0)
    return hashlib.sha256(ast.dump(tree, include_attributes=False).encode("utf-8")).hexdigest()


def _environment() -> dict[str, Any]:
    """记录真实依赖，不公开本机绝对路径。"""
    return {"python": sys.version.split()[0], "imported_mars_version": mars.__version__,
            "installed_distribution_version": importlib.metadata.version("mars-risk"),
            "import_source": "current MARS package; code fingerprints recorded",
            "dependencies": {name: importlib.metadata.version(name) for name in (
                "numpy", "pandas", "polars", "scikit-learn", "pyarrow", "openpyxl", "xlsxwriter", "matplotlib")}}


def _table(headers: list[str], rows: list[list[Any]]) -> str:
    """从实际有限证据生成 Markdown 紧凑预览。"""
    lines = ["| " + " | ".join(headers) + " |", "| " + " | ".join(["---"] * len(headers)) + " |"]
    for row in rows:
        cells = ["—" if value is None else str(value).replace("|", "\\|").replace("\n", " ") for value in row]
        lines.append("| " + " | ".join(cells) + " |")
    return "\n".join(lines)


def _number(value: float | None, scale: float = 1, suffix: str = "") -> str:
    """按真实指标单位格式化，不补未观测值。"""
    return "—" if value is None or not math.isfinite(value) else f"{value * scale:.4f}{suffix}"


def _previews(output: Path) -> None:
    """正文和首页卡片只使用同批 JSON 中的实际结果。"""
    blocks: dict[str, str] = {}
    for number in range(1, 8):
        path = output / f"case-{number}.json"
        if not path.exists():
            continue
        case = json.loads(path.read_text(encoding="utf-8"))
        finding = case["findings"]
        if number == 1:
            partitions = finding["labels_by_split"]
            blocks["case1"] = _table(["分区", "全部人数", "已观测", "坏率"], [[r["dataset_flag"], r["total_count"], r["observed_count"], _number(r["observed_bad_rate"], 100, "%")] for r in partitions])
            blocks["card1"] = _table(["合并样本", "无标签样本", "完整候选字段"], [[finding["input_row_count"], partitions[-1]["total_count"], len(finding["analyzed_features"])]])
        elif number == 2:
            metrics = finding["split_metrics"]
            values = []
            for feature in FEATURES:
                evidence = {metric: next(r for r in rows if r["feature"] == feature) for metric, rows in metrics.items()}
                values.append([feature, _number(evidence["ks"]["0train"]), _number(evidence["ks"]["1val"]), _number(evidence["ks"]["2test"]), _number(evidence["ks"]["3oot"]), _number(evidence["psi"]["4_oot_undue"])])
            blocks["case2"] = "TARGET；KS 为 0–100 点，PSI 无量纲。0train 拟合，3oot 的日期为模拟日期。\n\n" + _table(["字段", "Train KS", "Val KS", "Test KS", "OOT KS", "无标签 PSI"], values)
            blocks["card2"] = _table(["主图字段", "Train 参考人数"], [[finding["primary_feature"], finding["reference_row_count"]]])
        elif number == 3:
            audit = [r for r in finding["stats_audit"] if r["status"] == "Dropped"][:10]
            blocks["case3"] = _table(["字段", "状态", "阶段", "原因"], [[r["feature"], r["status"], r["stage"], r["reason"]] for r in audit])
            blocks["case3"] += "\n\n" + _table(["原始字段对", "Spearman", "绝对相关"], [[r["feature_a"] + " / " + r["feature_b"], _number(r["correlation"]), _number(r["abs_correlation"])] for r in finding["raw_pairs"][:4]])
            blocks["card3"] = _table(["Stats 候选→保留", "Linear 候选→保留"], [[f"{len(finding['candidate_features'])} → {len(finding['stats_selected'])}", f"{len(finding['linear_candidate_features'])} → {len(finding['linear_selected'])}"]])
        elif number == 4:
            rows = finding["normal_cells"]
            blocks["case4"] = "0train / TARGET / 来源分 2 固定 b1；坏率按已观测人数计算。\n\n" + _table(["来源分 3 箱", "全部人数", "已观测", "坏率", "Δ 行基线 pp", "Lift"], [[r["y_bin"], r["sample_count"], r["observed_sample_count"], _number(r["bad_rate"], 100, "%"), _number(r["delta_vs_row"], 100), _number(r["lift_vs_overall"])] for r in rows])
            rates = [r["bad_rate"] for r in rows if r["bad_rate"] is not None]
            blocks["card4"] = _table(["同一 b1 层最低坏率", "最高坏率"], [[_number(min(rates), 100, "%") if rates else None, _number(max(rates), 100, "%") if rates else None]])
        elif number == 5:
            origins = finding["candidate_origins"]
            def label(rule_id: str, evidence: dict[str, Any] = origins) -> str:
                """用真实交叉格子解释稳定规则身份。"""
                region = evidence[rule_id]["region"]
                return f"来源分2 {region['x_bin']} × 来源分3 {region['y_bin']}"
            blocks["case5"] = _table(["规则", "状态", "阶段", "原因"], [[label(r["rule_id"]), r["status"], r["rejection_stage"], r["reason"]] for r in finding["candidates"]])
            blocks["case5"] += "\n\n固定规则独立复核；sample_count 为已观测人数。\n\n" + _table(["分区", "规则", "全部命中", "已观测命中", "坏率", "Lift"], [[r["slice"], label(r["rule_id"]), r["total_count"], r["sample_count"], _number(r["event_rate"], 100, "%"), _number(r["lift"])] for r in finding["independent"]])
            automatic = finding["automatic_generation"]["queries"][0]["rows"][0]
            summary = finding["summary_rows"][0]
            blocks["card5"] = _table(["格子候选→保留", "自动候选→保留"], [[f"{summary['candidate_count']} → {summary['selected_count']}", f"{automatic['candidate_count']} → {automatic['selected_count']}"]])
        elif number == 6:
            first, second, empty = case["queries"][:3]
            blocks["case6"] = _table(["检查", "实际结果"], [["第一页/下一页", f"{first['returned_rows']}/{second['returned_rows']}"], ["合法空结果", empty["returned_rows"]], ["无效列", finding["invalid_request"]["exception"]], ["字符/预算", f"{finding['context_chars']}/{finding['max_chars']}"], ["固定策略回放一致", finding["policy_replay_matches"]]])
            blocks["card6"] = _table(["恢复对象", "字符预算"], [[finding["snapshot_type"], finding["max_chars"]]])
        else:
            blocks["case7"] = _table(["格式", "共享文件", "用途"], [[r["format"], r["path"], r["use"]] for r in finding["formats"]])
            blocks["card7"] = _table(["实际格式数", "Excel 与快照一致"], [[len(finding["formats"]), finding["excel_key_values_match"]]])
    (output / "previews.txt").write_text("\n\n".join(f"--8<-- [start:{name}]\n{body}\n--8<-- [end:{name}]" for name, body in blocks.items()) + "\n", encoding="utf-8")


def _allowlist(output: Path) -> list[Path]:
    """打包仅显式允许案例材料，不递归扫描任意文件或原始数据。"""
    allowed = set(TEXT_ARTIFACTS + SCRIPT_FILES + [f"case-{i}.json" for i in range(1, 8)] + [f"query-{i}.json" for i in range(1, 8)])
    allowed.update(f"{stem}.{suffix}" for stem in REPORT_STEMS for suffix in ("marsreport", "html", "xlsx"))
    allowed.add("manifest.json")
    return [output / name for name in sorted(allowed) if (output / name).is_file()]


def _finalize(output: Path, repository_root: Path | None = None) -> None:
    """汇总真实数据配置、报告身份与允许清单，在 ZIP 中排除个人 Skill/CSV/ID。"""
    preparation = json.loads((output / "preparation-summary.json").read_text(encoding="utf-8"))
    cases = {str(i): json.loads((output / f"case-{i}.json").read_text(encoding="utf-8")) for i in range(1, 8) if (output / f"case-{i}.json").exists()}
    _previews(output)
    _write_json(output / "summary.json", {"seed": preparation["seed"], "rows": preparation["rows"],
        "splits": {r["dataset_flag"]: r["total_count"] for r in preparation["partition_stats"]},
        "cases": {i: {"report_id": payload["report_id"], **payload["findings"]} for i, payload in cases.items()}})
    (output / "REPRODUCE.txt").write_text(
        "# Home Credit 2018 共享案例复现\n\n"
        "本包包含全量真实特征/原始标签的聚合统计；APPLY_DATE 是确定性模拟日期。\n"
        "保存报告的查询可使用兼容 MARS 0.0.28 及依赖，在任意解压目录运行（无需 CSV）：\n"
        "python task_cases.py --phase consume --output-dir .\n\n"
        "实际重算需安装本次修改后的当前源码（含规则无标签评估修复），并自行取得 Kaggle 三文件；不随本包发布：\n"
        "python task_cases.py --case all --data-dir /path/to/raw_data --output-dir reproduced\n"
        "也支持 --train-csv、--test-csv、--columns-description-csv。\n"
        "固定 seed=20261001，目标分区人数30000；输入指纹与真实分区见 preparation-summary.json。\n"
        "Train 用于拟合与 PSI，Val 用于规则筛选，Test/OOT 固定规则独立复核；July 无标签。\n"
        "--phase finalize 只收集已计算产物；ZIP 使用显式允许清单。\n"
        "ReportSnapshot 不恢复原始宽表或任意分析器；只加载可信报告。\n", encoding="utf-8")
    for name in SCRIPT_FILES:
        source_path = Path(__file__).with_name(name)
        if source_path.resolve() != (output / name).resolve():
            shutil.copyfile(source_path, output / name)
    for path in output.glob("*.html"):
        content = path.read_bytes()
        normalized = re.sub(rb"[ \t]+(?=\r?$)", b"", content, flags=re.MULTILINE)
        if normalized != content:
            path.write_bytes(normalized)
    repository = repository_root if repository_root is not None else next((p for p in Path(__file__).resolve().parents if (p / "pyproject.toml").exists() and (p / "src/mars").is_dir()), None)
    source_commit = None
    if repository is not None and shutil.which("git"):
        result = subprocess.run(["git", "rev-parse", "HEAD"], cwd=repository, check=False, capture_output=True, text=True)
        if result.returncode == 0:
            source_commit = result.stdout.strip()
    modules = ["analysis/score_cross.py", "analysis/score_cross_view.py", "analysis/profiler.py",
               "analysis/evaluator.py", "analysis/_risk_profile.py", "feature/selection/stats.py",
               "analysis/_evaluation/report_parts.py", "analysis/_evaluation/metrics.py",
               "compute/binning.py", "reporting/_metadata.py",
               "feature/selection/linear.py", "reporting/_query.py", "reporting/_artifact.py",
               "rule/workflow.py", "rule/evaluator.py", "rule/report.py", "rule/_statistics.py", "rule/_report_semantics.py"]
    environment = json.loads((output / "generation-environment.json").read_text(encoding="utf-8"))
    artifacts = [{"path": p.name, "size_bytes": p.stat().st_size, "sha256": hashlib.sha256(p.read_bytes()).hexdigest()} for p in _allowlist(output) if p.name != "manifest.json"]
    manifest = {"schema_version": 2, "source_commit": source_commit,
        "source_status": "known git checkout" if source_commit else "unknown commit; installed module fingerprints recorded",
        "command": "python task_cases.py --case all --data-dir examples/home-credit-2018/data/raw_data --output-dir current-cases",
        "data_config": {"source": "home-credit-2018", "rows": preparation["rows"], "seed": preparation["seed"],
                        "target_group_rows": preparation["partition_config"]["target_group_rows"], "reference": "0train",
                        "n_score_bins": 4, "scope_dimensions": ["dataset_flag", "APPLY_DATE simulated month", "TARGET"], "preparation": preparation},
        "dependencies": environment["dependencies"], "python": environment["python"], "generation_environment": environment,
        "code_fingerprints": {**{f"docs/snippets/{name}": _fingerprint(Path(__file__).with_name(name)) for name in SCRIPT_FILES},
                              **{f"src/mars/{name}": _fingerprint(Path(mars.__file__).resolve().parent / name) for name in modules}},
        "reports": {i: case["report_id"] for i, case in cases.items()}, "artifacts": artifacts,
        "zip_allowlist": sorted({"manifest.json", *(p.name for p in _allowlist(output))})}
    _write_json(output / "manifest.json", manifest)
    with zipfile.ZipFile(output / "cases.zip", "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for path in _allowlist(output):
            archive.write(path, arcname=path.name)


def main() -> None:
    """正式全量生成与只消费快照分进程；调试 fixture 不覆盖默认公开目录。"""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--case", choices=["all", *map(str, range(1, 8))], default="all")
    parser.add_argument("--data-dir", type=Path, default=DEFAULT_DATA_DIR)
    parser.add_argument("--train-csv", type=Path)
    parser.add_argument("--test-csv", type=Path)
    parser.add_argument("--columns-description-csv", type=Path)
    parser.add_argument("--seed", type=int, default=DEFAULT_SEED)
    parser.add_argument("--target-group-rows", type=int, default=30000)
    parser.add_argument("--fixture", action="store_true", help="仅测试非全量输入，必须显式另设输出目录")
    parser.add_argument("--output-dir", type=Path, default=Path("docs/assets/cases"))
    parser.add_argument("--phase", choices=["generate", "compute", "consume", "finalize"], default="generate")
    parser.add_argument("--repository-root", type=Path)
    args = parser.parse_args()
    output = args.output_dir
    if args.fixture and output.resolve() == Path("docs/assets/cases").resolve():
        parser.error("--fixture 必须使用另一个 --output-dir，不能覆盖公开资产。")
    output.mkdir(parents=True, exist_ok=True)
    if args.phase == "finalize":
        _finalize(output, args.repository_root)
        return
    if args.phase == "consume":
        _consume(output)
        _delivery(output)
        return
    cases = set(range(1, 8)) if args.case == "all" else {int(args.case)}
    if args.phase == "generate":
        if cases.intersection({1, 2, 3, 4, 5}):
            command = [sys.executable, str(Path(__file__).resolve()), "--phase", "compute", "--case", args.case,
                       "--data-dir", str(args.data_dir), "--seed", str(args.seed), "--target-group-rows", str(args.target_group_rows), "--output-dir", str(output.resolve())]
            for name in ("train_csv", "test_csv", "columns_description_csv"):
                if getattr(args, name) is not None:
                    command.extend(["--" + name.replace("_", "-"), str(getattr(args, name))])
            if args.fixture:
                command.append("--fixture")
            subprocess.run(command, check=True)
        if cases.intersection({4, 5, 6, 7}):
            subprocess.run([sys.executable, str(Path(__file__).resolve()), "--phase", "consume", "--output-dir", str(output.resolve())], check=True)
        _finalize(output, args.repository_root)
        print(f"已完成 Home Credit 案例 {args.case}；真实特征/标签，模拟日期；无 LLM。")
        return
    started = time.perf_counter()
    inputs = prepare_home_credit_data(args.data_dir, train_csv=args.train_csv, test_csv=args.test_csv,
        columns_description_csv=args.columns_description_csv, config=HomeCreditConfig(seed=args.seed, target_group_rows=args.target_group_rows), require_full_dataset=not args.fixture)
    timings: dict[str, Any] = {"preparation_seconds": time.perf_counter() - started, "rows": inputs.df.height,
                               "full_dataset": not args.fixture, "measurement": "wall time; no peak memory measurement"}
    _write_json(output / "preparation-summary.json", inputs.summary)
    _write_json(output / "feature-metadata.json", inputs.feature_metadata)
    _write_json(output / "generation-environment.json", _environment())
    cross: Report | None = None
    if cases.intersection({4, 5}):
        started = time.perf_counter()
        cross = _score_cross(inputs, output)
        timings["score_cross_seconds"] = time.perf_counter() - started
    for number, name, function in ((1, "profile", _profile), (2, "binning", _binning), (3, "selection", _selection)):
        if number in cases:
            started = time.perf_counter()
            print(f"开始 {name}: rows={inputs.df.height}, candidates={len(inputs.fcols)}", flush=True)
            function(inputs, output)
            timings[f"{name}_seconds"] = time.perf_counter() - started
    if 5 in cases:
        assert cross is not None
        started = time.perf_counter()
        _rules(inputs, cross, output)
        timings["rules_seconds"] = time.perf_counter() - started
    _write_json(output / "generation-timings.json", timings)


if __name__ == "__main__":
    main()
