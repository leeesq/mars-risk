"""同一份客群报告的原生消费、静态快照与导出数值回归。"""

from __future__ import annotations

import json
import subprocess
import sys
from datetime import date
from pathlib import Path

import openpyxl
import polars as pl
import pytest
from polars.testing import assert_frame_equal

from mars.analysis import profile_risk, profile_stats
from mars.reporting import load_report
from mars.reporting._artifact import snapshot_report


def test_shared_report_and_fixed_rules_in_fresh_processes(risk: object, tmp_path: Path) -> None:
    """真实 shared 快照在新进程消费，正式规则独立恢复且不得重新拟合。"""
    report = risk.report
    report_path = tmp_path / "shared.marsreport"
    rule_path = tmp_path / "rules.json"
    report.save(report_path, overwrite=True)
    risk.save_binning_rules(rule_path)
    report.write_excel(str(tmp_path / "native.xlsx"), segments=["seg2"], include_all=False)
    native = openpyxl.load_workbook(tmp_path / "native.xlsx", data_only=True)
    assert any(sheet.max_column == 22 and sheet.max_row > 1 for sheet in native)
    native.close()
    snapshot_report(report).write_excel(str(tmp_path / "original.xlsx"))
    original_png = tmp_path / "original.png"
    figure = report.build_risk_trend_figures("x", segments=["seg2"], include_all=False)[0]
    figure.savefig(original_png, dpi=150)
    import matplotlib.pyplot as plt

    plt.close(figure)
    expected_context = report.to_ai_context(
        tables=["summary"], segments=["seg2"], include_all=False, limit=1, max_chars=16000,
    )
    with pytest.raises(ValueError, match="max_chars"):
        report.to_ai_context(tables=["summary"], segments=["seg2"], include_all=False, max_chars=512)
    script = r'''
import json, sys
from pathlib import Path
from mars.reporting import load_report

root = Path(sys.argv[1])
r = load_report(root / "shared.marsreport")
page = r.query_page("summary", segments=["seg2"], include_all=False, limit=1)
assert page["returned_rows"] == 1
assert page["data"].equals(r.get_table(page["reference"]["table"], **page["reference"]["query"]))
context = r.to_ai_context(tables=["summary"], segments=["seg2"], include_all=False, limit=1, max_chars=16000)
assert len(context) <= 16000
try:
    r.to_ai_context(tables=["summary"], segments=["seg2"], include_all=False, max_chars=512)
except ValueError as exc:
    assert "max_chars" in str(exc)
else:
    raise AssertionError("budget should reject an undersized context")
evidence = json.loads(context)["evidence"][0]
assert r.get_table(evidence["reference"], **evidence["query"]).height == evidence["returned_rows"]
(root / "context.json").write_text(context, encoding="utf-8")
r.get_table("detail").write_parquet(root / "restored-detail.parquet")
r.get_table("calculation_status").write_parquet(root / "restored-status.parquet")
r.write_excel(str(root / "restored.xlsx"))
r.write_html(str(root / "restored.html"), segments=["seg2"], include_all=False)
figures = r.build_risk_trend_figures("x", segments=["seg2"], include_all=False, max_plots=1)
assert len(figures) == 1
figures[0].savefig(root / "restored.png", dpi=150)
assert (root / "restored.png").read_bytes().startswith(b"\x89PNG\r\n\x1a\n")
'''
    result = subprocess.run([sys.executable, "-c", script, str(tmp_path)], capture_output=True, text=True)
    assert result.returncode == 0, result.stderr
    assert json.loads((tmp_path / "context.json").read_text(encoding="utf-8")) == json.loads(expected_context)
    assert_frame_equal(pl.read_parquet(tmp_path / "restored-detail.parquet"), report.get_table("detail"))
    assert_frame_equal(pl.read_parquet(tmp_path / "restored-status.parquet"), report.get_table("calculation_status"))
    before = openpyxl.load_workbook(tmp_path / "original.xlsx", data_only=True)
    after = openpyxl.load_workbook(tmp_path / "restored.xlsx", data_only=True)
    assert before.sheetnames == after.sheetnames
    for left, right in zip(before, after):
        assert list(left.values) == list(right.values)
    before.close()
    after.close()
    assert "seg2" in (tmp_path / "restored.html").read_text(encoding="utf-8")
    restored_png = tmp_path / "restored.png"
    from PIL import Image, ImageChops

    with Image.open(original_png) as left_image, Image.open(restored_png) as right_image:
        assert ImageChops.difference(left_image.convert("RGB"), right_image.convert("RGB")).getbbox() is None

    rules_script = r'''
import sys
from datetime import date
from pathlib import Path
import polars as pl
from polars.testing import assert_frame_equal
from mars.analysis import MarsBinEvaluator, MarsRiskProfile
from mars.reporting import load_report

root = Path(sys.argv[1])
rules = MarsRiskProfile.load_binning_rules(root / "rules.json")
def forbidden(*args, **kwargs):
    raise AssertionError("fixed rules must not fit")
type(rules["all"]).fit = forbidden
frame = pl.DataFrame({"s": ["seg1"] * 40 + ["seg2"] * 40 + ["seg3"] * 40,
    "g": (["Train"] * 20 + ["Val"] * 20) * 3, "dt": [date(2026, 1, 1)] * 120,
    "x": [float(i % 10) for i in range(120)], "y": [int(i % 10 >= 5) for i in range(120)]})
run = MarsBinEvaluator().evaluate(frame, target="y", features=["x"], segment_col="s",
    group_col="g", time_col="dt", benchmark_df=frame.filter(pl.col("g") == "Train"),
    binner=rules["all"])
original = load_report(root / "shared.marsreport")
keys = ["scope_id", "y", "feature", "mars_group", "bin_index"]
assert_frame_equal(run.report.get_table("detail").sort(keys), original.get_table("detail").sort(keys))
'''
    result = subprocess.run([sys.executable, "-c", rules_script, str(tmp_path)], capture_output=True, text=True)
    assert result.returncode == 0, result.stderr


@pytest.fixture
def risk() -> object:
    frame = pl.DataFrame({
        "s": ["seg1"] * 40 + ["seg2"] * 40 + ["seg3"] * 40,
        "g": (["Train"] * 20 + ["Val"] * 20) * 3,
        "dt": [date(2026, 1, 1)] * 120,
        "x": [float(i % 10) for i in range(120)],
        "y": [int(i % 10 >= 5) for i in range(120)],
    })
    return profile_risk(frame, target="y", features=["x"], segment_col="s", group_col="g",
                        time_col="dt", benchmark_df=frame.filter(pl.col("g") == "Train"),
                        method="quantile", n_bins=3)


def test_native_and_snapshot_figure_scope_limit_and_dates(risk: object, tmp_path: Path) -> None:
    import matplotlib.pyplot as plt

    report = risk.report
    figures = report.build_risk_trend_figures("x", segments=["seg1", "seg2", "seg3"],
                                             include_all=True, max_plots=2)
    assert len(figures) == 2
    plt.close("all")
    paths = report.save_risk_trend_images(tmp_path, "x", segments=["seg2"], include_all=False,
                                         image_format="svg", max_plots=1)
    assert len(paths) == 1 and "seg_" in paths[0].name
    fragment = report.render_risk_trends_html("x", segments=["seg2"], include_all=False,
                                             max_plots=1)
    assert fragment.html.count('class="mars-risk-trend"') == 1
    report.save(tmp_path / "risk.marsreport")
    restored = load_report(tmp_path / "risk.marsreport")
    saved = restored.build_risk_trend_figures("x", segments=["seg2"], include_all=False)
    assert len(saved) == 1
    plt.close("all")


def test_native_excel_single_scope_keeps_template_and_values(risk: object, tmp_path: Path) -> None:
    report = risk.report
    before = report.detail_table.clone()
    with pytest.raises(ValueError, match="exactly one scope"):
        report.write_excel(str(tmp_path / "bad.xlsx"), segments=["seg2"])
    report.write_excel(str(tmp_path / "seg2.xlsx"), segments=["seg2"], include_all=False)
    workbook = openpyxl.load_workbook(tmp_path / "seg2.xlsx", data_only=True)
    sheets = [sheet for sheet in workbook if sheet.max_column == 22 and sheet.max_row > 1]
    assert sheets
    detail = sheets[0]
    headers = [cell.value for cell in detail[1]]
    assert len(headers) == 22 and "scope_id" not in headers
    assert "count" in headers
    count_index = headers.index("count")
    expected = report.get_table("detail", segments=["seg2"], include_all=False)
    assert sorted(row[count_index] for row in detail.iter_rows(min_row=2, values_only=True)) == sorted(expected["count"].to_list())
    assert report.detail_table.equals(before)
    workbook.close()


def test_html_keeps_scope_selection_and_actual_total_chart_budget(risk: object, tmp_path: Path) -> None:
    destination = tmp_path / "risk.html"
    risk.report.write_html(str(destination), segments=["seg1", "seg2", "seg3"], max_plots=2)
    from html import unescape

    source = unescape(destination.read_text(encoding="utf-8"))
    assert 'id="scope-select"' in source
    assert source.count('class="mars-chart-card"') == 2
    assert source.count("data-scope-id=") == 4


def test_stats_native_scope_views_keep_overview_identity(tmp_path: Path) -> None:
    report = profile_stats(pl.DataFrame({"s": ["seg1", "seg2"], "x": [1.0, 2.0]}),
                           features=["x"], segment_col="s", metrics=["mean", "missing"])
    all_view = report.show_overview().data
    segment_view = report.show_overview(segments=["seg2"], include_all=False).data
    assert len(all_view) == len(segment_view) == 1
    assert segment_view["mean"].iloc[0] == 2.0
    report.write_excel(str(tmp_path / "stats.xlsx"), segments=["seg2"], include_all=False)
    report.write_html(str(tmp_path / "stats.html"), segments=["seg2"], include_all=False)


def test_missing_scope_reference_is_unavailable_in_native_figure() -> None:
    import matplotlib.pyplot as plt

    frame = pl.DataFrame({"s": ["seg1"] * 20 + ["seg2"] * 20,
                          "g": ["Train", "Val"] * 20,
                          "dt": [date(2026, 1, 1)] * 40,
                          "x": [float(i % 10) for i in range(40)], "y": [i % 2 for i in range(40)]})
    report = profile_risk(frame, target="y", features=["x"], segment_col="s", group_col="g",
                           time_col="dt", benchmark_df=frame.filter(pl.col("s") == "seg1"),
                           risk_corr_baseline="benchmark", method="quantile", n_bins=3).report
    figures = report.build_risk_trend_figures("x", segments=["seg2"], include_all=False)
    text = " ".join(item.get_text() for ax in figures[0].axes for item in ax.texts)
    assert "PSI: n.a." in text
    assert "RC:n.a." in text
    assert "PSI: 0.00" not in text and "RC:1.00" not in text
    plt.close("all")


def test_scoped_html_asset_mode_keeps_assets_after_temporary_pages_removed(risk: object, tmp_path: Path) -> None:
    risk.report.write_html(str(tmp_path / "asset.html"), segments=["seg2"], include_all=False,
                            chart_embed_mode="asset", max_plots=1)
    assert len(list((tmp_path / "asset_assets").rglob("*.png"))) == 1
    source = (tmp_path / "asset.html").read_text(encoding="utf-8")
    assert "asset_assets/seg_" in source
    assert "data-src" in source


def test_single_scope_report_defaults_to_its_saved_identity(risk: object, tmp_path: Path) -> None:
    from mars.feature._scopes import resolve_scopes
    from mars.reporting._scope_view import scope_views

    report = risk.report
    identity = resolve_scopes(report._scopes, ["seg2"], False)[0]
    report.report_meta["default_scope_id"] = identity
    views = scope_views(report)
    assert views[0][0] == identity
    assert views[0][1].summary_table.height == 1
    assert views[0][1].report_id == report.report_id
    assert views[0][1].source == report.source
    reference = views[0][1].report_meta["scope_table_references"]["summary"]
    assert report.get_table(reference["table"], **reference["query"]).height == 1
    report.write_excel(str(tmp_path / "default.xlsx"))
    import matplotlib.pyplot as plt

    figures = report.build_risk_trend_figures("x")
    text = " ".join(item.get_text() for item in figures[0].texts)
    assert "seg2" in text
    assert "nan%" not in text
    plt.close("all")
