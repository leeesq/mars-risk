"""手工整数参考验证固定分段、风险分母和无需明细的规则回放。"""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
import polars as pl
import pytest
from openpyxl import load_workbook
from polars.testing import assert_frame_equal
from scipy.stats import binomtest

from mars.analysis import (
    cross_scores,
    evaluate_score_policy,
    get_score_bin_definitions,
    get_score_cell,
    show_score_matrix,
    write_score_cross_html,
)
from mars.reporting import load_report


def _data() -> pl.DataFrame:
    return pl.DataFrame(
        {
            "x": [4.0, 4.0, 1.0, 1.0, None, 4.0, 1.0, 4.0],
            "y": [1.0, 4.0, 1.0, 4.0, 1.0, None, float("inf"), 1.0],
            "bad": [0, 1, 1, None, 1, 0, 1, 0],
            "late": [None, 0, 0, None, 1, None, 0, 0],
            "split": ["TEST"] * 7 + ["OOT"],
        }
    )


def _report(data: pl.DataFrame | pd.DataFrame | None = None, **kwargs: Any) -> Any:
    return cross_scores(
        _data() if data is None else data,
        score_x="x",
        score_y="y",
        targets=["bad", "late"],
        score_directions={"x": "lower_risk", "y": "higher_risk"},
        cutpoints={"x": [2.0], "y": [2.0]},
        group_col="split",
        min_observed=2,
        **kwargs,
    )


def test_cells_margins_status_and_wilson_match_integer_reference() -> None:
    report = _report()
    cells = report.get_table("cells", filters={"target": "bad", "group": "TEST"})
    assert cells.height == 16
    assert cells["sample_count"].sum() == 7
    assert cells["observed_sample_count"].sum() == 6
    assert cells["bad_sample_count"].sum() == 4
    overall = report.get_table("overall", filters={"target": "bad", "group": "TEST"}).to_dicts()[0]
    assert overall["bad_rate"] == pytest.approx(4 / 6)
    rows = report.get_table("row_summary", filters={"target": "bad", "group": "TEST"})
    assert rows["sample_count"].sum() == 7
    detail = get_score_cell(report, "b1", "b0", filters={"target": "bad", "group": "TEST"})["page"][
        "data"
    ].to_dicts()[0]
    assert detail["bad_rate"] == 0 and detail["status"] == "low_sample"
    assert detail["row_bad_rate"] == pytest.approx(1 / 3)
    assert detail["delta_vs_row"] == pytest.approx(-1 / 3)
    interval = binomtest(0, 1).proportion_ci(method="wilson")
    assert detail["unweighted_ci_lower"] == pytest.approx(interval.low, abs=1e-15)
    assert detail["unweighted_ci_upper"] == pytest.approx(interval.high)
    assert (
        get_score_cell(report, "b0", "b1", filters={"target": "bad", "group": "TEST"})["page"][
            "data"
        ]["status"][0]
        == "unobserved"
    )
    assert (
        get_score_cell(report, "missing", "missing", filters={"target": "bad", "group": "TEST"})[
            "page"
        ]["data"]["status"][0]
        == "empty"
    )
    assert overall["sample_share"] == 1
    assert (
        report.get_table("overall", filters={"target": "late", "group": "TEST"})[
            "observed_sample_count"
        ][0]
        == 4
    )


@pytest.mark.parametrize("kind", ["x_only", "y_only", "and", "or", "staircase"])
def test_policy_matches_row_boolean_reference_and_load(kind: str, tmp_path: Path) -> None:
    report = _report()
    rule: dict[str, Any] = {"type": kind}
    if kind in {"x_only", "and", "or"}:
        rule["x_max_risk_rank"] = 1
    if kind in {"y_only", "and", "or"}:
        rule["y_max_risk_rank"] = 1
    if kind == "staircase":
        rule["steps"] = {
            "b1": {"action": "accept", "y_max_risk_rank": 1},
            "b0": {"action": "reject"},
        }
    baseline = {"type": "x_only", "x_max_risk_rank": 1}
    policy = evaluate_score_policy(report, rule, baseline=baseline)
    data = _data().to_dicts()
    for label in ["bad", "late"]:
        for group in ["TEST", "OOT"]:
            subset = [r for r in data if r["split"] == group]
            retained: list[dict[str, Any]] = []
            for row in subset:
                x = row["x"] is not None and row["x"] > 2
                y = row["y"] is not None and np.isfinite(row["y"]) and row["y"] <= 2
                keep = (
                    x
                    if kind == "x_only"
                    else y
                    if kind == "y_only"
                    else x or y
                    if kind == "or"
                    else x and y
                )
                if keep:
                    retained.append(row)
            result = policy.get_table(
                "summary",
                filters={"target": label, "group": group, "rule": "candidate", "retained": True},
            ).to_dicts()[0]
            assert result["sample_count"] == len(retained)
            assert result["observed_sample_count"] == sum(r[label] is not None for r in retained)
            assert result["bad_sample_count"] == sum(r[label] == 1 for r in retained)
            regions = policy.get_table("regions", filters={"target": label, "group": group})
            assert regions.height == 4 and regions["sample_count"].sum() == len(subset)
    path = tmp_path / "cross.marsreport"
    report.save(path)
    replay = evaluate_score_policy(load_report(path), rule, baseline=baseline)
    for table in policy.describe()["tables"]:
        assert_frame_equal(replay.get_table(table), policy.get_table(table))
    policy.save(tmp_path / "policy.marsreport")
    assert_frame_equal(
        load_report(tmp_path / "policy.marsreport").get_table("summary"),
        policy.get_table("summary"),
    )


def test_fixed_reference_boundaries_specials_constant_and_real_scopes() -> None:
    reference = pl.DataFrame({"x": [1.0, 1.0, 1.0, 2.0, 3.0, 3.0], "y": [5.0] * 6})
    data = pl.DataFrame(
        {
            "x": [-100.0, 1.0, 2.0, 3.0, 100.0, None],
            "y": [5.0, 5.0, 5.0, 5.0, 5.0, 5.0],
            "g": ["TEST"] * 3 + ["OOT"] * 3,
            "date": ["2026-01-01"] * 3 + ["2026-02-01"] * 3,
        }
    )
    report = cross_scores(
        data,
        score_x="x",
        score_y="y",
        score_directions={"x": "lower_risk", "y": "higher_risk"},
        binning_reference=reference,
        group_col="g",
        time_col="date",
        n_bins=5,
    )
    definitions = get_score_bin_definitions(report)
    assert definitions["y"]["actual_n_bins"] == 1
    assert report.describe()["parameters"]["actual_scope_count"] == 2
    assert report.get_table("overall")["sample_count"].sum() == 6
    assert set(report.get_table("cells")["status"]) == {"not_requested"}
    assert report.get_table("overall")["observed_sample_count"].null_count() == 2
    restored_bins = cross_scores(
        data,
        score_x="x",
        score_y="y",
        score_directions={"x": "lower_risk", "y": "higher_risk"},
        bin_definitions=definitions,
    )
    assert get_score_bin_definitions(restored_bins) == definitions
    boundary = cross_scores(
        pl.DataFrame({"x": [0.0, 0.5, 1.0, -99.0, 2.0, float("nan")], "y": [0.0] * 6}),
        score_x="x",
        score_y="y",
        score_directions={"x": "higher_risk", "y": "higher_risk"},
        cutpoints={"x": [0.5, 0.5], "y": []},
        special_values={"x": [-99.0]},
        probability_scores=["x"],
    )
    counts = {r["x_bin"]: r["sample_count"] for r in boundary.get_table("row_summary").to_dicts()}
    assert counts == {"b0": 2, "b1": 1, "missing": 1, "invalid": 1, "s0": 1}


def test_weighted_risk_amounts_zero_denominators_and_invalid_inputs() -> None:
    data = pl.DataFrame(
        {
            "x": [1, 1, 4, 4],
            "y": [1, 1, 4, 4],
            "bad": [0, 1, 0, 1],
            "split": ["TEST"] * 4,
            "w": [1.0, 3.0, 0.0, 0.0],
            "amt": [10.0, 30.0, -1.0, None],
        }
    )
    report = cross_scores(
        data,
        score_x="x",
        score_y="y",
        targets=["bad"],
        group_col="split",
        score_directions={"x": "higher_risk", "y": "higher_risk"},
        cutpoints={"x": [2], "y": [2]},
        weights_col="w",
        amount_col="amt",
    )
    cell = get_score_cell(report, "b0", "b0")["page"]["data"].to_dicts()[0]
    assert cell["bad_rate"] == 0.75 and cell["amt_bad_rate"] == 0.75
    assert cell["weighted_ci_status"] == "unsupported"
    assert cell["unweighted_ci_upper"] == pytest.approx(
        binomtest(1, 2).proportion_ci(method="wilson").high
    )
    invalid = get_score_cell(report, "b1", "b1")["page"]["data"].to_dicts()[0]
    assert invalid["status"] == "invalid_denominator" and invalid["bad_rate"] is None
    replay = evaluate_score_policy(report, {"type": "x_only", "x_max_risk_rank": 2})
    assert (
        replay.get_table("summary", filters={"retained": True, "rule": "candidate"})["bad_rate"][0]
        == 0.75
    )
    with pytest.raises(ValueError, match="Weights"):
        cross_scores(
            data.with_columns(pl.lit(-1).alias("w")),
            score_x="x",
            score_y="y",
            score_directions={"x": "higher_risk", "y": "higher_risk"},
            weights_col="w",
        )
    with pytest.raises(ValueError, match="continuous"):
        evaluate_score_policy(report, {"type": "x_only", "x_threshold": 2.5})
    rejected = evaluate_score_policy(report, {"type": "x_only", "x_max_risk_rank": 0})
    assert (
        rejected.get_table("summary", filters={"retained": True, "rule": "candidate"})[
            "sample_count"
        ][0]
        == 0
    )
    assert rejected.get_table("changes")["candidate_bad_rate"][0] is None


@pytest.mark.parametrize("bad", [0, 1])
def test_all_good_bad_and_unobserved_labels_are_supported(bad: int) -> None:
    report = _report(
        _data().with_columns(pl.lit(bad).alias("bad"), pl.lit(None, dtype=pl.Int8).alias("late"))
    )
    assert set(report.get_table("overall", filters={"target": "bad"})["bad_rate"]) == {float(bad)}
    assert set(report.get_table("overall", filters={"target": "late"})["status"]) == {"unobserved"}
    if bad == 0:
        assert report.get_table("cells")["lift_vs_overall"].null_count() == len(
            report.get_table("cells")
        )


def test_wide_projection_agent_evidence_and_exports(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    data = _data().to_pandas()
    for i in range(25):
        data[f"unrelated_{i}"] = "unneeded"
    original = pl.from_pandas
    projected: list[list[str]] = []

    def check_projection(frame: pd.DataFrame, *args: Any, **kwargs: Any) -> pl.DataFrame:
        projected.append(list(frame.columns))
        return original(frame, *args, **kwargs)

    monkeypatch.setattr(pl, "from_pandas", check_projection)
    report = _report(
        data,
        feature_metadata={"x": {"display_name": "<script>危险</script>", "data_source": "main"}},
    )
    assert projected == [["x", "y", "bad", "late", "split"]]
    page = report.query_page("cells", features="y", sources="main", limit=3)
    assert page["total_rows"] == len(report.get_table("cells"))
    context = report.to_ai_context(tables=["cells"], features="x", limit=4, max_chars=16000)
    payload = json.loads(context)
    assert len(context) <= 16000 and payload["evidence"][0]["truncated"]
    assert payload["description"]["tables"]["cells"]["feature_scope"] == ["x", "y"]
    assert report.get_feature("y")["tables"]["cells"].height > 0
    with pytest.raises(ValueError, match="Unknown feature"):
        report.get_feature("bad")
    path = tmp_path / "cross.marsreport"
    report.save(path)
    restored = load_report(path)
    assert (
        "omitted special samples=3"
        in show_score_matrix(restored, filters={"group": "TEST", "target": "bad"}).to_html()
    )
    replay = evaluate_score_policy(
        restored, {"type": "or", "x_max_risk_rank": 1, "y_max_risk_rank": 1}
    )
    html = tmp_path / "cross.html"
    write_score_cross_html(restored, html, policy_reports=[replay])
    text = html.read_text(encoding="utf-8")
    assert "https://" not in text and "<script>危险</script>" not in text
    assert "addEventListener" in text and "cell_decisions" in text
    restored.write_html(str(tmp_path / "dispatch.html"))
    restored.write_excel(str(tmp_path / "cross.xlsx"))
    workbook = load_workbook(tmp_path / "cross.xlsx")
    assert any("cells" in name for name in workbook.sheetnames)
    assert any("bins" in name for name in workbook.sheetnames)


def test_axis_fitting_is_once_and_saved_bins_never_fit(monkeypatch: pytest.MonkeyPatch) -> None:
    from mars.analysis import score_cross as module

    calls: list[str] = []
    original = module._fit_axis

    def counted(frame: pl.DataFrame, score: str, *args: Any, **kwargs: Any) -> dict[str, Any]:
        calls.append(score)
        return original(frame, score, *args, **kwargs)

    monkeypatch.setattr(module, "_fit_axis", counted)
    report = cross_scores(
        _data(),
        score_x="x",
        score_y="y",
        targets=["bad", "late"],
        score_directions={"x": "lower_risk", "y": "higher_risk"},
        group_col="split",
    )
    assert calls == ["x", "y"]
    # 显式/保存定义只校验配置，没有 reference 分数拟合。
    definitions = get_score_bin_definitions(report)
    cross_scores(
        _data(),
        score_x="x",
        score_y="y",
        score_directions={"x": "lower_risk", "y": "higher_risk"},
        bin_definitions=definitions,
    )
    assert calls == ["x", "y"]
    show_score_matrix(report, filters={"target": "bad", "group": "TEST"})
    evaluate_score_policy(report, {"type": "x_only", "x_max_risk_rank": 1})
    assert calls == ["x", "y"]


def test_empty_input_explicit_bins_and_replay_preserve_empty_scope() -> None:
    report = cross_scores(
        pl.DataFrame({"x": [], "y": []}),
        score_x="x",
        score_y="y",
        score_directions={"x": "higher_risk", "y": "lower_risk"},
        cutpoints={"x": [], "y": []},
    )
    assert report.get_table("cells").height == 0 and report.get_table("bins").height == 6
    replay = evaluate_score_policy(report, {"type": "x_only", "x_max_risk_rank": 0})
    assert replay.get_table("summary").height == 0
