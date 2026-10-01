"""从公共 API 生成真实浏览器夹具，再在独立进程只加载快照重导出。

先运行 ``--phase generate``，该进程结束后运行 ``--phase export``。
所有数据、快照和 HTML 均写入显式指定的仓库外验收目录。
"""

from __future__ import annotations

import argparse
import json
import os
import platform
import subprocess
import sys
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
import polars as pl
from openpyxl import load_workbook

from mars.analysis import (
    cross_scores,
    evaluate_score_policy,
    get_score_bin_definitions,
    get_score_cell,
    profile_risk,
    profile_stats,
    write_score_cross_html,
)
from mars.feature import MarsLinearSelector
from mars.reporting import Report, load_report, show_correlation_matrix
from mars.rule import MarsRuleMiningSpec, mine_rules

_X = 'champion_credit_probability_with_long_english_identity_<&>_quoted"'
_Y = "challenger_probability_中文交叉模型_</script>_literal"
_TITLE = '真实浏览器验收：很长报告标题 <&> "引号" </script> 字面文本与固定分箱'
_GROUP = '开发样本<&>"：very_long_arbitrary_named_cohort_with_saved_cutpoints_0123456789'


def _write_json(path: Path, value: dict[str, Any]) -> None:
    """将公共 API 已序列化证据写成便于审查的 JSON。"""
    path.write_text(
        json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False), encoding="utf-8"
    )


def _read_json(path: Path) -> dict[str, Any]:
    """读取本工具上阶段生成的标准 JSON。"""
    value: dict[str, Any] = json.loads(path.read_text(encoding="utf-8"))
    return value


def _facts(report: Report) -> dict[str, Any]:
    """通过公共上下文序列化完整小表，保留原生 schema 和状态。"""
    names = list(report.describe()["tables"])
    context: dict[str, Any] = json.loads(
        report.to_ai_context(
            tables=names,
            limit=100000,
            queries={name: {"limit": 100000} for name in names},
            max_chars=50000000,
        )
    )
    tables: dict[str, list[dict[str, Any]]] = {
        evidence["reference"]: evidence["rows"] for evidence in context["evidence"]
    }
    assert set(tables) == set(names)
    assert all(len(tables[name]) == len(report.get_table(name)) for name in names)
    schemas: dict[str, dict[str, str]] = {}
    for name in names:
        frame = report.get_table(name)
        schema = frame.schema if isinstance(frame, pl.DataFrame) else frame.dtypes.to_dict()
        schemas[name] = {str(column): str(dtype) for column, dtype in schema.items()}
    return {"description": report.describe(), "tables": tables, "schemas": schemas}


def _scoped_data() -> pl.DataFrame:
    """扩充现有 3×4 手工计数夹具，保留稀疏 scope、缺失和零权重。"""
    return pl.DataFrame(
        {
            _X: [
                0.9,
                0.9,
                0.9,
                0.9,
                0.5,
                0.1,
                0.5,
                None,
                float("nan"),
                -99.0,
                float("inf"),
                float("-inf"),
                -7.0,
                1.2,
                0.9,
                0.1,
                0.5,
                0.9,
                0.1,
                0.5,
            ],
            _Y: [
                0.1,
                0.1,
                0.3,
                0.9,
                0.6,
                0.9,
                0.1,
                0.3,
                -99.0,
                float("-inf"),
                0.1,
                float("inf"),
                -7.0,
                -0.2,
                0.6,
                0.1,
                0.3,
                0.1,
                0.9,
                0.3,
            ],
            "bad": [0, 0, 1, None, 0, 1, 1, 1, 0, 1, 0, 1, None, 1, 0, None, 1, 0, 0, 1],
            "later": [None, 0, 0, None, 1, 0, None, 0, 1, None, 1, 0, 0, 1, None, 0, 1, 1, None, 0],
            "weight": [
                1.0,
                3.0,
                2.0,
                1.0,
                0.0,
                4.0,
                2.0,
                3.0,
                5.0,
                2.0,
                1.0,
                4.0,
                1.0,
                2.0,
                2.0,
                0.0,
                7.0,
                1.0,
                2.0,
                3.0,
            ],
            "amount": [
                10.0,
                30.0,
                20.0,
                None,
                -1.0,
                40.0,
                20.0,
                30.0,
                float("nan"),
                20.0,
                10.0,
                40.0,
                None,
                20.0,
                20.0,
                0.0,
                70.0,
                10.0,
                20.0,
                30.0,
            ],
            "cohort": [_GROUP] * 14 + ["观察</script>"] * 3 + [_GROUP] * 2 + ["仅二月/独立样本"],
            "date": ["2026-01-01"] * 14 + ["2026-02-01"] * 3 + ["2026-03-01"] * 2 + ["2026-02-01"],
        }
    )


def _scoped_report(*, weighted: bool, labels: bool = True) -> Report:
    """使用公共固定切点入口生成不同表现分母的报告。"""
    return cross_scores(
        _scoped_data(),
        score_x=_X,
        score_y=_Y,
        score_directions={_X: "lower_risk", _Y: "higher_risk"},
        targets=["bad", "later"] if labels else None,
        cutpoints={_X: [0.3, 0.6], _Y: [0.2, 0.5, 0.8]},
        missing_values={_X: [-99.0, float("inf")], _Y: [-99.0, float("-inf")]},
        special_values={_X: [-7.0], _Y: [-7.0]},
        probability_scores=[_X, _Y],
        group_col="cohort",
        time_col="date",
        time_grain="month",
        weights_col="weight" if weighted else None,
        amount_col="amount",
        min_observed=2,
        confidence_level=0.9,
        feature_metadata={
            _X: {
                "display_name": '旧模型长中文展示名称<&>"，固定低风险方向',
                "data_source": "champion",
            },
            _Y: {"display_name": "新模型 </script> 字面文本 & 概率", "data_source": "challenger"},
        },
        business_context={
            "dataset_id": "browser-contract-fixture",
            "sample_unit": "synthetic row",
            "currency": "CNY",
        },
    )


def _small_report(*, empty: bool = False, zero: bool = False) -> Report:
    """合法空输入、一轴单箱和整体有效零分别使用当前公共入口。"""
    if empty:
        frame = pl.DataFrame({"x": [], "y": []})
        options: dict[str, Any] = {"cutpoints": {"x": [], "y": []}}
    else:
        frame = pl.DataFrame(
            {
                "x": [0.1, 0.4, 0.8, 0.9],
                "y": [0.1, 0.4, 0.8, None],
                "bad": [0, 0, 0, 0] if zero else [0, 1, None, 0],
            }
        )
        options = {"cutpoints": {"x": [], "y": [0.5]}, "targets": ["bad"], "min_observed": 1}
    return cross_scores(
        frame,
        score_x="x",
        score_y="y",
        score_directions={"x": "lower_risk", "y": "higher_risk"},
        **options,
    )


def _medium_report() -> Report:
    """中等展示矩阵用于布局与连续操作，不作为容量基准。"""
    rng = np.random.default_rng(20261002)
    size = 480
    frame = pl.DataFrame(
        {
            "x": rng.uniform(size=size),
            "y": rng.uniform(size=size),
            "bad": [None if i % 17 == 0 else int(i % 5 == 0) for i in range(size)],
            "later": [None if i % 4 == 0 else int(i % 7 == 0) for i in range(size)],
            "cohort": ["开发"] * 160 + ["验证/长组名" * 8] * 160 + ["观察"] * 160,
            "date": ["2026-01-01"] * 80
            + ["2026-02-01"] * 80
            + ["2026-02-01"] * 160
            + ["2026-03-01"] * 160,
        }
    )
    return cross_scores(
        frame,
        score_x="x",
        score_y="y",
        score_directions={"x": "lower_risk", "y": "higher_risk"},
        cutpoints={"x": [i / 8 for i in range(1, 8)], "y": [i / 10 for i in range(1, 10)]},
        targets=["bad", "later"],
        group_col="cohort",
        time_col="date",
        time_grain="month",
        min_observed=3,
    )


def _expressions(report: Report) -> list[str]:
    """按保存箱数构造合法规则，避免假设固定五箱。"""
    definitions = get_score_bin_definitions(report)
    nx, ny = (definitions[axis]["actual_n_bins"] for axis in ("x", "y"))
    rules = [
        f"X <= X{min(2, nx)} AND Y <= Y{min(3, ny)}",
        "X = X1 OR X = X1 OR Y >= Y1",
        "x = 1 or X = X1 and y = y1",
        "(X = 1 OR X = X1) AND Y = Y1",
        "X < 1",
        "X = X1 AND Y = Y1",
        f"X = X{min(2, nx)} AND Y = Y{min(3, ny)}",
        f"X = X1 AND Y = Y{ny}",
        f"X != {min(2, nx)} AND Y = 1",
        f"X >= X{nx} OR Y >= Y{ny}",
        "X=1 OR X=1 AND Y=1",
        "(X=1 OR X=1) AND Y=1",
    ]
    if nx > 1:
        rules += ["X=1 OR X=2 AND Y=1", "(X=1 OR X=2) AND Y=1"]
    return list(dict.fromkeys(rules))


def _score_facts(report: Report) -> dict[str, Any]:
    """保存公共 policy 期望值、固定定义与每个正常格子的重放证据。"""
    facts = _facts(report)
    facts["bin_definitions"] = report.describe()["parameters"]["bin_definitions"]
    expression_facts: dict[str, Any] = {}
    for expression in _expressions(report):
        policy = evaluate_score_policy(report, {"type": "expression", "expression": expression})
        data = _facts(policy)
        expression_facts[expression] = {
            "summary": [
                row
                for row in data["tables"]["summary"]
                if row["rule"] == "candidate" and row["retained"]
            ],
            "tables": data["tables"],
        }
    facts["expressions"] = expression_facts
    facts["evidence"] = []
    for cell in facts["tables"]["cells"]:
        if cell["x_risk_rank"] is None or cell["y_risk_rank"] is None:
            continue
        scope = {key: cell[key] for key in ("target", "group", "period")}
        page = get_score_cell(report, cell["x_bin"], cell["y_bin"], filters=scope)["page"]
        reference = page["reference"]
        replay = report.query_page(reference["table"], **reference["query"])
        assert replay["data"].to_dicts() == [cell]
        facts["evidence"].append(
            {"reference": reference, "query": reference["query"], "rows": [cell]}
        )
    return facts


def _save_score(name: str, report: Report, output: Path) -> dict[str, Any]:
    """保存生成阶段事实，不在生成进程进行加载。"""
    folder = output / name
    folder.mkdir(exist_ok=True)
    snapshot, original = folder / "report.marsreport", folder / "original.json"
    report.save(snapshot, overwrite=True)
    _write_json(original, _score_facts(report))
    return {
        "snapshot": str(snapshot),
        "original": str(original),
        "expected": str(folder / "expected.json"),
        "html": str(folder / "restored.html"),
        "excel": str(folder / "restored.xlsx"),
        "title": _TITLE if name in {"weighted", "standard"} else f"浏览器验收 · {name}",
    }


def _representatives(output: Path) -> dict[str, Any]:
    """复用教程的公共画像、规则与 Linear 相关性入口。"""
    folder = output / "representatives"
    folder.mkdir(exist_ok=True)
    values = list(range(80))
    frame = pl.DataFrame(
        {
            "income": values,
            "debt": [v % 3 for v in values],
            "bad": [None if v % 11 == 0 else int(v >= 60) for v in values],
            "later": [None] * 80,
            "split": ["开发<&>"] * 40 + ["观察</script>"] * 40,
        }
    )
    metadata = {
        "income": {
            "display_name": '收入<&>"很长业务中文说明' * 3,
            "data_source": "application",
            "unit": "CNY/month",
        },
        "debt": {"display_name": "负债</script>字面文本", "data_source": "credit"},
    }
    context = {
        "dataset_id": "representative-browser-synthetic",
        "labels": {"bad": {"definition": '模拟违约<&>"'}, "later": {"definition": "模拟未表现"}},
        "currency": "CNY",
    }
    reports: dict[str, Any] = {
        "profile": profile_stats(
            frame,
            features=["income", "debt"],
            metrics=["missing", "mean"],
            group_col="split",
            feature_metadata=metadata,
            business_context=context,
        ),
        "risk": profile_risk(
            frame,
            target=["bad", "later"],
            features=["income", "debt"],
            group_col="split",
            n_bins=2,
            feature_metadata=metadata,
            business_context=context,
        ).report,
        "rule": mine_rules(
            frame,
            target="bad",
            validation_df=frame.with_columns((pl.col("income") + 7).alias("income")),
            features=["income", "debt"],
            seed_rules=["income >= 65 AND debt >= 0", "income >= 70"],
            generators=[],
            spec=MarsRuleMiningSpec(iou_threshold=1.0),
        ).to_report(feature_metadata=metadata, business_context=context),
    }
    corr_data = pd.DataFrame(
        {"income": values, "debt": [-v for v in values], "independent": [v % 5 for v in values]}
    )
    selector = MarsLinearSelector(corr_thr=0.95).fit(
        corr_data,
        pd.Series([int(v >= 60) for v in values]),
        feature_metadata=metadata,
        business_context=context,
    )
    reports["correlation"] = selector.get_correlation_report()
    entries: dict[str, Any] = {}
    for name, report in reports.items():
        snapshot = folder / f"{name}.marsreport"
        report.save(snapshot, overwrite=True)
        _write_json(folder / f"{name}-original.json", _facts(report))
        if name == "risk":
            report.write_html(str(folder / "risk-original.html"), include_charts=False)
        else:
            report.write_html(str(folder / f"{name}-original.html"))
        entries[name] = {
            "snapshot": str(snapshot),
            "original": str(folder / f"{name}-original.json"),
            "expected": str(folder / f"{name}-expected.json"),
            "html": str(folder / f"{name}.html"),
            "original_html": str(folder / f"{name}-original.html"),
        }
    return entries


def _generate(output: Path) -> None:
    """生成进程只计算、保存，退出后交给 export 阶段加载。"""
    output.mkdir(parents=True, exist_ok=True)
    reports: dict[str, Report] = {
        "standard": _scoped_report(weighted=False),
        "weighted": _scoped_report(weighted=True),
        "no_labels": _scoped_report(weighted=True, labels=False),
        "one_bin": _small_report(),
        "empty": _small_report(empty=True),
        "zero_overall": _small_report(zero=True),
        "medium": _medium_report(),
    }
    manifest: dict[str, Any] = {
        "environment": {
            "python": sys.version,
            "platform": platform.platform(),
            "polars": pl.__version__,
            "numpy": np.__version__,
            "pandas": pd.__version__,
            "generation_pid": os.getpid(),
        },
        "fixtures": {name: _save_score(name, report, output) for name, report in reports.items()},
        "representatives": _representatives(output),
    }
    try:
        cross_scores(
            pl.DataFrame({"x": [], "y": []}),
            score_x="x",
            score_y="y",
            score_directions={"x": "higher_risk", "y": "higher_risk"},
        )
    except ValueError as error:
        manifest["empty_without_explicit_bins"] = {
            "accepted": False,
            "exception": type(error).__name__,
            "message": str(error),
        }
    else:
        manifest["empty_without_explicit_bins"] = {"accepted": True}
    _write_json(output / "manifest.json", manifest)
    print(
        json.dumps(
            {
                "phase": "generate",
                "fixtures": list(reports),
                "manifest": str(output / "manifest.json"),
            },
            ensure_ascii=False,
        )
    )


def _export(output: Path) -> None:
    """新进程只通过 load_report 读聚合快照，核对并生成浏览器产物。"""
    manifest = _read_json(output / "manifest.json")
    results: dict[str, Any] = {}
    for name, entry in manifest["fixtures"].items():
        report = load_report(entry["snapshot"])
        original, restored = _read_json(Path(entry["original"])), _score_facts(report)
        assert restored == original, name
        definitions = get_score_bin_definitions(report)
        max_x = min(2, definitions["x"]["actual_n_bins"])
        max_y = min(3, definitions["y"]["actual_n_bins"])
        baseline = {"type": "x_only", "x_max_risk_rank": max_x}
        policies = (
            []
            if name in {"empty", "no_labels"}
            else [
                evaluate_score_policy(
                    report,
                    {"type": "expression", "expression": f"X <= X{max_x} AND Y <= Y{max_y}"},
                    baseline=baseline,
                ),
                evaluate_score_policy(
                    report,
                    {
                        "type": "staircase",
                        "steps": {"b0": {"action": "accept", "y_max_risk_rank": max_y}},
                    },
                    baseline=baseline,
                ),
            ]
        )
        restored["policies"] = [_facts(policy) for policy in policies]
        write_score_cross_html(
            report, entry["html"], report_name=entry["title"], policy_reports=policies
        )
        report.write_excel(entry["excel"])
        workbook = load_workbook(entry["excel"], read_only=True)
        cell_sheet = next(sheet for sheet in workbook if sheet.title.endswith("_cells"))
        headers = next(cell_sheet.values)
        assert "sample_count" in headers and "bad_rate" in headers
        assert cell_sheet.max_row == len(restored["tables"]["cells"]) + 1
        workbook.close()
        _write_json(Path(entry["expected"]), restored)
        results[name] = {
            "identity_description_schema_tables_status_rules_evidence_equal": True,
            "scope_rows": len(restored["tables"]["overall"]),
            "normal_cell_evidence_replayed": len(restored["evidence"]),
            "excel_headers_and_row_count": True,
        }
    for name, entry in manifest["representatives"].items():
        report = load_report(entry["snapshot"])
        restored = _facts(report)
        assert restored == _read_json(Path(entry["original"])), name
        report.write_html(entry["html"], report_name=f"真实浏览器验收 · {name} <&> </script>")
        if name == "correlation":
            matrix = show_correlation_matrix(report)
            matrix_path = Path(entry["html"]).with_name("correlation-matrix.html")
            matrix_path.write_text(matrix.to_html(), encoding="utf-8")
            entry["matrix_html"] = str(matrix_path)
            restored["matrix"] = matrix.data.to_dict(orient="index")
        _write_json(Path(entry["expected"]), restored)
    # 保留教程现有 CLI 的独立新进程路径，不在验收工具中另写重导出算法。
    snippet_output = output / "snippet-reexport"
    rule = "X <= X2 AND Y <= Y3"
    command = [
        sys.executable,
        "docs/snippets/correlation_and_score_cross.py",
        "--report",
        manifest["fixtures"]["weighted"]["snapshot"],
        "--rule",
        rule,
        "--output",
        str(snippet_output),
    ]
    completed = subprocess.run(
        command,
        cwd=Path(__file__).resolve().parents[2],
        check=True,
        capture_output=True,
        text=True,
        encoding="utf-8",
    )
    (output / "fixtures-snippet.log").write_text(
        completed.stdout + completed.stderr, encoding="utf-8"
    )
    expected_rule = _read_json(Path(manifest["fixtures"]["weighted"]["expected"]))
    replay_policy = load_report(snippet_output / "score-policy.marsreport")
    assert _facts(replay_policy)["tables"] == expected_rule["expressions"][rule]["tables"]
    manifest["snippet_reexport"] = {
        "html": str(snippet_output / "score-cross.html"),
        "excel": str(snippet_output / "score-cross.xlsx"),
        "snapshot": manifest["fixtures"]["weighted"]["snapshot"],
        "policy_snapshot": str(snippet_output / "score-policy.marsreport"),
        "rule": rule,
        "command": command,
        "passed": True,
    }
    history = manifest["environment"].setdefault("export_pid_history", [])
    previous = manifest["environment"].get("export_pid")
    for pid in [previous, os.getpid()]:
        if pid is not None and pid not in history:
            history.append(pid)
    manifest["environment"]["export_pid"] = os.getpid()
    manifest["restore_results"] = results
    _write_json(output / "manifest.json", manifest)
    print(json.dumps({"phase": "export", "results": results}, ensure_ascii=False))


def main() -> None:
    """执行独立的生成或加载阶段，不在浏览器模拟 Python 环境。"""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--phase", choices=["generate", "export"], required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    if args.phase == "generate":
        _generate(args.output.resolve())
    else:
        _export(args.output.resolve())


if __name__ == "__main__":
    main()
