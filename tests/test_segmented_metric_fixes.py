"""分客群指标适用人口、独立参考与局部空值的数值回归。"""
from __future__ import annotations

import re
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl
import polars as pl
import pytest
from polars.testing import assert_frame_equal

from mars.analysis import MarsBinEvaluator, MarsRiskProfile, profile_risk
from mars.compute import to_polars_frame
from mars.reporting import load_report
from mars.reporting._binning_html import _BinningHtmlRenderer
from mars.utils.html import format_html_value


def population() -> pl.DataFrame:
    """提供各期均有三箱、两类标签的确定性人口。"""
    return pl.DataFrame([
        {"s": s, "g": g, "x": float(i), "z": float(i), "y": int(i >= 30), "w": 1.0, "amt": 10.0}
        for s in ("seg1", "seg2") for g in ("01", "02") for i in range(60)
    ])


@pytest.mark.parametrize("mode", ["shared", "per_segment"])
@pytest.mark.parametrize("ks_method", ["raw", "binned"])
@pytest.mark.parametrize("labels", ["zero", "one", "unobserved", "zero_bad_weight"])
def test_undefined_separation_preserves_population_and_fixed_rules(
    mode: str, ks_method: str, labels: str,
) -> None:
    frame = population()
    benchmark = frame.filter(pl.col("g") == "01")
    label = pl.lit(None, dtype=pl.Int64) if labels == "unobserved" else pl.lit(1 if labels == "one" else 0)
    if labels == "zero_bad_weight":
        frame = frame.with_columns(
            pl.when((pl.col("s") == "seg2") & (pl.col("y") == 1)).then(0.0).otherwise(1.0).alias("w")
        )
    else:
        frame = frame.with_columns(pl.when(pl.col("s") == "seg2").then(label).otherwise(pl.col("y")).alias("y"))
    run = profile_risk(
        frame, target="y", features=["x"], segment_col="s", group_col="g", benchmark_df=benchmark,
        binning_scope=mode, ks_method=ks_method, method="quantile", n_bins=3, n_jobs=1,
        weights_col="w", amount_col="amt",
    )
    report = run.report
    summary = report.get_table("summary", segments=["seg2"], include_all=False)
    assert summary.select("iv", "ks", "auc").null_count().row(0) == (1, 1, 1)
    detail = report.get_table("detail", segments=["seg2"], include_all=False)
    for column in ["iv_bin", "ks_bin", "auc_bin"]:
        assert detail[column].null_count() == detail.height
    for metric in ["iv", "ks", "auc"]:
        if metric in report.trend_tables:
            trend = report.get_table(f"trend.{metric}", segments=["seg2"], include_all=False)
            assert trend.select("01", "02", "Total").null_count().row(0) == (trend.height,) * 3
    status = report.get_table("calculation_status", segments=["seg2"], include_all=False)
    assert ("no_observed_labels" if labels == "unobserved" else "single_observed_class") in status["reason"].to_list()
    assert summary["psi_max"].item() is not None
    if labels != "unobserved":
        assert set(detail["bad_rate"].drop_nulls()) == {1.0 if labels == "one" else 0.0}
        assert detail.filter((pl.col("bin_index") == 9999) & (pl.col("mars_group") == "Total"))["amt_bad_rate"].item() == (0.5 if labels == "zero_bad_weight" else 1.0 if labels == "one" else 0.0)
        if labels == "zero":
            assert detail["lift_amt"].null_count() == detail.height
        elif labels == "zero_bad_weight":
            assert detail["lift_amt"].drop_nulls().max() > 0

    valid = report.get_table("summary", segments=["seg1"], include_all=False)
    assert valid["iv"].item() > 0
    assert valid["ks"].item() > 0
    assert run.scopes is not None
    identity = report.get_table("scopes").filter(pl.col("display_name") == "seg2")["scope_id"].item()
    rule = run.binners_by_scope[identity]
    mapping = dict(rule.bin_woes_["x"])
    assert len(set(mapping.values())) > 1
    # 当前客群的统计未定义不能改写从开发人口学到的 WOE。
    assert rule.bin_woes_["x"] == mapping


def test_single_class_group_keeps_valid_total_and_other_target() -> None:
    frame = population().with_columns(pl.col("y").alias("other"))
    frame = frame.with_columns(
        pl.when((pl.col("s") == "seg2") & (pl.col("g") == "02")).then(1).otherwise(pl.col("y")).alias("y")
    )
    run = profile_risk(frame, target=["y", "other"], features=["x"], segment_col="s", group_col="g", method="quantile", n_bins=3, n_jobs=1)
    detail = run.report.get_table("detail", segments=["seg2"], include_all=False)
    single = detail.filter((pl.col("y") == "y") & (pl.col("mars_group") == "02"))
    assert single["ks_bin"].null_count() == single.height
    assert detail.filter((pl.col("y") == "y") & (pl.col("mars_group") == "Total"))["ks_bin"].null_count() == 0
    assert detail.filter(pl.col("y") == "other")["ks_bin"].null_count() == 0


@pytest.mark.parametrize("baseline", ["total", "first_group", "benchmark"])
def test_missing_psi_benchmark_only_nulls_its_own_metric(baseline: str) -> None:
    frame = population()
    benchmark = frame.filter((pl.col("s") == "seg1") & (pl.col("g") == "01"))
    report = profile_risk(frame, target="y", features=["x"], segment_col="s", group_col="g", benchmark_df=benchmark, risk_corr_baseline=baseline, method="quantile", n_bins=3, n_jobs=1).report
    summary = report.get_table("summary", segments=["seg2"], include_all=False)
    assert summary["psi_max"].item() is None
    assert summary["rc_min"].item() == (None if baseline == "benchmark" else 1.0)
    statuses = report.get_table("calculation_status", segments=["seg2"], include_all=False)
    rc = statuses.filter(pl.col("metric") == "risk_corr")
    if baseline == "benchmark":
        assert rc["reason"].to_list() == ["missing_reference_population"]
    else:
        assert rc.is_empty()
    reference = report.get_table("risk_corr_reference", segments=["seg2"], include_all=False)
    if baseline != "benchmark":
        assert reference["base_br"].drop_nulls().n_unique() == 3


def test_missing_first_group_keeps_valid_explicit_psi_reference() -> None:
    benchmark = population().filter(pl.col("g") == "01")
    frame = population().filter(~((pl.col("s") == "seg2") & (pl.col("g") == "01")))
    report = profile_risk(frame, target="y", features=["x"], segment_col="s", group_col="g", benchmark_df=benchmark, risk_corr_baseline="first_group", method="quantile", n_bins=3, n_jobs=1).report
    summary = report.get_table("summary", segments=["seg2"], include_all=False)
    assert summary["psi_max"].item() == 0.0
    assert summary["rc_min"].item() is None
    status = report.get_table("calculation_status", segments=["seg2"], include_all=False)
    assert status.filter(pl.col("metric") == "psi").is_empty()
    assert status.filter(pl.col("metric") == "risk_corr")["reason"].to_list() == ["missing_reference_population"]


@pytest.mark.parametrize("kind", ["missing", "special", "zero_weight"])
@pytest.mark.parametrize("include", [False, True])
@pytest.mark.parametrize("backend", ["polars", "pandas"])
def test_psi_population_is_validated_per_group(kind: str, include: bool, backend: str) -> None:
    frame = population()
    benchmark = frame.filter(pl.col("g") == "01")
    affected = (pl.col("s") == "seg2") & (pl.col("g") == "02")
    if kind == "zero_weight":
        frame = frame.with_columns(pl.when(affected).then(0.0).otherwise(1.0).alias("w"))
    else:
        frame = frame.with_columns(pl.when(affected).then(None if kind == "missing" else -999.0).otherwise(pl.col("x")).alias("x"))
    report = profile_risk(
        frame.to_pandas() if backend == "pandas" else frame, target="y", features=["x", "z"], segment_col="s", group_col="g",
        benchmark_df=benchmark, method="quantile", n_bins=3, n_jobs=1, weights_col="w", special_values=[-999.0],
        psi_include_missing=include, psi_include_special=include,
    ).report
    trend = to_polars_frame(report.get_table("trend.psi", segments=["seg2"], include_all=False))
    row = trend.filter(pl.col("feature") == "x").row(0, named=True)
    undefined = kind == "zero_weight" or not include
    assert (row["02"] is None) == undefined
    assert row["01"] == 0.0
    assert row["Total"] is not None and row["Total"] >= 0.0
    if undefined:
        assert row["Total"] == 0.0
    assert to_polars_frame(report.get_table("summary", segments=["seg2"], include_all=False))["psi_max"].null_count() == 0
    status = to_polars_frame(report.get_table("calculation_status", segments=["seg2"], include_all=False))
    status = status.filter((pl.col("feature") == "x") & (pl.col("metric") == "psi"))
    assert status.height == int(undefined)
    if undefined:
        assert status["group"].to_list() == ["02"]
        assert status["reason"].item() == ("zero_current_weight" if kind == "zero_weight" else "no_effective_current_population")
    if kind != "zero_weight":
        assert trend.filter(pl.col("feature") == "z")["02"].item() == 0.0


def test_local_nulls_roundtrip_and_excel_binned_ks(tmp_path: Path) -> None:
    frame = population().with_columns(pl.when(pl.col("s") == "seg2").then(0).otherwise(pl.col("y")).alias("y"), pl.lit(date(2026, 1, 1)).alias("date"))
    run = profile_risk(frame, target="y", features=["x"], segment_col="s", group_col="g", time_col="date", method="quantile", n_bins=3, n_jobs=1, ks_method="raw")
    path = tmp_path / "undefined.marsreport"
    run.report.save(path)
    restored = load_report(path)
    assert_frame_equal(run.report.get_table("detail"), restored.get_table("detail"))
    run.report.write_excel(str(tmp_path / "undefined.xlsx"), segments=["seg2"], include_all=False)
    workbook = openpyxl.load_workbook(tmp_path / "undefined.xlsx", data_only=True)
    detail = next(sheet for sheet in workbook if sheet.max_column == 22 and sheet.max_row > 1)
    headers = [cell.value for cell in detail[1]]
    for name in ["ks", "ks_bin", "auc_bin", "iv_bin"]:
        if name in headers:
            index = headers.index(name)
            assert all(row[index] is None for row in detail.iter_rows(min_row=2, values_only=True))
    workbook.close()
    html = _BinningHtmlRenderer._build_grouped_pivot_section_html(
        run.report.get_table("detail", segments=["seg2"], include_all=False).to_pandas(),
        group_col="mars_group", feature_sources={"x": "UNMAPPED"},
    )
    iv_cells = re.findall(r'data-col="iv\|[^"]*"[^>]*>([^<]*)</td>', html)
    assert iv_cells and set(iv_cells) == {format_html_value(float("nan"))}
    rules = tmp_path / "rules.json"
    run.save_binning_rules(rules)
    fixed = MarsRiskProfile.load_binning_rules(rules)
    before = {scope: rule.to_dict() for scope, rule in fixed.items()}
    MarsBinEvaluator().evaluate(frame, target="y", features=["x"], segment_col="s", group_col="g", binner=fixed["all"])
    assert before == {scope: rule.to_dict() for scope, rule in fixed.items()}
    figures = restored.build_risk_trend_figures("x", segments=["seg2"], include_all=False)
    text = " ".join(item.get_text() for ax in figures[0].axes for item in ax.texts)
    assert "KS:n.a." in text or "KS: n.a." in text
    assert "IV: n.a." in text and "AUC: n.a." in text
    assert "IV: 0.00" not in text and "AUC: 1.00" not in text
    import matplotlib.pyplot as plt
    plt.close("all")


def test_plain_single_class_still_rejects_and_supervised_fit_stays_local() -> None:
    frame = population().with_columns(pl.lit(0).alias("y"))
    with pytest.raises(ValueError, match="2 observed classes"):
        profile_risk(frame, target="y", features=["x"], method="quantile", n_jobs=1)
    frame = population().with_columns(pl.when(pl.col("s") == "seg2").then(0).otherwise(pl.col("y")).alias("y"))
    run = profile_risk(frame, target="y", features=["x"], segment_col="s", binning_scope="per_segment", method="cart", n_jobs=1)
    statuses = run.report.get_table("calculation_status", segments=["seg2"], include_all=False)
    assert statuses["reason"].to_list() == ["single_class_fit"]



@pytest.mark.parametrize("kind,reason", [
    ("unobserved", "no_observed_reference_labels"),
    ("constant", "constant_reference_bad_rate"),
    ("zero_weight", "zero_reference_weight"),
])
def test_rc_reference_failure_reasons_are_independent(kind: str, reason: str) -> None:
    frame = population()
    benchmark = frame.filter(pl.col("g") == "01")
    if kind == "zero_weight":
        benchmark = benchmark.with_columns(pl.when(pl.col("s") == "seg2").then(0.0).otherwise(pl.col("w")).alias("w"))
    else:
        benchmark = benchmark.with_columns(pl.when(pl.col("s") == "seg2").then(None if kind == "unobserved" else 0).otherwise(pl.col("y")).alias("y"))
    report = profile_risk(frame, target="y", features=["x"], segment_col="s", group_col="g", benchmark_df=benchmark, weights_col="w", risk_corr_baseline="benchmark", method="quantile", n_bins=3, n_jobs=1).report
    summary = report.get_table("summary", segments=["seg2"], include_all=False)
    assert summary["rc_min"].item() is None
    assert (summary["psi_max"].item() is None) == (kind == "zero_weight")
    status = report.get_table("calculation_status", segments=["seg2"], include_all=False).filter(pl.col("metric") == "risk_corr")
    assert set(status["reason"]) == {reason}


def test_plain_psi_uses_the_same_population_guard_and_null_sum() -> None:
    frame = population().filter(pl.col("s") == "seg1")
    benchmark = frame.filter(pl.col("g") == "01")
    frame = frame.with_columns(pl.lit(None, dtype=pl.Float64).alias("x"))
    report = profile_risk(frame, target="y", features=["x", "z"], group_col="g", benchmark_df=benchmark, method="quantile", n_bins=3, n_jobs=1).report
    summary = report.get_table("summary")
    assert summary.filter(pl.col("feature") == "x")["psi_max"].item() is None
    assert summary.filter(pl.col("feature") == "z")["psi_max"].item() == 0.0
    trend = report.get_table("trend.psi").filter(pl.col("feature") == "x")
    assert trend.select("01", "02", "Total").null_count().row(0) == (1, 1, 1)
    status = report.get_table("calculation_status").filter(pl.col("metric") == "psi")
    assert set(status["group"]) == {"01", "02", "Total"}



def test_index_ordered_ks_uses_only_its_normal_bin_population() -> None:
    frame = population()
    benchmark = frame.filter(pl.col("g") == "01")
    frame = frame.with_columns(pl.when((pl.col("s") == "seg2") & (pl.col("y") == 0)).then(None).otherwise(pl.col("x")).alias("x"))
    report = profile_risk(frame, target="y", features=["x"], segment_col="s", group_col="g", benchmark_df=benchmark, ordered_metric_sort_by="bin_index", method="quantile", n_bins=3, n_jobs=1).report
    summary = report.get_table("summary", segments=["seg2"], include_all=False)
    assert summary["ks"].item() is None and summary["auc"].item() is None
    assert summary["iv"].item() > 0
    assert report.get_table("detail", segments=["seg2"], include_all=False)["ks_bin"].null_count() == report.get_table("detail", segments=["seg2"], include_all=False).height
    status = report.get_table("calculation_status", segments=["seg2"], include_all=False)
    assert set(status.filter(pl.col("metric") == "ordered_class_separation")["reason"]) == {"single_observed_class_in_normal_bins"}



@pytest.mark.parametrize("backend", ["polars", "pandas"])
def test_rc_empty_normal_group_has_own_status(backend: str) -> None:
    frame = population()
    benchmark = frame.filter(pl.col("g") == "01")
    frame = frame.with_columns(pl.when((pl.col("s") == "seg2") & (pl.col("g") == "02")).then(None).otherwise(pl.col("x")).alias("x"))
    report = profile_risk(frame.to_pandas() if backend == "pandas" else frame, target="y", features=["x"], segment_col="s", group_col="g", benchmark_df=benchmark, risk_corr_baseline="benchmark", method="quantile", n_bins=3, n_jobs=1).report
    trend = to_polars_frame(report.get_table("trend.risk_corr", segments=["seg2"], include_all=False))
    assert trend["01"].item() == trend["Total"].item() == 1.0
    assert trend["02"].item() is None
    status = to_polars_frame(report.get_table("calculation_status", segments=["seg2"], include_all=False))
    rc = status.filter(pl.col("metric") == "risk_corr")
    assert rc["group"].to_list() == ["02"]
    assert rc["reason"].to_list() == ["no_normal_bins"]
    reference = to_polars_frame(report.get_table("risk_corr_reference", segments=["seg2"], include_all=False))
    assert reference["base_br"].drop_nulls().n_unique() == 3



def test_rc_disjoint_valid_reference_bins_are_not_unobserved_reference() -> None:
    def make_frame(start: int, stop: int, threshold: int) -> pl.DataFrame:
        """整体支持全部切点，seg2 在当前与基准占据互不相交的箱。"""
        return pl.DataFrame([
            {"s": "seg1", "g": "02", "x": float(i), "y": int(i >= 50)} for i in range(101)
        ] + [
            {"s": "seg2", "g": "02", "x": float(i), "y": int(i >= threshold)} for i in range(start, stop)
        ])
    frame = make_frame(75, 100, 87)
    benchmark = make_frame(0, 25, 12)
    report = profile_risk(frame, target="y", features=["x"], segment_col="s", group_col="g", benchmark_df=benchmark, risk_corr_baseline="benchmark", method="uniform", n_bins=8, n_jobs=1).report
    reference = report.get_table("risk_corr_reference", segments=["seg2"], include_all=False)
    assert reference["base_br"].drop_nulls().len() >= 2
    status = report.get_table("calculation_status", segments=["seg2"], include_all=False).filter(pl.col("metric") == "risk_corr")
    assert set(status["group"]) == {"02", "Total"}
    assert set(status["reason"]) == {"insufficient_valid_bin_pairs"}
    summary = report.get_table("summary", segments=["seg2"], include_all=False)
    assert summary["rc_min"].item() is None and summary["psi_max"].item() is not None



def test_feature_start_rc_reference_does_not_require_explicit_benchmark() -> None:
    frame = pl.DataFrame([
        {"s": s, "dt": datetime(2026, 1, 1) + timedelta(days=day),
         "x": None if day < 4 else float(index % 2), "y": 0 if day < 4 else index % 2}
        for s in ("seg1", "seg2") for day in range(8) for index in range(20)
    ])
    report = profile_risk(frame, target="y", features=["x"], segment_col="s", time_col="dt", time_grain="month", method="quantile", n_bins=2, n_jobs=1, feature_start_aware_reference=True, risk_corr_baseline="benchmark").report
    assert report.get_table("summary")["rc_min"].to_list() == pytest.approx([1.0, 1.0, 1.0], rel=0, abs=1e-12)
    assert report.get_table("summary")["psi_max"].to_list() == [0.0, 0.0, 0.0]
    assert report.get_table("calculation_status").filter(pl.col("metric") == "risk_corr").is_empty()
    assert all(meta["risk_corr_reference_source"] == "feature_start_reference" for meta in report.report_meta["scope_metadata"].values())
