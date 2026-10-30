"""画像 PSI 局部空人口在查询、快照与静态导出中的消费合同。"""

from __future__ import annotations

import json
import os
import subprocess
import sys
from html.parser import HTMLParser
from pathlib import Path
from typing import Any

import openpyxl
import pandas as pd
import polars as pl
import pytest
from polars.testing import assert_frame_equal

from mars.analysis import MarsDataProfiler
from mars.feature._scopes import resolve_scopes
from mars.reporting import MarsProfileReport, ReportSnapshot, load_report
from mars.reporting._query import ReportFrame
from mars.reporting._scope_view import scope_views


def _polars(frame: ReportFrame) -> pl.DataFrame:
    """只转换有限查询结果，比较跨后端的统计身份与空值。"""
    return frame if isinstance(frame, pl.DataFrame) else pl.from_pandas(frame)


class _HtmlTables(HTMLParser):
    """读取实际导出表头、单元格及缺失排序属性，递归读取客群 iframe。"""

    def __init__(self) -> None:
        super().__init__()
        self.tables: list[list[list[dict[str, str]]]] = []
        self._table: list[list[dict[str, str]]] | None = None
        self._row: list[dict[str, str]] | None = None
        self._cell: dict[str, str] | None = None

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        attributes = {key: value or "" for key, value in attrs}
        if tag == "iframe" and "srcdoc" in attributes:
            child = _HtmlTables()
            child.feed(attributes["srcdoc"])
            self.tables.extend(child.tables)
        if tag == "table":
            self._table = []
        elif tag == "tr" and self._table is not None:
            self._row = []
        elif tag in {"td", "th"} and self._row is not None:
            self._cell = {**attributes, "text": ""}

    def handle_data(self, data: str) -> None:
        if self._cell is not None:
            self._cell["text"] += data

    def handle_endtag(self, tag: str) -> None:
        if tag in {"td", "th"} and self._cell is not None and self._row is not None:
            self._row.append(self._cell)
            self._cell = None
        elif tag == "tr" and self._row is not None and self._table is not None:
            self._table.append(self._row)
            self._row = None
        elif tag == "table" and self._table is not None:
            self.tables.append(self._table)
            self._table = None


@pytest.fixture(params=["polars", "pandas"])
def population_report(request: pytest.FixtureRequest) -> MarsProfileReport:
    """两客群只有 seg2/x/02 无适用人口，其 Total 与其他结果均有效。"""
    frame = pl.DataFrame([
        {"s": segment, "g": group, "x": float(i), "z": float(i)}
        for segment in ("seg1", "seg2") for group in ("01", "02") for i in range(20)
    ])
    benchmark = frame.filter(pl.col("g") == "01")
    current = frame.with_columns(
        pl.when((pl.col("s") == "seg2") & (pl.col("g") == "02"))
        .then(None).otherwise(pl.col("x")).alias("x")
    )
    return MarsDataProfiler(psi_batch_size=2).generate_profile(
        current.to_pandas() if request.param == "pandas" else current,
        features=["x", "z"], metrics=["psi"], segment_col="s", group_col="g",
        benchmark_df=benchmark.to_pandas() if request.param == "pandas" else benchmark,
        feature_metadata={
            "x": {"display_name": "局部空人口", "data_source": "source_x", "unit": "元"},
            "z": {"display_name": "正常特征", "data_source": "source_z"},
        },
    )


def _selected_query() -> dict[str, Any]:
    """固定可重放业务查询，身份与原数值列同时保留。"""
    return {
        "features": "x", "segments": ["seg2"], "include_all": False,
        "columns": ["feature", "display_name", "data_source", "dtype", "01", "02", "total"],
    }


def _assert_queries(report: MarsProfileReport | ReportSnapshot) -> None:
    """同一统计与局部状态用于查询、分页、JSON 证据和引用重放。"""
    query = _selected_query()
    selected = _polars(report.get_table("stats.psi", **query))
    row = selected.row(0, named=True)
    assert selected.height == 1
    assert row["feature"] == "x" and row["display_name"] == "局部空人口"
    assert row["data_source"] == "source_x" and row["dtype"] == "Float64"
    assert row["02"] is None
    assert selected.schema["02"] == pl.Float64
    assert row["01"] == pytest.approx(0.0, abs=1e-12)
    assert row["total"] == pytest.approx(0.0, abs=1e-12)
    statuses = _polars(report.get_table("calculation_status", features="x",
                                       segments=["seg2"], include_all=False))
    assert statuses.select("feature", "group", "metric", "status", "reason").to_dicts() == [{
        "feature": "x", "group": "02", "metric": "psi", "status": "not_computed",
        "reason": "no_effective_current_population",
    }]
    assert statuses["scope_id"].to_list() == selected["scope_id"].to_list()
    for options in ({"features": "z", "segments": ["seg2"], "include_all": False},
                    {"features": "x", "segments": ["seg1"], "include_all": False},
                    {"features": "x", "include_all": True}):
        assert len(report.get_table("calculation_status", **options)) == 0
    page = report.query_page("stats.psi", **query, limit=1)
    assert (page["total_rows"], page["returned_rows"], page["omitted_rows"], page["next_offset"]) == (1, 1, 0, None)
    assert_frame_equal(_polars(page["data"]), selected)
    reference = page["reference"]
    assert reference["report_id"] == report.report_id
    assert_frame_equal(_polars(report.get_table(reference["table"], **reference["query"])), selected)
    status_query: dict[str, Any] = {
        "features": "x", "segments": ["seg2"], "include_all": False,
        "filters": {"metric": "psi", "group": "02"},
    }
    text = report.to_ai_context(
        tables=["stats.psi", "calculation_status"],
        queries={"stats.psi": {**query, "limit": 1}, "calculation_status": status_query},
        max_chars=30000,
    )
    assert len(text) <= 30000
    assert '"02":null' in text.replace(" ", "")
    context = json.loads(text)
    assert context["omitted"] == []
    assert context["evidence"][0]["rows"] == selected.to_dicts()
    assert context["evidence"][1]["rows"] == statuses.to_dicts()
    for evidence in context["evidence"]:
        replay = report.get_table(evidence["reference"], **evidence["query"])
        assert _polars(replay).to_dicts() == evidence["rows"]
    empty = report.query_page("stats.psi", **query, offset=1, limit=1)
    assert empty["total_rows"] == 1 and empty["returned_rows"] == 0
    budgeted = report.to_ai_context(tables=["stats.psi"], queries={"stats.psi": query}, max_chars=12000)
    assert len(budgeted) <= 12000
    assert json.loads(budgeted)["evidence"][0]["rows"][0]["02"] is None
    description = report.describe()
    assert description["calculation_state"]["table"] == "calculation_status"
    assert description["tables"]["calculation_status"]["grain"] == "scope_id/target/feature/group/metric"
    assert "no_effective_current_population" in json.dumps(description)


def test_local_population_queries_and_snapshot_keep_identity_dtype_and_evidence(
    population_report: MarsProfileReport, tmp_path: Path,
) -> None:
    report = population_report
    path = tmp_path / "population.marsreport"
    report.save(path, overwrite=True)
    restored = load_report(path)
    assert restored.report_id == report.report_id
    for name in ("stats.psi", "calculation_status"):
        original = report.get_table(name)
        loaded = restored.get_table(name)
        assert type(original) is type(loaded)
        if isinstance(original, pd.DataFrame):
            pd.testing.assert_frame_equal(original, loaded)
        else:
            assert_frame_equal(original, loaded)
    for consumer in (report, restored):
        _assert_queries(consumer)


@pytest.mark.parametrize("columns", [None, ["feature", "01"], ["01"], ["02"], ["total"]])
def test_local_population_trend_caption_only_explains_displayed_cells(
    population_report: MarsProfileReport, tmp_path: Path, columns: list[str] | None,
) -> None:
    path = tmp_path / "trend.marsreport"
    population_report.save(path, overwrite=True)
    for consumer in (population_report, load_report(path)):
        # 恢复报告只使用已保存统计构造领域展示视图。
        view = scope_views(consumer, ["seg2"], False)[0][1]
        styler = view.show_trend("psi", features="x", columns=columns)
        visible = columns is None or "02" in columns
        assert ("not_computed/no_effective_current_population" in styler.caption) is visible
        context = json.loads(view.to_ai_context(
            tables=["stats.psi"], features="x", columns=columns, max_chars=30000,
        ))
        assert bool(context["description"]["parameters"]["diagnostics"]) is visible
        if visible:
            assert styler.data["02"].isna().all()
        assert "no_effective_current_population" not in view.show_trend("psi", features="z").caption
        assert "no_effective_current_population" not in view.show_trend("psi", features=[]).caption
        assert "no_effective_current_population" not in view.show_trend("psi", limit=0).caption
        other = scope_views(consumer, ["seg1"], False)[0][1]
        assert "no_effective_current_population" not in other.show_trend("psi", features="x").caption
    # 原始多客群入口还需沿用范围筛选，不能只验证拆分后的视图。
    direct = population_report.show_trend("psi", features="x", columns=columns,
                                          segments=["seg2"], include_all=False)
    assert ("no_effective_current_population" in direct.caption) is (columns is None or "02" in columns)


def _assert_html(path: Path) -> None:
    """检查实际 HTML 单元格的 null、零、状态及原始身份。"""
    parser = _HtmlTables()
    parser.feed(path.read_text(encoding="utf-8"))
    trends = [table for table in parser.tables if table and {"feature", "01", "02", "total"}.issubset(
        {cell["text"] for cell in table[0]}
    )]
    assert len(trends) == 1
    table = trends[0]
    headers = [cell["text"] for cell in table[0]]
    rows = [{column: cell for column, cell in zip(headers, row)} for row in table[1:]]
    x = next(row for row in rows if row["feature"]["text"] == "x")
    assert x["02"]["data-sort-type"] == "missing"
    assert x["02"]["data-sort-value"] == "" and x["02"]["text"] == "—"
    assert x["01"]["data-sort-type"] == x["total"]["data-sort-type"] == "number"
    assert float(x["01"]["data-sort-value"]) == pytest.approx(0.0, abs=1e-12)
    assert float(x["total"]["data-sort-value"]) == pytest.approx(0.0, abs=1e-12)
    states = [table for table in parser.tables if table and {"group", "metric", "status", "reason"}.issubset(
        {cell["text"] for cell in table[0]}
    )]
    assert len(states) == 1 and len(states[0]) == 2
    state = dict(zip([cell["text"] for cell in states[0][0]], [cell["text"] for cell in states[0][1]]))
    assert {key: state[key] for key in ("feature", "group", "metric", "status", "reason")} == {
        "feature": "x", "group": "02", "metric": "psi", "status": "not_computed",
        "reason": "no_effective_current_population",
    }


def _assert_excel(path: Path, *, generic: bool, scope_id: str) -> None:
    """检查工作簿真实单元格，缺失保持空白且局部状态可读。"""
    workbook = openpyxl.load_workbook(path, data_only=True)
    try:
        trend_name = next(name for name in workbook.sheetnames
                          if name.endswith("stats.psi")) if generic else "Trend_Psi"
        state_name = next(name for name in workbook.sheetnames
                          if name.endswith("calculation_status")) if generic else "CalculationStatus"
        trend = workbook[trend_name]
        headers = [cell.value for cell in trend[1]]
        records = [dict(zip(headers, row)) for row in trend.iter_rows(min_row=2, values_only=True)]
        x = next(row for row in records if row["feature"] == "x" and (
            row.get("scope_id", scope_id) == scope_id
        ))
        assert x["02"] is None
        assert x["01"] == pytest.approx(0.0, abs=1e-12)
        assert x["total"] == pytest.approx(0.0, abs=1e-12)
        state = workbook[state_name]
        headers = [cell.value for cell in state[1]]
        states = [dict(zip(headers, row)) for row in state.iter_rows(min_row=2, values_only=True)]
        assert len(states) == 1
        assert {key: states[0][key] for key in ("feature", "group", "metric", "status", "reason")} == {
            "feature": "x", "group": "02", "metric": "psi", "status": "not_computed",
            "reason": "no_effective_current_population",
        }
    finally:
        workbook.close()


def test_local_population_native_and_saved_html_excel_cells(
    population_report: MarsProfileReport, tmp_path: Path,
) -> None:
    report = population_report
    scope_id = resolve_scopes(report._scopes, ["seg2"], False)[0]
    report.save(tmp_path / "export.marsreport", overwrite=True)
    restored = load_report(tmp_path / "export.marsreport")
    for name, consumer in (("native", report), ("saved", restored)):
        html = tmp_path / f"{name}.html"
        consumer.write_html(str(html), segments=["seg2"], include_all=False)
        _assert_html(html)
        excel = tmp_path / f"{name}.xlsx"
        if isinstance(consumer, MarsProfileReport):
            consumer.write_excel(str(excel), segments=["seg2"], include_all=False)
        else:
            consumer.write_excel(str(excel))
        _assert_excel(excel, generic=name == "saved", scope_id=scope_id)


def test_population_snapshot_new_process_queries_and_exports_without_source_data(
    population_report: MarsProfileReport, tmp_path: Path,
) -> None:
    path = tmp_path / "standalone.marsreport"
    population_report.save(path, overwrite=True)
    script = tmp_path / "consume.py"
    script.write_text(
        "from pathlib import Path\n"
        "import json, sys\n"
        "from mars.reporting import load_report\n"
        "from mars.reporting._scope_view import scope_views\n"
        "report = load_report(sys.argv[1])\n"
        "query = dict(features='x', segments=['seg2'], include_all=False)\n"
        "page = report.query_page('stats.psi', **query, limit=1)\n"
        "context = json.loads(report.to_ai_context(tables=['stats.psi', 'calculation_status'], "
        "queries={'stats.psi': query, 'calculation_status': query}, max_chars=30000))\n"
        "row = context['evidence'][0]['rows'][0]\n"
        "assert row['02'] is None and abs(row['total']) < 1e-12\n"
        "assert context['evidence'][1]['rows'][0]['group'] == '02'\n"
        "assert context['evidence'][1]['rows'][0]['status'] == 'not_computed'\n"
        "reference = page['reference']\n"
        "replayed = report.get_table(reference['table'], **reference['query'])\n"
        "assert len(replayed) == 1\n"
        "view = scope_views(report, ['seg2'], False)[0][1]\n"
        "assert 'no_effective_current_population' in view.show_trend('psi', features='x', columns=['02']).caption\n"
        "assert 'no_effective_current_population' not in view.show_trend('psi', features='x', columns=['01']).caption\n"
        "report.write_html(str(Path(sys.argv[2]) / 'standalone.html'), segments=['seg2'], include_all=False)\n"
        "report.write_excel(str(Path(sys.argv[2]) / 'standalone.xlsx'))\n"
        "print(json.dumps({'report_id': report.report_id, 'row': row}, ensure_ascii=False))\n",
        encoding="utf-8",
    )
    env = {**os.environ, "PYTHONUTF8": "1", "POLARS_MAX_THREADS": "4", "MPLBACKEND": "Agg"}
    env["PYTHONPATH"] = str(Path(__file__).resolve().parents[1] / "src")
    completed = subprocess.run([sys.executable, str(script), str(path), str(tmp_path)],
                               cwd=tmp_path, env=env, capture_output=True, text=True,
                               encoding="utf-8", timeout=120, check=False)
    assert completed.returncode == 0, completed.stdout + completed.stderr
    payload = json.loads(completed.stdout.splitlines()[-1])
    assert payload["report_id"] == population_report.report_id and payload["row"]["02"] is None
    _assert_html(tmp_path / "standalone.html")
    scope_id = resolve_scopes(population_report._scopes, ["seg2"], False)[0]
    _assert_excel(tmp_path / "standalone.xlsx", generic=True, scope_id=scope_id)


def test_historical_snapshot_reads_saved_psi_without_recomputing(tmp_path: Path) -> None:
    # 旧快照只承担原统计读取合同，加载不能把保存的历史数值改写为新合同的 null。
    historical = MarsProfileReport(
        pl.DataFrame({"feature": ["x"], "dtype": ["Float64"]}),
        dq_tables={},
        stats_tables={"psi": pl.DataFrame({
            "feature": ["x"], "dtype": ["Float64"],
            "01": [0.0], "02": [11.538983029798928], "total": [0.0],
        })},
        report_meta={"psi_batch_size": 2, "psi_include_missing": False},
    )
    path = tmp_path / "historical.marsreport"
    historical.save(path, overwrite=True)
    restored = load_report(path)
    result = _polars(restored.get_table("stats.psi", features="x"))
    assert result["02"][0] == 11.538983029798928
    assert "calculation_status" not in restored.describe()["tables"]


@pytest.mark.parametrize("backend", ["polars", "pandas"])
def test_unsegmented_empty_reference_keeps_json_null_and_total_status(
    backend: str, tmp_path: Path,
) -> None:
    frame = pl.DataFrame({
        "g": ["01"] * 20 + ["02"] * 20,
        "x": pl.Series([None] * 20 + list(map(float, range(20))), dtype=pl.Float64),
        "z": list(map(float, range(20))) * 2,
    })
    report = MarsDataProfiler(psi_batch_size=2).generate_profile(
        frame.to_pandas() if backend == "pandas" else frame,
        features=["x", "z"], metrics=["psi"], group_col="g",
    )
    path = tmp_path / "no-reference.marsreport"
    report.save(path, overwrite=True)
    for consumer in (report, load_report(path)):
        result = _polars(consumer.get_table("stats.psi", features="x"))
        assert result.select("01", "02", "total").row(0) == (None, None, None)
        assert all(result.schema[column] == pl.Float64 for column in ("01", "02", "total"))
        status = _polars(consumer.get_table("calculation_status", features="x"))
        assert set(status["group"]) == {"01", "02", "Total"}
        assert set(status["reason"]) == {"no_effective_reference_population"}
        context = json.loads(consumer.to_ai_context(
            tables=["stats.psi"], features="x", max_chars=20000,
        ))
        row = context["evidence"][0]["rows"][0]
        assert all(row[column] is None for column in ("01", "02", "total"))
    assert "no_effective_reference_population" in report.show_trend("psi", features="x", columns=["total"]).caption
    assert "no_effective_reference_population" not in report.show_trend("psi", features="z").caption


@pytest.mark.parametrize("feature,segment,columns,limit,has_diagnostic", [
    ("x", "seg2", ["02"], 1, True),
    ("x", "seg2", ["01"], 1, False),
    ("x", "seg2", ["total"], 1, False),
    ("z", "seg2", ["02"], 1, False),
    ("x", "seg1", ["02"], 1, False),
    ("x", "seg2", ["02"], 0, False),
])
def test_ai_population_diagnostics_follow_actual_evidence_with_hidden_feature(
    population_report: MarsProfileReport, tmp_path: Path,
    feature: str, segment: str, columns: list[str], limit: int, has_diagnostic: bool,
) -> None:
    path = tmp_path / "projected.marsreport"
    population_report.save(path, overwrite=True)
    for consumer in (population_report, load_report(path)):
        text = consumer.to_ai_context(
            tables=["stats.psi"], features=feature, segments=[segment], include_all=False,
            columns=columns, limit=limit, max_chars=30000,
        )
        context = json.loads(text)
        assert bool(context["description"]["parameters"]["diagnostics"]) is has_diagnostic
        evidence = context["evidence"][0]
        assert all("feature" not in row for row in evidence["rows"])
        if has_diagnostic:
            assert evidence["identities"] == [{"feature": "x"}]
            diagnostic = context["description"]["parameters"]["diagnostics"][0]
            assert diagnostic["features"] == ["x"] and diagnostic["group"] == "02"
            assert diagnostic["scope_id"] == evidence["rows"][0]["scope_id"]
            assert diagnostic["reason"] == "no_effective_current_population"
        assert consumer.describe()["parameters"]["diagnostics"]


def test_ai_population_projection_preserves_existing_global_diagnostics(
    population_report: MarsProfileReport,
) -> None:
    legacy = {"component": "psi", "features": ["z"], "reason": "legacy_fit_failure"}
    population_report.report_meta["diagnostics"].append(legacy)
    context = json.loads(population_report.to_ai_context(
        tables=["stats.psi"], features="z", segments=["seg1"], include_all=False,
        columns=["01"], max_chars=30000,
    ))
    assert context["description"]["parameters"]["diagnostics"] == [legacy]


@pytest.mark.parametrize("backend", ["polars", "pandas"])
def test_actual_total_group_does_not_explain_valid_total_summary(
    backend: str, tmp_path: Path,
) -> None:
    frame = pl.DataFrame({
        "g": ["Total"] * 20 + ["01"] * 20,
        "x": pl.Series([None] * 20 + list(map(float, range(20))), dtype=pl.Float64),
    })
    benchmark = pl.DataFrame({"x": list(map(float, range(20)))})
    report = MarsDataProfiler().generate_profile(
        frame.to_pandas() if backend == "pandas" else frame,
        features=["x"], metrics=["psi"], group_col="g",
        benchmark_df=benchmark.to_pandas() if backend == "pandas" else benchmark,
    )
    path = tmp_path / "total-label.marsreport"
    report.save(path, overwrite=True)
    restored = load_report(path)
    for consumer in (report, restored):
        table = _polars(consumer.get_table("stats.psi", features="x"))
        assert table["Total"][0] is None and table["total"][0] == pytest.approx(0.0, abs=1e-12)
        for columns, expected in ((["total"], False), (["Total"], True), (["01"], False)):
            context = json.loads(consumer.to_ai_context(
                tables=["stats.psi"], features="x", columns=columns, max_chars=30000,
            ))
            assert bool(context["description"]["parameters"]["diagnostics"]) is expected
            if isinstance(consumer, MarsProfileReport):
                assert ("no_effective_current_population" in consumer.show_trend(
                    "psi", features="x", columns=columns,
                ).caption) is expected
