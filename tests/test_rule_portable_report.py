"""规则报告公共消费、桥接关联及跨进程恢复回归。"""

from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path
from typing import Any

import numpy as np
import openpyxl
import polars as pl
import pytest
from polars.testing import assert_frame_equal

from mars.agent import MarsAgentSession, MarsRiskAgent
from mars.reporting import Report, ReportSnapshot, load_report
from mars.reporting._query import ReportFrame
from mars.rule import MarsRule, MarsRuleMiningResult, MarsRuleMiningSpec, MarsRuleReport, mine_rules


def _result(*, strategy: str = "ranked", empty: bool = False) -> MarsRuleMiningResult:
    # 独立行与相同总体机制；辅助目标包含未表现值。
    def sample(start: int) -> pl.DataFrame:
        values = np.arange(start, start + 240)
        return pl.DataFrame(
            {
                "income": values % 120,
                "salary": values % 2,
                "unused": values,
                "bad": (values % 120 >= 85).astype(int),
                "aux": [None if v % 7 == 0 else int(v % 120 >= 80) for v in values],
                "month": ["2026-04-15" if v % 2 else "2026-05-15" for v in values],
                "segment": ["A" if v % 2 else "B" for v in values],
            }
        )

    return mine_rules(
        sample(0),
        target="bad",
        validation_df=sample(10000),
        aux_targets=["aux"],
        features=["income", "salary", "unused"],
        group_col="segment",
        time_col="month",
        time_grain="month",
        seed_rules=[]
        if empty
        else ["income >= 90 AND salary >= 0", "income >= 100", "income < 10"],
        generators=[],
        spec=MarsRuleMiningSpec(selection_strategy=strategy, iou_threshold=1.0),
    )


def _report(*, strategy: str = "ranked") -> MarsRuleReport:
    return _result(strategy=strategy).to_report(
        feature_metadata={
            "income": {"display_name": "月收入", "data_source": "application"},
            "salary": {"display_name": "月收入", "data_source": "bank"},
            "unused": {"data_source": "application"},
        },
        business_context={"dataset_id": "synthetic", "labels": {"bad": {"definition": "模拟违约"}}},
    )


@pytest.mark.parametrize("strategy", ["ranked", "cascade"])
def test_relation_preserves_statistical_rows_and_independent_conditions(
    tmp_path: Path, strategy: str
) -> None:
    original = _report(strategy=strategy)
    path = tmp_path / "rules.marsreport"
    original.save(path)
    restored = load_report(path)
    assert isinstance(original, Report) and isinstance(restored, ReportSnapshot)
    assert original.describe() == restored.describe()
    for report in (original, restored):
        row = report.get_table(
            "candidates", filters={"rule_id": MarsRule("income >= 90 AND salary >= 0").rule_id}
        ).row(0, named=True)
        rule_id = row["rule_id"]
        wanted = report.get_table("evaluation", filters={"rule_id": rule_id})
        assert_frame_equal(report.get_table("evaluation", features="salary"), wanted)
        # 特征和来源允许命中同一规则的不同成员；并集绝不使指标重复。
        assert_frame_equal(
            report.get_table("evaluation", features="income", sources="bank"), wanted
        )
        assert_frame_equal(
            report.get_table("evaluation", features=["income", "salary"]),
            report.get_table("evaluation", features="income"),
        )
        assert report.get_table("evaluation", features="unused").is_empty()
        assert report.get_feature("unused")["evidence_status"] == "no_evidence"
        assert {v["feature"] for v in report.search_features("月收入")} == {"income", "salary"}
        assert report.get_feature("income", limit=1)["omitted_rows"]["evaluation"] > 0
        with pytest.raises(ValueError, match="Unknown feature"):
            report.get_feature("missing")
        for name in ("summary",):
            with pytest.raises(ValueError):
                report.get_table(name, features="income")
        with pytest.raises(ValueError, match="Unknown feature sources"):
            report.get_table("evaluation", sources="manual")
        with pytest.raises(ValueError):
            report.get_table("evaluation", filters={"rule_id": {"op": "eval", "value": "bad"}})
        query = {
            "features": "salary",
            "filters": {"dataset": "validation", "target": "aux", "group": "hit"},
            "columns": ["rule_id", "dataset", "target", "slice", "sample_count"],
            "limit": 1,
        }
        page = report.query_page("slices", **query)
        assert page["total_rows"] > 1
        assert page["next_offset"] == 1
        assert_frame_equal(
            page["data"], report.get_table(page["reference"]["table"], **page["reference"]["query"])
        )
        assert all(v == "aux" for v in page["data"]["target"])
    assert original.to_ai_context(
        queries={"evaluation": {"features": "salary", "limit": 2}}, max_chars=16000
    ) == restored.to_ai_context(
        queries={"evaluation": {"features": "salary", "limit": 2}}, max_chars=16000
    )
    assert original.get_table("rule_explanations")["dataset"].unique().to_list() == ["validation"]
    if strategy == "cascade":
        assert "generation_round" in restored.get_table("candidates").columns
        assert restored.get_table("rules")["selection_round"].null_count() == 0


def test_queries_do_not_recompute_and_budget_is_final_json(monkeypatch: pytest.MonkeyPatch) -> None:
    report = _report()

    def forbidden(*args: Any, **kwargs: Any) -> None:
        raise AssertionError("查询不得重新挖掘或分析")

    monkeypatch.setattr("mars.rule.workflow.mine_rules", forbidden)
    monkeypatch.setattr("mars.rule.workflow.analyze_rule_set", forbidden)
    report.get_feature("salary")
    text = report.to_ai_context(
        queries={"evaluation": {"features": "salary", "limit": 100}}, max_chars=7000
    )
    payload = json.loads(text)
    assert len(text) <= 7000 and payload["omitted"]
    assert payload["evidence"][0]["report_id"] == report.report_id
    assert "candidates" not in [
        e["reference"] for e in json.loads(report.to_ai_context())["evidence"]
    ]
    defaults = json.loads(report.to_ai_context())["evidence"]
    assert [e["reference"] for e in defaults] == ["summary", "rule_explanations"]
    assert defaults[1]["returned_rows"] <= 3
    assert all("expression" not in row for row in defaults[1]["rows"])
    projected = json.loads(
        report.to_ai_context(
            queries={"evaluation": {"features": "salary", "columns": ["sample_count"], "limit": 1}}
        )
    )
    assert set(projected["description"]["feature_metadata"]) == {"income", "salary"}
    with pytest.raises(ValueError, match="max_chars"):
        report.to_ai_context(max_chars=512)
    for invalid in (None, -1, True, "3"):
        with pytest.raises(ValueError, match="limit"):
            report.to_ai_context(limit=invalid)
        with pytest.raises(ValueError, match="limit"):
            report.to_ai_context(queries={"candidates": {"limit": invalid}})


@pytest.mark.parametrize("empty", [False, True])
def test_no_rules_keeps_audit_and_state(tmp_path: Path, empty: bool) -> None:
    if empty:
        with pytest.warns(UserWarning):
            result = _result(empty=True)
    else:
        data = pl.DataFrame({"income": list(range(100)), "bad": [int(i >= 80) for i in range(100)]})
        validation = data.with_columns((1 - pl.col("bad")).alias("bad"))
        with pytest.warns(UserWarning):
            result = mine_rules(
                data,
                target="bad",
                validation_df=validation,
                seed_rules=["income >= 80"],
                generators=[],
            )
    report = result.to_report()
    report.save(tmp_path / "empty.marsreport")
    restored = load_report(tmp_path / "empty.marsreport")
    assert restored.get_table("summary")["status"][0] == "no_rules"
    assert restored.get_table("rules").is_empty()
    assert restored.describe()["parameters"]["analysis_states"]["interactions"] == "not_computed"
    assert restored.get_table("candidates").height == (0 if empty else 1)
    if not empty:
        assert restored.get_table("evaluation", features="income").height > 0
        assert restored.get_table("candidates")["rejection_stage"][0] == "validation_filter"


def test_explicit_analysis_grain_scope_and_cumulative_membership(tmp_path: Path) -> None:
    from mars.rule import MarsRule, MarsRuleSet
    from mars.rule.analysis import analyze_rule_set

    data = pl.DataFrame(
        {
            "x": list(range(100)),
            "z": [i % 2 for i in range(100)],
            "bad": [int(i >= 70) for i in range(100)],
        }
    )
    rules = MarsRuleSet([MarsRule("x >= 80"), MarsRule("z == 1")])
    analysis = analyze_rule_set(rules, data, target="bad", bootstrap_repeats=10)
    definitions = pl.DataFrame(
        [{"rule_id": r.rule_id, "expression": r.expression} for r in rules.rules]
    )
    report = MarsRuleReport(
        detail_tables={
            "rules": definitions,
            "interactions": analysis.interaction_table,
            "cumulative": analysis.cumulative_table,
            "bootstrap": analysis.bootstrap_table,
        },
        metadata={"advanced_analysis": dict(analysis.metadata)},
    )
    report.save(tmp_path / "advanced.marsreport")
    for current in (report, load_report(tmp_path / "advanced.marsreport")):
        assert current.get_table("interactions", features=["x", "z"]).height == 1
        assert current.get_table("interactions", features="z").height == 1
        # 累计第 2 步依旧涉及第一条规则，不能只匹配 added_rule_id。
        assert current.get_table("cumulative", features="x").height == 2
        assert current.get_table("cumulative", features="z").height == 1
        assert current.describe()["parameters"]["advanced_analysis"]["bootstrap_repeats"] == 10
        assert (
            current.describe()["tables"]["interactions"]["fields"]["rule_b"]["unit"] == "identifier"
        )
        assert (
            current.describe()["tables"]["cumulative"]["fields"]["cumulative_sample_count"]["unit"]
            == "count"
        )


def test_new_process_exports_nonfinite_and_agent_registration(tmp_path: Path) -> None:
    report = _report()
    # 仅测试统计字段的非有限保存；不保存原宽表。
    details = {
        **report.detail_tables,
        "numeric_evidence": pl.DataFrame({"metric": [0.0, None, float("nan"), float("inf")]}),
    }
    report = MarsRuleReport(
        report.summary_table,
        details,
        report.metadata,
        feature_metadata=report.feature_metadata,
        business_context={**report.business_context, "unsafe": "<script>alert(1)</script>"},
    )
    path = tmp_path / "report.marsreport"
    report.save(path)
    report.write_html(tmp_path / "original.html")
    report.write_excel(tmp_path / "original.xlsx")
    original_html = (tmp_path / "original.html").read_text(encoding="utf-8")
    assert "&lt;script&gt;alert(1)" in original_html and "<script>alert(1)" not in original_html
    original_workbook = openpyxl.load_workbook(tmp_path / "original.xlsx", read_only=True)
    original_sheet = original_workbook["evaluation"]
    rows = list(original_sheet.values)
    assert rows[0] == tuple(report.get_table("evaluation").columns)
    # Excel 存储数值保留约 15 位有效数字；身份/维度精确比较，指标容许浮点尾差。
    for actual, expected in zip(rows[1], report.get_table("evaluation").row(0)):
        assert (
            actual == pytest.approx(expected, rel=1e-14)
            if isinstance(expected, float)
            else actual == expected
        )
    assert original_sheet.max_row == report.get_table("evaluation").height + 1
    original_workbook.close()
    identifier = report.report_id
    del report, details
    code = """import json, sys
from pathlib import Path
from mars.reporting import load_report
r = load_report(sys.argv[1])
assert r.report_id == sys.argv[2]
assert r.get_table('evaluation', features='income', sources='bank').height > 0
assert r.query_page('slices', features='salary', limit=1)['next_offset'] == 1
assert len(r.to_ai_context(max_chars=8000)) <= 8000
r.write_html(str(Path(sys.argv[1]).with_suffix('.html')))
r.write_excel(str(Path(sys.argv[1]).with_suffix('.xlsx')))
"""
    subprocess.run(
        [sys.executable, "-c", code, str(path), identifier],
        check=True,
        capture_output=True,
        text=True,
    )
    restored = load_report(path)
    assert restored.get_table("numeric_evidence")["metric"].is_nan().sum() == 1
    assert restored.get_table("numeric_evidence")["metric"][0] == 0.0
    document = path.with_suffix(".html").read_text(encoding="utf-8")
    assert "&lt;script&gt;alert(1)" in document and "<script>alert(1)" not in document
    workbook = openpyxl.load_workbook(path.with_suffix(".xlsx"), read_only=True)
    sheet = next(s for s in workbook if s.title.endswith("_evaluation"))
    assert "sample_count" in next(sheet.values)
    assert sheet.max_row == restored.get_table("evaluation").height + 1
    workbook.close()
    session = MarsAgentSession()
    handle = session.register_report(restored)
    answer = MarsRiskAgent().execute_tool(
        "get_report_table",
        {
            "report_id": handle,
            "table": "evaluation",
            "features": ["salary"],
            "sources": ["application"],
            "limit": 2,
        },
        session=session,
    )
    assert answer.success and answer.data["rows"]
    assert answer.data["persistent_report_id"] == identifier


def test_benchmark_and_relation_validation(tmp_path: Path) -> None:
    report = MarsRuleReport.from_benchmark({"seconds": 1.25, "peak_memory_mb": 20.0})
    report.save(tmp_path / "bench.marsreport")
    restored = load_report(tmp_path / "bench.marsreport")
    description = restored.describe()
    assert description["report_type"] == "rule_benchmark"
    assert description["tables"]["benchmark"]["fields"]["seconds"]["unit"] == "seconds"
    assert description["tables"]["benchmark"]["fields"]["peak_memory_mb"]["unit"] == "MB"
    assert "rules" not in description["tables"]
    assert "qualification" not in description["parameters"]
    with pytest.raises(ValueError):
        restored.get_table("benchmark", features="income")
    description = _report().describe()
    description["tables"]["evaluation"]["feature_relation"]["table"] = "missing"
    with pytest.raises(ValueError, match="feature_relation"):
        ReportSnapshot(_report()._query_tables(), description)


def test_large_audit_is_paged_before_json(monkeypatch: pytest.MonkeyPatch) -> None:
    # 不执行挖掘，仅构造较大的已计算审计；强制 JSON 边界看不到全表。
    definitions = [MarsRule(f"income > {i} AND salary >= 0") for i in range(1500)]
    candidates = pl.DataFrame(
        [
            {"rule_id": r.rule_id, "expression": r.expression, "status": "rejected"}
            for r in definitions
        ]
    )
    report = MarsRuleReport(
        pl.DataFrame({"status": ["no_rules"], "candidate_count": [1500]}),
        {"candidates": candidates},
    )
    from mars.reporting._serialization import table_rows

    def bounded_rows(frame: ReportFrame) -> list[dict[str, Any]]:
        assert len(frame) <= 3, "JSON 编码前必须先分页"
        return table_rows(frame)

    monkeypatch.setattr("mars.reporting._query.table_rows", bounded_rows)
    context = json.loads(
        report.to_ai_context(
            queries={
                "candidates": {
                    "features": ["income", "salary"],
                    "columns": ["rule_id", "status"],
                    "limit": 3,
                }
            },
            max_chars=8000,
        )
    )
    evidence = context["evidence"][0]
    assert evidence["total_rows"] == 1500
    assert evidence["returned_rows"] == 3 and evidence["next_offset"] == 3
    assert context["omitted"][0]["rows"] == 1497
    assert report.get_table("candidates", features=["income", "salary"]).height == 1500


def test_relation_query_retains_pandas_backend(tmp_path: Path) -> None:
    original = _report()
    tables = {name: table.to_pandas() for name, table in original._query_tables().items()}
    description = original.describe()
    for name, frame in tables.items():
        for column, dtype in frame.dtypes.items():
            description["tables"][name]["fields"][column]["dtype"] = str(dtype)
    report = ReportSnapshot(tables, description)
    report.save(tmp_path / "pandas.marsreport")
    restored = load_report(tmp_path / "pandas.marsreport")
    first = report.get_table("evaluation", features="salary", sources="application")
    second = restored.get_table("evaluation", features="income", sources="bank")
    assert first.equals(second)
    assert restored.get_feature("unused")["evidence_status"] == "no_evidence"
