"""既有 snapshot 静态 Excel 的当前数值、持久化与独立进程交付。"""

from __future__ import annotations

import json
import math
import os
import subprocess
import sys
import zipfile
from copy import deepcopy
from datetime import date, datetime, timezone
from numbers import Real
from pathlib import Path
from types import SimpleNamespace
from typing import Any
from unittest.mock import MagicMock
from xml.etree import ElementTree

import pandas as pd
import polars as pl
import pytest
from openpyxl import load_workbook

import mars
import mars.analysis.evaluator as evaluator_module
from mars.analysis import profile_risk
from mars.compute import to_pandas_frame
from mars.feature.binning import MarsNativeBinner
from mars.reporting import MarsBinningReport, ReportSnapshot, snapshot_report
from mars.reporting._binning_excel import _BinningExcelWriter


def _assert_excel_records(records: list[list[Any]], expected: pd.DataFrame) -> None:
    """逐单元格核对 Excel 边界，缺失为空、文字保持原文、指标仍为数值。"""
    assert list(records[0]) == list(expected.columns)
    assert len(records) == len(expected) + 1
    for column_index, column in enumerate(expected.columns):
        for actual_row, expected_value in zip(records[1:], expected[column]):
            actual = actual_row[column_index]
            if pd.isna(expected_value):
                assert actual is None
            elif isinstance(expected_value, str):
                assert isinstance(actual, str) and actual == expected_value
            elif isinstance(expected_value, Real) and not isinstance(expected_value, bool):
                assert isinstance(actual, Real) and not isinstance(actual, bool)
                assert actual == pytest.approx(expected_value, rel=1e-12, abs=1e-12)
            else:
                assert actual == expected_value


@pytest.mark.parametrize("native", [True, False])
def test_excel_preserves_business_strings_as_literal_cells(tmp_path: Path, native: bool) -> None:
    """真实工作簿的业务文字不成为公式，也不丢失前导零或日期形文字。"""
    texts = ["=1+1", "普通中文", '带"引号', "00123", "+1+1", "-1+1", "@SUM(A1)", "2026-01-02", "'literal", "1E10"]
    frame = pl.DataFrame({**{f"x{i}": list(range(12)) for i in range(len(texts))}, "bad": [0, 1] * 6})
    report = profile_risk(
        frame, target="bad", features=[f"x{i}" for i in range(len(texts))], n_bins=2,
        feature_metadata={f"x{i}": {"display_name": text} for i, text in enumerate(texts)},
    ).report
    before = snapshot_report(report)
    path = tmp_path / "literal.xlsx"
    if native:
        report.write_excel(str(path), engine="openpyxl")
    else:
        before.write_excel(str(path))
    for data_only in (False, True):
        workbook = load_workbook(path, data_only=data_only)
        try:
            sheet = workbook["FeatureMetadata"]
            assert [sheet.cell(i + 2, 2).value for i in range(len(texts))] == texts
            assert all(sheet.cell(i + 2, 2).data_type == "s" for i in range(len(texts)))
        finally:
            workbook.close()
    for name in before.describe()["tables"]:
        pd.testing.assert_frame_equal(to_polars(report.get_table(name)).to_pandas(), to_polars(before.get_table(name)).to_pandas())
    assert report.feature_metadata == before.feature_metadata


def test_native_excel_formats_actual_export_units_and_opens_summary(tmp_path: Path) -> None:
    report = profile_risk(
        pl.DataFrame({"x": list(range(12)), "bad": [0, 1] * 6}),
        target="bad", features=["x"], n_bins=2, ks_method="raw",
    ).report
    path = tmp_path / "units.xlsx"
    report.write_excel(str(path), engine="openpyxl")
    workbook = load_workbook(path)
    try:
        assert workbook.active.title == "最终汇总"
        sheet = workbook["最终汇总"]
        headers = [cell.value for cell in sheet[1]]
        expected = {"ks": "#,##0.00", "auc": "#,##0.00", "iv": "#,##0.00", "missing_min": "0.00%", "missing_max": "0.00%"}
        for name, number_format in expected.items():
            assert sheet.cell(2, headers.index(name) + 1).number_format == number_format
        trend = workbook["最终KS趋势"]
        for column in range(5, trend.max_column + 1):
            assert trend.cell(2, column).number_format == "#,##0.00"
        assert sheet.freeze_panes == "A2" and sheet.auto_filter.ref is None
    finally:
        workbook.close()


def test_xlwings_writes_literal_text_and_dates_without_local_timezone_shift(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(sys, "platform", "win32")
    worksheet = MagicMock()
    region = worksheet.range.return_value
    rows = [["=1+1", pd.Timestamp("2026-10-15 12:30"), date(2026, 10, 15), 1.25, None, "'literal"]]
    _BinningExcelWriter._write_xlwings_detail_rows(
        worksheet, rows, start_row=2, first_col=1, last_col=6,
    )
    actual = region.api.Value[0]
    assert actual == (
        "=1+1", datetime(2026, 10, 15, 12, 30, tzinfo=timezone.utc),
        datetime(2026, 10, 15, tzinfo=timezone.utc), 1.25, None, "''literal",
    )
    assert region.number_format == "@"
    assert isinstance(rows[0][1], pd.Timestamp) and rows[0][1].tzinfo is None


def test_native_excel_formats_counts_weights_amounts_and_preserves_template_formula(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    report = profile_risk(
        pl.DataFrame({"=identity": list(range(12)), "bad": [0, 1] * 6, "w": [1.25] * 12}),
        target="bad", features=["=identity"], n_bins=2, weights_col="w",
        feature_metadata={"=identity": {"data_source": "=source"}},
    ).report
    # 使用真实报告目录给附表增加明确指标，不改变生产导出目录或算法。
    summary = to_pandas_frame(report.summary_table)
    summary["count"] = pd.Series([1000.0])
    summary["bad"] = pd.Series([1.25])
    summary["tot_amt"] = pd.Series([12345.5])
    summary["bad_rate"] = pd.Series([1.0])
    report = MarsBinningReport(
        summary, report.trend_tables, report.detail_table,
        group_col=report.group_col, detail_group_col=report.detail_group_col,
        report_meta=report.report_meta, feature_metadata=report.feature_metadata,
    )
    assert report.describe()["tables"]["summary"]["fields"]["bad"]["unit"] == "weight_sum"
    template = _BinningExcelWriter._resolve_excel_template_path("mars_bin_report_linux.xlsx")
    workbook = load_workbook(template)
    workbook.create_sheet("TemplateFormula")["A1"] = "=1+1"
    custom_template = tmp_path / "template.xlsx"
    workbook.save(custom_template)
    workbook.close()
    monkeypatch.setattr(_BinningExcelWriter, "_resolve_excel_template_path", classmethod(lambda cls, file_name: custom_template))
    path = tmp_path / "formats.xlsx"
    report.write_excel(str(path), engine="openpyxl")
    workbook = load_workbook(path, data_only=False)
    try:
        assert workbook["TemplateFormula"]["A1"].data_type == "f"
        assert workbook["TemplateFormula"]["A1"].value == "=1+1"
        sheet = workbook["最终汇总"]
        headers = [cell.value for cell in sheet[1]]
        expected = {"count": (1000, "#,##0"), "bad": (1.25, "#,##0.00"), "tot_amt": (12345.5, "#,##0.00"), "bad_rate": (1, "0.00%")}
        for name, (value, number_format) in expected.items():
            cell = sheet.cell(2, headers.index(name) + 1)
            assert cell.value == value and cell.number_format == number_format and cell.data_type == "n"
        assert sheet["A2"].value == "=identity" and sheet["A2"].data_type == "s"
        assert sheet["D2"].value == "=source" and sheet["D2"].data_type == "s"
    finally:
        workbook.close()


def _assert_current_excel(path: Path, report: ReportSnapshot) -> None:
    """以不重算、不刷新方式逐表核对当前公共展示数据。"""
    workbook = load_workbook(path, data_only=True)
    try:
        for index, name in enumerate(report.describe()["tables"]):
            frame = report.show_table(name)
            sheet = workbook[f"{index:03d}_{name}"[:31]]
            records = list(sheet.values)
            assert records[0] == tuple(frame.columns)
            assert len(records) == len(frame) + 1
            for actual, expected in zip(records[1:], frame.itertuples(index=False, name=None)):
                for found, original in zip(actual, expected):
                    if pd.isna(original):
                        assert found is None
                    elif isinstance(original, float):
                        assert found == pytest.approx(original)
                    else:
                        assert found == original
    finally:
        workbook.close()
    # 静态交付所有当前值均是字面单元格，不依赖公式或透视缓存刷新。
    formulas = load_workbook(path, data_only=False)
    try:
        assert not any(cell.data_type == "f" for sheet in formulas for row in sheet for cell in row)
    finally:
        formulas.close()


@pytest.mark.parametrize("backend", ["pandas", "polars"])
def test_snapshot_excel_delivers_current_risk_values_after_new_process_load(
    backend: str, tmp_path: Path
) -> None:
    """静态文件和新进程文件都可直接读取本次风险值与中文类别。"""
    frame = pl.DataFrame(
        {
            "income": [1.0, 2.0, float("nan"), 4.0, 5.0, 6.0, 7.0, 8.0],
            "category": ["甲", "甲", None, "乙", "乙", "乙", "甲", "乙"],
            "bad": [0, 0, 0, None, 1, 1, None, 1],
            "weight": [1.0, 2.0, 1.0, 0.0, 2.0, 1.0, 1.0, 3.0],
        }
    )
    report = profile_risk(
        frame.to_pandas() if backend == "pandas" else frame,
        features=["income", "category"], target="bad", weights_col="weight", n_bins=2,
        feature_metadata={"income": {"display_name": "月收入"}},
    ).report
    snapshot = snapshot_report(report)
    summary = snapshot.get_table("summary")
    rows: list[dict[str, Any]] = (
        summary.to_dict("records") if isinstance(summary, pd.DataFrame) else summary.to_dicts()
    )
    income = next(row for row in rows if row["feature"] == "income")
    assert math.isfinite(income["ks"]) and income["ks"] > 0
    assert math.isfinite(income["iv"]) and income["iv"] > 0
    output = tmp_path / "static.xlsx"
    snapshot.write_excel(str(output))
    _assert_current_excel(output, snapshot)
    archive = tmp_path / "risk.marsreport"
    report.save(archive)
    restored_output = tmp_path / "restored.xlsx"
    script = (
        "import sys; from mars.reporting import load_report; "
        "report=load_report(sys.argv[1]); report.write_excel(sys.argv[2]); "
        "print(report.report_id)"
    )
    environment = dict(os.environ)
    environment["PYTHONPATH"] = str(Path(mars.__file__).resolve().parents[1])
    completed = subprocess.run(
        [sys.executable, "-c", script, str(archive), str(restored_output)],
        cwd=tmp_path, env=environment, capture_output=True, text=True, encoding="utf-8", check=True,
    )
    assert completed.stdout.strip() == report.report_id
    _assert_current_excel(restored_output, snapshot)
    # 既有透视选项仍可写出原结构；这里只查 XML，不将它称为原生 Excel 刷新。
    pivot_output = tmp_path / "pivot.xlsx"
    report.write_excel(str(pivot_output), engine="openpyxl")
    with zipfile.ZipFile(pivot_output) as workbook_archive:
        pivots = [name for name in workbook_archive.namelist() if name.startswith("xl/pivotTables/") and name.endswith(".xml")]
        caches = [name for name in workbook_archive.namelist() if name.startswith("xl/pivotCache/pivotCacheDefinition") and name.endswith(".xml")]
        assert pivots and caches
        assert all(ElementTree.fromstring(workbook_archive.read(name)).attrib.get("refreshOnLoad") == "1" for name in caches)


@pytest.mark.parametrize("backend", ["pandas", "polars"])
def test_raw_and_binned_reuse_template_without_refit_and_keep_distinct_ks(
    backend: str, tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    """同一固定规则的双模式报告保留各自 KS，原生 Excel 一致显示分箱值。"""
    frame = pl.DataFrame({
        "x": list(range(12)) * 4,
        "category": ["甲", "乙"] * 24,
        "bad": [0, 1] * 6 + [0, 0, 1, 1] * 3 + [0, 0, 0, 1] * 3 + [None] * 12,
        "second": [0, 0, 1, 1] * 6 + [0, 1] * 6 + [None] * 12,
        "group": ["observed"] * 12 + ["validation"] * 12 + [None] * 12 + ["undue"] * 12,
        "weight": [1.0, 2.0, 1.0, 3.0] * 12,
        "amount": list(range(1, 49)),
        "dt": ["2024-01-01"] * 24 + ["2024-02-01"] * 24,
    })
    input_frame = frame.to_pandas() if backend == "pandas" else frame
    options: dict[str, Any] = {
        "target": ["bad", "second"], "features": ["x", "category"], "group_col": "group",
        "n_bins": 2, "method": "quantile", "ordered_metric_sort_by": "bin_index",
        "weights_col": "weight", "amount_col": "amount",
        "time_col": "dt",
        "benchmark_df": input_frame,
        "feature_metadata": {"x": {"display_name": "连续字段", "data_source": "申请"}},
        "business_context": {"labels": {"bad": {"definition": "已表现坏样本"}}},
    }
    binned_run = profile_risk(input_frame, **options, ks_method="binned")

    def forbid_refit(*args: Any, **kwargs: Any) -> None:
        raise AssertionError("Excel 消费已经计算的报告，不得重新 fit。")

    # profile_risk 无公开 binner 参数；复用其已经拟合的规则并阻断所有再次拟合。
    monkeypatch.setattr(evaluator_module, "build_binner", lambda **kwargs: binned_run.binner)
    monkeypatch.setattr(MarsNativeBinner, "fit", forbid_refit)
    raw_run = profile_risk(input_frame, **options, ks_method="raw")
    assert raw_run.binner is binned_run.binner
    reports = {"binned": binned_run.report, "raw": raw_run.report}
    binned_summary = to_polars(reports["binned"].summary_table)
    raw_summary = to_polars(reports["raw"].summary_table)
    total_filter = (pl.col("feature") == "x") & (pl.col("target") == "bad")
    assert raw_summary.filter(total_filter)["ks"][0] != pytest.approx(
        binned_summary.filter(total_filter)["ks"][0],
    )
    # 数值字段 raw、分类字段回退、所有非 KS 指标仍由源报告决定。
    pd.testing.assert_frame_equal(
        raw_summary.drop("ks").to_pandas(), binned_summary.drop("ks").to_pandas(),
    )
    assert reports["raw"].report_meta["ks_source_by_feature"] == {"x": "raw", "category": "binned"}
    pd.testing.assert_frame_equal(
        raw_summary.filter(pl.col("feature") == "category").to_pandas(),
        binned_summary.filter(pl.col("feature") == "category").to_pandas(),
    )

    original: dict[str, Any] = {}
    for method, report in reports.items():
        original[method] = {
            "describe": deepcopy(report.describe()),
            "meta": deepcopy(report.report_meta),
            "tables": {name: to_polars(report.get_table(name)).clone() for name in report.describe()["tables"]},
            "objects": (report.summary_table, report.detail_table, report.trend_tables, report.report_meta),
            "context": report.to_ai_context(tables=["summary"], features="x", limit=2, max_chars=24000),
        }
        # 先 HTML 后 Excel 与先 Excel 后查询都保留同一 Python 口径。
        if method == "raw":
            report.write_html(str(tmp_path / "before_excel.html"), max_plots=0)

    template = load_workbook(_BinningExcelWriter._resolve_excel_template_path("mars_bin_report_linux.xlsx"))
    try:
        template_styles: dict[str, dict[str, Any]] = {
            name: {coordinate: deepcopy(template[name][coordinate]._style) for coordinate in coordinates}
            for name, coordinates in {
                "分组明细": ["A1", "Q1", "V1", "A2", "Q2", "V2"],
                "KS": ["A4", "B4"], "综合指标": ["A5", "D5", "F5"],
            }.items()
        }
        pivot_locations = {
            name: [pivot.location.ref for pivot in template[name]._pivots]
            for name in ["KS", "综合指标"]
        }
        native_style = deepcopy(next(iter(template["分组明细"].tables.values())).tableStyleInfo)
    finally:
        template.close()

    native_frames: dict[str, dict[str, pd.DataFrame]] = {}
    for method, report in reports.items():
        output = tmp_path / f"{method}.xlsx"
        report.write_excel(str(output), engine="openpyxl")
        workbook = load_workbook(output, data_only=True)
        try:
            assert {"分组透视", "综合指标", "KS", "分组明细", "最终汇总", "最终KS趋势"} <= set(workbook.sheetnames)
            detail = workbook["分组明细"]
            native = to_polars(report.detail_table)
            assert detail.max_row == native.height + 1
            assert next(iter(detail.tables.values())).ref == f"A1:V{native.height + 1}"
            assert detail.freeze_panes == "A2"
            headers = list(next(detail.values))
            assert headers[:6] == ["y", "feature", "trend", "mars_group", "bin_index", "bin_label"]
            assert len(headers) == 22 and headers[16] == "ks_bin" and headers[21] == "ks"
            native_table = next(iter(detail.tables.values()))
            assert [column.name for column in native_table.tableColumns] == headers
            assert native_table.tableStyleInfo == native_style
            for name, styles in template_styles.items():
                for coordinate, style in styles.items():
                    assert workbook[name][coordinate]._style == style
            for name, locations in pivot_locations.items():
                assert [pivot.location.ref for pivot in workbook[name]._pivots] == locations
            assert "display_name" not in headers and "dtype" not in headers and "data_source" not in headers
            actual = pd.DataFrame(list(detail.values)[1:], columns=headers)
            expected_detail = native.to_pandas()
            expected_detail["ks"] = expected_detail.groupby(
                ["feature", "y", report.detail_group_col], dropna=False, observed=True,
            )["ks_bin"].transform("max")
            pd.testing.assert_frame_equal(
                actual, expected_detail.loc[:, headers], check_dtype=False, check_exact=False,
            )
            assert actual.loc[actual["mars_group"] == "undue", "ks"].isna().all()
            if method == "raw":
                assert any("源报告" in name and "诊断" in name for name in workbook.sheetnames)
            pivot = workbook["KS"]._pivots[0]
            assert pivot.cache.cacheFields[pivot.dataFields[0].fld].name == "ks"
            assert pivot.cache.cacheSource.worksheetSource.name == "detail"
            assert pivot.cache.refreshOnLoad and pivot.cache.recordCount == 0
            assert pivot.cache.records is None
            assert "分箱" in pivot.dataFields[0].name
            combined_ks = [
                field for combined in workbook["综合指标"]._pivots for field in combined.dataFields
                if combined.cache.cacheFields[field.fld].name == "ks"
            ]
            assert combined_ks and all(field.subtotal == "average" for field in combined_ks)
            assert all("分箱" in field.name and "avg" in field.name for field in combined_ks)
            final = workbook["最终汇总"]
            final_headers = list(next(final.values))
            assert final_headers[:4] == ["feature", "display_name", "dtype", "data_source"]
            x_row = next(row for row in range(2, final.max_row + 1) if final.cell(row, 1).value == "x")
            assert final.cell(x_row, 2).value == "连续字段"
            assert final.cell(x_row, 4).value == "申请"
            final_table = next(iter(final.tables.values()))
            assert final.freeze_panes == "A2" and final_table.autoFilter.ref == final_table.ref
            # 工作表 AutoFilter 与 ListObject filter 重叠会令真实 Excel 拒绝打开。
            assert final.auto_filter.ref is None
            final_frame = pd.DataFrame(list(final.values)[1:], columns=final_headers)
            pd.testing.assert_frame_equal(
                final_frame, binned_summary.to_pandas(), check_dtype=False, check_exact=False,
            )
            trend = workbook["最终KS趋势"]
            trend_rows = list(trend.values)
            trend_frame = pd.DataFrame(trend_rows[1:], columns=trend_rows[0])
            expected_trend = to_polars(report.trend_tables["ks"]).to_pandas()
            grouped_ks = expected_detail.drop_duplicates(["feature", "y", report.detail_group_col])
            first_target = grouped_ks.loc[grouped_ks["y"] == "bad"]
            for group in ["null", "observed", "undue", "validation", "Total"]:
                group_mask = (
                    first_target[report.detail_group_col].isna()
                    if group == "null" else first_target[report.detail_group_col].eq(group)
                )
                matched = first_target.loc[group_mask]
                expected_trend[group] = expected_trend[["feature"]].merge(
                    matched[["feature", "ks"]], how="left", on="feature", sort=False,
                )["ks"].to_numpy()
            _assert_excel_records(trend_rows, expected_trend)
            assert list(trend_frame.columns) == list(to_polars(report.trend_tables["ks"]).columns)
            assert "target" not in trend_frame  # 保留既有首目标趋势契约。
            assert trend_frame["undue"].isna().all()
            semantics = " ".join(str(value) for row in workbook["KS口径"].values for value in row)
            assert "source_report_ks_method" in semantics and method in semantics
            assert "exported_ks_method" in semantics and "binned" in semantics
            assert "exported_ks_source" in semantics and "ks_bin" in semantics
            assert "0–100" in semantics or "0-100" in semantics
            native_frames[method] = {"detail": actual, "summary": final_frame, "trend": trend_frame}
        finally:
            workbook.close()
        _assert_native_excel_relationships(output, native.height)
        _assert_report_unchanged(report, original[method])

    for name in ["detail", "summary", "trend"]:
        pd.testing.assert_frame_equal(native_frames["raw"][name], native_frames["binned"][name])
    raw_report = reports["raw"]
    second_output = tmp_path / "raw_second.xlsx"
    raw_report.write_excel(str(second_output), engine="openpyxl")
    workbook = load_workbook(second_output, data_only=True)
    try:
        for name, sheet_name in [("detail", "分组明细"), ("summary", "最终汇总"), ("trend", "最终KS趋势")]:
            rows = list(workbook[sheet_name].values)
            pd.testing.assert_frame_equal(
                pd.DataFrame(rows[1:], columns=rows[0]), native_frames["raw"][name],
            )
    finally:
        workbook.close()
    _assert_report_unchanged(raw_report, original["raw"])

    # 原生导出后的通用快照和新进程静态导出继续交付源 raw 结果。
    snapshot = snapshot_report(raw_report)
    snapshot_output = tmp_path / "raw_snapshot.xlsx"
    snapshot.write_excel(str(snapshot_output))
    _assert_current_excel(snapshot_output, snapshot)
    archive = tmp_path / "raw.marsreport"
    raw_report.save(archive)
    restored_output = tmp_path / "raw_restored.xlsx"
    environment = dict(os.environ)
    environment["PYTHONPATH"] = str(Path(mars.__file__).resolve().parents[1])
    environment["PYTHONIOENCODING"] = "utf-8"
    script = (
        "import sys; from mars.reporting import load_report; "
        "r=load_report(sys.argv[1]); r.write_excel(sys.argv[2]); "
        "print(r.to_ai_context(tables=['summary'],features='x',limit=2,max_chars=24000))"
    )
    completed = subprocess.run(
        [sys.executable, "-c", script, str(archive), str(restored_output)],
        cwd=tmp_path, env=environment, capture_output=True, text=True, encoding="utf-8", check=True,
    )
    assert json.loads(completed.stdout) == json.loads(original["raw"]["context"])
    _assert_current_excel(restored_output, snapshot)
    _assert_report_unchanged(raw_report, original["raw"])


def _assert_report_unchanged(report: MarsBinningReport, before: dict[str, Any]) -> None:
    """导出前后公共表、语义、身份和原始对象均保持不变。"""
    assert report.describe() == before["describe"]
    assert report.report_meta == before["meta"]
    for actual, original in zip(
        (report.summary_table, report.detail_table, report.trend_tables, report.report_meta),
        before["objects"],
    ):
        assert actual is original
    for name, expected in before["tables"].items():
        pd.testing.assert_frame_equal(to_polars(report.get_table(name)).to_pandas(), expected.to_pandas())
    assert report.to_ai_context(
        tables=["summary"], features="x", limit=2, max_chars=24000,
    ) == before["context"]


def _assert_native_excel_relationships(path: Path, detail_rows: int) -> None:
    """检查 22 列原生表、透视字段及 ZIP 内关系完整，不冒充 Excel 交互刷新。"""
    namespace = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if name.startswith("xl/pivotCache/pivotCacheDefinition") and name.endswith(".xml"):
                cache = ElementTree.fromstring(archive.read(name))
                fields = cache.find("s:cacheFields", namespace)
                assert fields is not None and len(fields) == 22
                assert fields[16].attrib["name"] == "ks_bin" and fields[21].attrib["name"] == "ks"
                source = cache.find("s:cacheSource/s:worksheetSource", namespace)
                assert source is not None and source.attrib["name"] == "detail"
                assert cache.attrib["refreshOnLoad"] == "1"
            elif name.endswith(".rels"):
                from posixpath import normpath

                parent = Path(name).parent.parent.as_posix()
                for relation in ElementTree.fromstring(archive.read(name)):
                    if relation.attrib.get("TargetMode") == "External":
                        continue
                    target = relation.attrib["Target"]
                    resolved = target.lstrip("/") if target.startswith("/") else normpath(f"{parent}/{target}")
                    assert resolved in archive.namelist(), (name, resolved)
            elif name.startswith("xl/worksheets/") and name.endswith(".xml"):
                sheet = ElementTree.fromstring(archive.read(name))
                if sheet.find("s:tableParts", namespace) is not None:
                    assert sheet.find("s:autoFilter", namespace) is None, name
        native_table = ElementTree.fromstring(archive.read("xl/tables/table1.xml"))
        assert native_table.attrib["ref"] == f"A1:V{detail_rows + 1}"


def to_polars(frame: pd.DataFrame | pl.DataFrame) -> pl.DataFrame:
    """测试中按真实原后端比较同一报告数值。"""
    return pl.from_pandas(frame) if isinstance(frame, pd.DataFrame) else frame


def test_native_excel_builds_binned_ks_without_source_final_ks(tmp_path: Path) -> None:
    """原生视图只依赖逐箱 KS，源 raw ks 是否物化不影响导出口径。"""
    report = profile_risk(
        pl.DataFrame({"x": list(range(12)), "bad": [0, 1] * 6}),
        target="bad", features=["x"], ks_method="raw", n_bins=2,
    ).report
    incomplete = MarsBinningReport(
        report.summary_table, report.trend_tables, report.detail_table.drop("ks"),
        detail_group_col=report.detail_group_col, report_meta=report.report_meta,
    )
    output = tmp_path / "without_source_ks.xlsx"
    incomplete.write_excel(str(output), engine="openpyxl")
    workbook = load_workbook(output, data_only=True)
    try:
        rows = list(workbook["分组明细"].values)
        actual = pd.DataFrame(rows[1:], columns=rows[0])
        assert actual["ks"].tolist() == pytest.approx([report.detail_table["ks_bin"].max()] * len(actual))
    finally:
        workbook.close()


@pytest.mark.parametrize("missing", ["ks_bin", "feature", "y", "mars_group"])
def test_native_excel_rejects_missing_export_dimension_or_bin_ks(
    missing: str, tmp_path: Path,
) -> None:
    """必要字段缺失时明确报错，不能静默用源 raw KS 混合导出。"""
    report = profile_risk(
        pl.DataFrame({"x": list(range(12)), "bad": [0, 1] * 6}),
        target="bad", features=["x"], ks_method="raw", n_bins=2,
    ).report
    report._detail = report.detail_table.drop(missing)
    with pytest.raises(ValueError, match="y|target" if missing == "y" else missing):
        report.write_excel(str(tmp_path / f"missing_{missing}.xlsx"), engine="openpyxl")


@pytest.mark.parametrize("backend", ["pandas", "polars"])
def test_export_view_uses_actual_null_groups_and_preserves_order(backend: str) -> None:
    """真实分组列的 null 与分类未观测组合不会串值，全部未定义仍为空。"""
    summary = pd.DataFrame({
        "feature": ["b", "a", "b", "a"], "target": ["second", "first", "first", "second"],
        "ks": [91.0, 92.0, 93.0, 94.0], "iv": [0.1, 0.2, 0.3, 0.4],
    })
    detail = pd.DataFrame({
        "feature": ["a", "b", "a", "b", "a", "a", "b", "b", "a", "a"],
        "y": ["first", "second", "first", "second", "second", "first", "first", "second", "first", "second"],
        "actual_bucket": [None, "g", "Total", "Total", "Total", None, "Total", "g", "Total", "g"],
        "mars_group": ["wrong"] * 10,
        "ks_bin": [11.0, 23.0, 7.0, 13.0, 19.0, 17.0, 29.0, 31.0, 5.0, None],
        "ks": [99.0] * 10, "count": list(range(10)),
    })
    detail["feature"] = pd.Categorical(detail["feature"], categories=["a", "b", "unseen"])
    detail["y"] = pd.Categorical(detail["y"], categories=["first", "second", "absent"])
    detail["actual_bucket"] = pd.Categorical(detail["actual_bucket"], categories=["Total", "g", "unseen"])
    detail.index = [8, 3, 11, 4, 7, 14, 0, 2, 12, 6]
    trend = pd.DataFrame({
        "feature": ["b", "a"], "display_name": ["乙", "甲"], "dtype": ["Float64"] * 2,
        "data_source": ["来源乙", "来源甲"], "g": [88.0, None], "null": [None, 77.0],
        "Total": [66.0, 55.0],
    })
    convert = pl.from_pandas if backend == "polars" else lambda frame: frame
    report = MarsBinningReport(
        convert(summary), {"ks": convert(trend)}, convert(detail),
        group_col="business_group", detail_group_col="actual_bucket",
        report_meta={"ks_method": "raw", "targets": ["first", "second"],
                     "ks_source_by_feature": {"a": "raw", "b": "raw"}, "raw_ks_diagnostics": []},
    )
    before = deepcopy(report.describe())
    original_trend = to_pandas_frame(report.trend_tables["ks"])
    original_detail = to_polars(report.detail_table).clone()
    exported, additional = _BinningExcelWriter(report)._prepare_excel_export()
    assert exported["ks"].tolist()[:-1] == [17.0, 31.0, 7.0, 13.0, 19.0, 17.0, 29.0, 31.0, 7.0]
    assert pd.isna(exported["ks"].iloc[-1])
    assert exported["count"].tolist() == list(range(10))
    if backend == "pandas":
        assert list(exported.index) == list(detail.index)
    assert exported["ks_bin"].tolist()[:-1] == detail["ks_bin"].tolist()[:-1]
    assert additional["最终汇总"]["ks"].tolist() == [13.0, 7.0, 29.0, 19.0]
    assert additional["最终汇总"]["iv"].tolist() == summary["iv"].tolist()
    exported_trend = additional["最终KS趋势"]
    pd.testing.assert_frame_equal(
        exported_trend.loc[:, trend.columns[:4]], original_trend.loc[:, trend.columns[:4]],
    )
    # 聚合趋势中的 Float64 不能证明原输入 dtype，可信源仍应保持未知。
    assert exported_trend["dtype"].isna().all()
    pd.testing.assert_frame_equal(to_pandas_frame(report.trend_tables["ks"]), original_trend)
    assert exported_trend["Total"].tolist() == [29.0, 7.0]
    assert pd.isna(exported_trend["g"]).all()
    assert pd.isna(exported_trend["null"].iloc[0]) and exported_trend["null"].iloc[1] == 17.0
    assert report.describe() == before
    pd.testing.assert_frame_equal(to_polars(report.detail_table).to_pandas(), original_detail.to_pandas())


@pytest.mark.parametrize("method", ["raw", "binned"])
def test_native_excel_keeps_undefined_total_ks_empty(method: str, tmp_path: Path) -> None:
    """整份无标签报告的分箱 KS 保持未定义，汇总和趋势不能填 0。"""
    report = profile_risk(
        pl.DataFrame({"x": [1, 2, 3, 4], "bad": [None] * 4}),
        target="bad", features=["x"], ks_method=method, n_bins=2,
    ).report
    path = tmp_path / f"undefined_{method}.xlsx"
    report.write_excel(str(path), engine="openpyxl")
    workbook = load_workbook(path, data_only=True)
    try:
        tables = [("分组明细", ["ks", "ks_bin"]), ("最终汇总", ["ks"])]
        if "ks" in report.trend_tables:
            tables.append(("最终KS趋势", ["Total"]))
        else:
            assert "最终KS趋势" not in workbook.sheetnames
        for name, columns in tables:
            rows = list(workbook[name].values)
            actual = pd.DataFrame(rows[1:], columns=rows[0])
            assert actual[columns].isna().all().all()
    finally:
        workbook.close()


def test_openpyxl_and_xlwings_consume_same_prepared_export(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    """模拟原生引擎只验证传入值，共享准备过程同时提供所有 Excel 入口。"""
    features = ["unknown", "literal_none", "literal_nan", "literal_na"]
    report = profile_risk(
        pl.DataFrame({**{feature: list(range(12)) for feature in features}, "bad": [0, 1] * 6}),
        target="bad", features=features, ks_method="raw", n_bins=2,
        feature_metadata={
            "literal_none": {"display_name": "None"},
            "literal_nan": {"display_name": "nan"},
            "literal_na": {"display_name": "<NA>"},
        },
    ).report
    source = snapshot_report(report)
    prepared: list[tuple[pd.DataFrame, dict[str, pd.DataFrame]]] = []
    original_prepare = _BinningExcelWriter._prepare_excel_export

    def prepare(writer: _BinningExcelWriter) -> tuple[pd.DataFrame, dict[str, pd.DataFrame]]:
        result = original_prepare(writer)
        prepared.append(result)
        return result

    monkeypatch.setattr(_BinningExcelWriter, "_prepare_excel_export", prepare)
    output = tmp_path / "openpyxl.xlsx"
    report.write_excel(str(output), engine="openpyxl")

    # 不启动用户 Excel；此模拟只证明 xlwings 获得同一导出数据及刷新调用。
    workbook, app = MagicMock(), MagicMock()
    app.books.open.return_value = workbook
    sheets: dict[str, Any] = {}
    detail = MagicMock()
    detail.name = "分组明细"
    detail.tables = [MagicMock()]
    detail.used_range.last_cell.row = 1
    sheets[detail.name] = detail

    def add_sheet(name: str) -> Any:
        sheet = MagicMock()
        sheet.name = name
        sheets[name] = sheet
        return sheet

    workbook.sheets.__getitem__.side_effect = sheets.__getitem__
    workbook.sheets.add.side_effect = add_sheet
    monkeypatch.setitem(sys.modules, "xlwings", SimpleNamespace(App=lambda **kwargs: app))
    written: dict[str, list[list[Any]]] = {}

    def write_rows(
        worksheet: Any, rows: list[list[Any]], *, start_row: int, first_col: int, last_col: int,
    ) -> None:
        written.setdefault(worksheet.name, []).extend(rows)

    refresh = MagicMock()
    monkeypatch.setattr(_BinningExcelWriter, "_write_xlwings_detail_rows", staticmethod(write_rows))
    monkeypatch.setattr(_BinningExcelWriter, "_refresh_xlwings_pivots", staticmethod(refresh))
    report.write_excel(str(tmp_path / "xlwings.xlsx"), engine="xlwings")
    assert len(prepared) == 2
    pd.testing.assert_frame_equal(prepared[0][0], prepared[1][0])
    for name, expected in prepared[0][1].items():
        pd.testing.assert_frame_equal(expected, prepared[1][1][name])
        if expected.shape[1]:
            records = written[name]
            _assert_excel_records(records, expected)
    records = written["分组明细"]
    assert len(records[0]) == 22 and records[0][16] == "ks_bin" and records[0][21] == "ks"
    _assert_excel_records(records, prepared[0][0][records[0]])
    metadata = prepared[0][1]["FeatureMetadata"].set_index("feature")["display_name"]
    assert pd.isna(metadata["unknown"])
    assert metadata.loc[["literal_none", "literal_nan", "literal_na"]].tolist() == ["None", "nan", "<NA>"]
    for name in source.describe()["tables"]:
        pd.testing.assert_frame_equal(
            to_pandas_frame(report.get_table(name)), to_pandas_frame(source.get_table(name)),
        )
    refresh.assert_called_once_with(workbook)
    workbook.save.assert_called_once_with(str(tmp_path / "xlwings.xlsx"))
    workbook.close.assert_called_once()
    app.quit.assert_called_once()


def test_xlwings_refresh_labels_binned_ks_and_preserves_average(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Windows 原生刷新绑定导出 ks，综合指标继续使用模板 average 聚合。"""
    monkeypatch.setattr(sys, "platform", "win32")
    average = SimpleNamespace(SourceName="ks", Function=-4106, Caption="旧标题")
    maximum = SimpleNamespace(SourceName="ks", Function=-4136, Caption="旧标题")
    fields = MagicMock(Count=2)
    fields.Item.side_effect = lambda index: [average, maximum][index - 1]
    pivot = MagicMock()
    pivot.DataFields = fields
    pivots = MagicMock(Count=1)
    pivots.Item.return_value = pivot
    sheet = MagicMock()
    sheet.api.PivotTables.return_value = pivots
    workbook = MagicMock()
    workbook.sheets = [sheet]
    _BinningExcelWriter._refresh_xlwings_pivots(workbook)
    assert average.Caption == "avg(分箱 KS，0–100点)"
    assert maximum.Caption == "分箱 KS（0–100点）"
    assert average.Function == -4106 and maximum.Function == -4136
    assert pivot.PivotCache().MissingItemsLimit == 0
    pivot.PivotCache().Refresh.assert_called_once()
    pivot.RefreshTable.assert_called_once()
    workbook.app.api.CalculateFullRebuild.assert_called_once()
    maximum.SourceName = "ks_bin"
    with pytest.raises(ValueError, match="ks_bin"):
        _BinningExcelWriter._refresh_xlwings_pivots(workbook)
