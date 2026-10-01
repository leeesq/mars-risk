"""真实聚合口径、跨进程快照及独立 Agent 消费的闭环验证。"""

from __future__ import annotations

import json
import os
import runpy
import subprocess
import sys
from pathlib import Path
from typing import Any

import polars as pl
import pytest
from openpyxl import load_workbook
from polars.testing import assert_frame_equal

from mars.analysis import (
    cross_scores,
    evaluate_score_policy,
    get_score_cell,
    write_score_cross_html,
)
from mars.reporting import load_report


def _scoped_input() -> pl.DataFrame:
    """手工验收样本包含真实缺失、特殊箱、权重和非笛卡尔 scope。"""
    return pl.DataFrame(
        {
            "old_prob": [0.1, 0.1, 0.1, 0.4, 0.4, 0.9, None, -99.0, 1.2, 0.9, 0.1, 0.4],
            "new_prob": [0.1, 0.4, 0.9, 0.4, None, 0.9, 0.4, 0.1, 0.1, 0.4, 0.1, 0.4],
            "bad": [0, 1, None, 0, 1, 1, 0, 1, 1, 0, None, 0],
            "later": [None, 0, None, 1, 0, 0, 1, 1, 0, None, None, 1],
            "weight": [1.0, 3.0, 2.0, 0.0, 2.0, 5.0, 1.0, 4.0, 1.0, 2.0, 0.0, 7.0],
            "amount": [10.0, 30.0, 20.0, -1.0, None, 50.0, 10.0, 40.0, 10.0, 20.0, 0.0, 70.0],
            "cohort": ["验证 & <样本>"] * 9 + ["观察</script>", "观察</script>", "另一个组"],
            "date": ["2026-01-01"] * 9 + ["2026-02-01"] * 3,
        }
    )


def _scoped_report(*, weighted: bool) -> Any:
    """通过用户 API 产生 3×4 常规箱报告，不在前端生成统计。"""
    return cross_scores(
        _scoped_input(),
        score_x="old_prob",
        score_y="new_prob",
        score_directions={"old_prob": "higher_risk", "new_prob": "higher_risk"},
        probability_scores=["old_prob", "new_prob"],
        targets=["bad", "later"],
        cutpoints={"old_prob": [0.3, 0.6], "new_prob": [0.2, 0.5, 0.8]},
        special_values={"old_prob": [-99.0]},
        group_col="cohort",
        time_col="date",
        time_grain="month",
        weights_col="weight" if weighted else None,
        amount_col="amount",
        min_observed=2,
        confidence_level=0.9,
        feature_metadata={
            "old_prob": {"display_name": "旧模型 <script>字段</script>", "data_source": "旧模型"},
            "new_prob": {"display_name": "新模型 & P(bad)", "data_source": "新模型"},
        },
        business_context={"dataset_id": "hand-counted-contract-fixture", "sample_unit": "测试行"},
    )


@pytest.mark.parametrize("weighted", [False, True])
def test_scopes_margins_and_expression_share_one_denominator(weighted: bool) -> None:
    report = _scoped_report(weighted=weighted)
    assert report.describe()["parameters"]["actual_scope_count"] == 3
    for scope in report.get_table("overall").to_dicts():
        filters = {key: scope[key] for key in ("target", "group", "period")}
        cells = report.get_table("cells", filters=filters)
        rows = report.get_table("row_summary", filters=filters)
        columns = report.get_table("column_summary", filters=filters)
        for field in ["sample_count", "observed_sample_count", "bad_sample_count", "tot_amt"]:
            assert cells[field].sum() == scope[field]
            assert rows[field].sum() == scope[field]
            assert columns[field].sum() == scope[field]
        if weighted:
            for field in ["weight_sum", "observed_weight_sum", "bad_weight_sum"]:
                assert cells[field].sum() == scope[field]
                assert rows[field].sum() == scope[field]
                assert columns[field].sum() == scope[field]
        normal = cells.filter(pl.col("x_risk_rank").is_not_null() & pl.col("y_risk_rank").is_not_null())
        assert normal["sample_share"].sum() == pytest.approx(normal["sample_count"].sum() / scope["sample_count"])
        for cell in cells.to_dicts():
            row = rows.filter(pl.col("x_bin") == cell["x_bin"]).to_dicts()[0]
            assert cell["row_bad_rate"] == row["bad_rate"]
            if cell["bad_rate"] is not None and row["bad_rate"] is not None:
                assert cell["delta_vs_row"] == pytest.approx(cell["bad_rate"] - row["bad_rate"])
            else:
                assert cell["delta_vs_row"] is None
        policy = evaluate_score_policy(report, {"type": "expression", "expression": "X >= X1 OR Y >= 1"})
        result = policy.get_table("summary", filters={**filters, "rule": "candidate", "retained": True}).to_dicts()[0]
        assert result["sample_count"] == normal["sample_count"].sum()
        assert result["sample_share"] == pytest.approx(normal["sample_count"].sum() / scope["sample_count"])
        numerator = normal["bad_weight_sum"].sum() if weighted else normal["bad_sample_count"].sum()
        denominator = normal["observed_weight_sum"].sum() if weighted else normal["observed_sample_count"].sum()
        if denominator:
            assert result["bad_rate"] == pytest.approx(numerator / denominator)
            if scope["bad_rate"]:
                assert result["lift_vs_overall"] == pytest.approx(result["bad_rate"] / scope["bad_rate"])
        else:
            assert result["bad_rate"] is None and result["lift_vs_overall"] is None
    initial = report.get_table("overall", filters={"target": "bad", "group": "验证 & <样本>"}).to_dicts()[0]
    assert initial["sample_count"] == 9 and initial["observed_sample_count"] == 8
    assert initial["bad_sample_count"] == 5
    assert initial["bad_rate"] == pytest.approx(15 / 17 if weighted else 5 / 8)


def test_weighted_snapshot_new_process_exports_and_agent_evidence(tmp_path: Path) -> None:
    report = _scoped_report(weighted=True)
    path = tmp_path / "weighted-scopes.marsreport"
    report.save(path)
    restored = load_report(path)
    assert restored.report_id == report.report_id
    assert restored.describe() == report.describe()
    for name in report.describe()["tables"]:
        assert_frame_equal(restored.get_table(name), report.get_table(name))
    rules = {"type": "expression", "expression": "(X = X1) OR (X = X2 AND Y <= Y2)"}
    first = evaluate_score_policy(report, rules)
    replay = evaluate_score_policy(restored, rules)
    for name in first.describe()["tables"]:
        assert_frame_equal(replay.get_table(name), first.get_table(name))
    replay.save(tmp_path / "policy.marsreport")
    assert load_report(tmp_path / "policy.marsreport").describe() == replay.describe()
    filters = {"target": "bad", "group": "验证 & <样本>"}
    evidence = get_score_cell(restored, "b0", "b0", filters=filters)
    assert evidence["page"]["reference"]["report_id"] == report.report_id
    cell = evidence["page"]["data"].to_dicts()[0]
    assert cell["bad_rate"] == 0.0 and cell["lift_vs_overall"] == 0.0
    assert cell["weighted_ci_status"] == "unsupported"
    write_score_cross_html(restored, tmp_path / "original.html", policy_reports=[replay])
    restored.write_excel(str(tmp_path / "cross.xlsx"))
    workbook = load_workbook(tmp_path / "cross.xlsx", read_only=True)
    assert any(name.endswith("_cells") for name in workbook.sheetnames)
    assert any(name.endswith("_bins") for name in workbook.sheetnames)
    fields = restored.describe()["tables"]["cells"]["fields"]
    assert fields["delta_vs_row"]["unit"] == "ratio_difference"
    assert fields["lift_vs_overall"]["unit"] == "dimensionless"
    assert fields["overall_bad_rate"]["unit"] == "ratio"
    assert "包含特殊箱" in fields["overall_bad_rate"]["meaning"]
    assert "分母" in fields["scope_sample_count"]["meaning"]
    policy_tables = replay.describe()["tables"]
    assert policy_tables["summary"]["grain"] == "target/group/period/rule/retained"
    assert "权重口径" in policy_tables["changes"]["fields"]["candidate_bad_rate"]["meaning"]
    script = (
        "from pathlib import Path; from mars.reporting import load_report; "
        "from mars.analysis import evaluate_score_policy; "
        "r=load_report(Path(__import__('sys').argv[1])); "
        "r.write_html(__import__('sys').argv[2]); "
        "p=evaluate_score_policy(r, {'type':'expression','expression':'X <= 2 AND Y != Y4'}); "
        "print(r.report_id); print(p.get_table('summary').height)"
    )
    completed = subprocess.run(
        [sys.executable, "-c", script, str(path), str(tmp_path / "new-process.html")],
        cwd=tmp_path,
        env={**os.environ, "PYTHONPATH": str(Path(__file__).resolve().parents[1] / "src")},
        capture_output=True,
        text=True,
        check=True,
    )
    assert report.report_id in completed.stdout
    html = (tmp_path / "new-process.html").read_text(encoding="utf-8")
    assert "<script>字段</script>" not in html and "观察</script>" not in html
    assert "\\u003c/script>" in html
    context = json.loads(restored.to_ai_context(queries={"cells": {"filters": filters, "limit": 2}}, max_chars=20000))
    assert context["evidence"][0]["report_id"] == report.report_id
    assert context["evidence"][0]["reference"] == "cells"
    assert context["evidence"][0]["query"]["filters"] == filters
    snippet = Path(__file__).resolve().parents[1] / "docs/snippets/correlation_and_score_cross.py"
    namespace = runpy.run_path(str(snippet))
    namespace["export_saved"](path, tmp_path / "saved-export", "X <= X2 AND Y <= Y3")
    assert (tmp_path / "saved-export/score-cross.html").is_file()
    assert load_report(tmp_path / "saved-export/score-policy.marsreport").describe()["parameters"]["parent_report_id"] == report.report_id


@pytest.mark.skipif(sys.version_info < (3, 10), reason="mars.agent requires Python >=3.10")
def test_agent_cell_and_rule_queries_keep_all_five_dimensions(tmp_path: Path) -> None:
    from mars.agent import MarsAgentSession, MarsRiskAgent

    report = _scoped_report(weighted=True)
    path = tmp_path / "agent-cross.marsreport"
    report.save(path)
    restored = load_report(path)
    filters = {"target": "bad", "group": "验证 & <样本>"}
    cell = get_score_cell(restored, "b0", "b0", filters=filters)["page"]["data"].to_dicts()[0]
    session = MarsAgentSession()
    handle = session.register_report(restored)
    result = MarsRiskAgent().execute_tool(
        "get_report_table",
        {
            "report_id": handle,
            "table": "cells",
            "filters": {**filters, "period": cell["period"], "x_bin": "b0", "y_bin": "b0"},
        },
        session=session,
    )
    assert result.success, result.error_message
    assert result.data["persistent_report_id"] == report.report_id
    assert result.data["rows"] == [cell]
    replay = evaluate_score_policy(restored, {"type": "expression", "expression": "X <= X2"})
    policy_handle = session.register_report(replay)
    rule_filters = {**filters, "period": cell["period"], "rule": "candidate", "retained": True}
    rule_result = MarsRiskAgent().execute_tool(
        "get_report_table",
        {"report_id": policy_handle, "table": "summary", "filters": rule_filters},
        session=session,
    )
    assert rule_result.success, rule_result.error_message
    assert rule_result.data["rows"] == replay.get_table("summary", filters=rule_filters).to_dicts()
    overflow = MarsRiskAgent().execute_tool(
        "get_report_table",
        {"report_id": handle, "table": "cells", "filters": {str(i): i for i in range(9)}},
        session=session,
    )
    assert not overflow.success and "too many fields" in overflow.error_message
