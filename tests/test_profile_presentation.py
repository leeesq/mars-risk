"""画像比较表的真实 Notebook 渲染与现有导出契约。"""

from __future__ import annotations

from pathlib import Path
from typing import Any

import pandas as pd
import polars as pl
import pytest
from pandas.testing import assert_frame_equal

from mars.analysis import profile_stats
from mars.compute import to_pandas_frame
from mars.reporting import MarsProfileReport


def _comparison_report(backend: str, grouped: bool) -> MarsProfileReport:
    """构造可比较、不可比较、缺字段和全空类别的真实比较结果。"""
    current = pl.DataFrame(
        {
            "category": ["甲", "新类别", "甲", None],
            "number": [1, 2, 3, 4],
            "empty_category": pl.Series([None] * 4, dtype=pl.Utf8),
            "current_only": [True, False, True, False],
            "month": ["开发", "观察", "开发", "观察"],
        }
    )
    benchmark = pl.DataFrame(
        {
            "category": ["甲", "乙"],
            "number": ["one", "two"],
            "empty_category": pl.Series([None, None], dtype=pl.Utf8),
            "benchmark_only": [1, 2],
            "month": ["开发", "开发"],
        }
    )
    return profile_stats(
        current.to_pandas() if backend == "pandas" else current,
        metrics=["schema", "unseen"],
        benchmark_df=benchmark.to_pandas() if backend == "pandas" else benchmark,
        categorical_features=["number"],
        group_col="month" if grouped else None,
    )


@pytest.mark.parametrize("backend", ["pandas", "polars"])
@pytest.mark.parametrize("grouped", [False, True])
@pytest.mark.parametrize("metric", ["schema", "unseen"])
@pytest.mark.parametrize("explicit", [False, True])
def test_comparison_notebook_renders_native_text_and_numbers(
    backend: str, grouped: bool, metric: str, explicit: bool, tmp_path: Path
) -> None:
    """原生比较值在 Notebook、HTML 和 Excel 中保持一致。"""
    report = _comparison_report(backend, grouped)
    before = to_pandas_frame(report.get_table(f"comparison.{metric}"))
    options: dict[str, Any] = {"sort_by": "feature", "sort_ascending": True} if explicit else {}
    styled = report.show_trend(metric, **options)
    html = styled.to_html()
    assert "String" in html and "Int64" in html
    assert "incompatible" in html and "dtype families are not comparable" in html
    assert "current_only" in html and "benchmark_only" in html
    assert_frame_equal(before, to_pandas_frame(report.get_table(f"comparison.{metric}")))
    displayed = styled.data.set_index("feature").loc[before["feature"], before.columns[1:]]
    expected_display = before.copy()
    expected_display["display_name"] = expected_display["display_name"].fillna(expected_display["feature"])
    assert_frame_equal(displayed.reset_index(), expected_display, check_dtype=True)
    if explicit:
        assert styled.data["feature"].tolist() == sorted(before["feature"].tolist())
    elif metric == "schema":
        assert styled.data["feature"].tolist() == sorted(before["feature"].tolist(), reverse=True)
    else:
        totals = styled.data["total"].dropna().tolist()
        assert totals == sorted(totals, reverse=True)
        assert styled.data.set_index("feature").loc["category", "total"] == pytest.approx(1 / 3)
    if metric == "unseen":
        assert "no_reference_values" in html
        assert "background-color:" in html and "0.33" in html
    path = tmp_path / "comparison.html"
    report.write_html(str(path))
    exported = path.read_text(encoding="utf-8")
    assert "dtype families are not comparable" in exported
    excel = tmp_path / "comparison.xlsx"
    report.write_excel(str(excel))
    actual = pd.read_excel(excel, sheet_name=f"Compare_{metric.capitalize()}")
    assert_frame_equal(actual.isna(), before.isna())
    # Excel 空单元格恢复为 NaN；按同一缺失值边界核对全部原统计与字典文本。
    assert_frame_equal(
        actual.astype(object).where(actual.notna(), None),
        before.astype(object).where(before.notna(), None), check_dtype=False,
    )


@pytest.mark.parametrize("backend", ["pandas", "polars"])
@pytest.mark.parametrize("metric", ["schema", "unseen"])
def test_comparison_empty_selection_and_unknown_sort_are_explicit(
    backend: str, metric: str
) -> None:
    """空结果仍能渲染，用户显式请求不存在的排序列清楚失败。"""
    report = _comparison_report(backend, grouped=False)
    styled = report.show_trend(metric, limit=0)
    assert styled.data.empty
    assert "<table" in styled.to_html() and "status" in styled.to_html()
    for column in ("not_a_column", "total" if metric == "schema" else "not_a_column"):
        with pytest.raises(ValueError, match="Unknown columns"):
            report.show_trend(metric, sort_by=column)


def test_shared_styler_keeps_nullable_numeric_boolean_and_all_empty_columns() -> None:
    """显式渐变子集不能把文本、Boolean 或全空列转成数字。"""
    report = MarsProfileReport(pd.DataFrame(), dq_tables={}, stats_tables={})
    frame = pd.DataFrame(
        {
            "feature": ["first", "second"],
            "current_dtype": ["Int64", "Int64"],
            "status": ["comparable", "not_computed"],
            "reason": [None, "no_target"],
            "count": pd.Series([1, None], dtype="Int64"),
            "rate": pd.Series([0.25, None], dtype="Float64"),
            "all_empty": pd.Series([None, None], dtype="Float64"),
            "comparable": pd.Series([True, None], dtype="boolean"),
        }
    )
    styled = report._get_styler(
        frame, title="Nullable comparison", cmap="Blues",
        subset_cols=list(frame.columns),
    )
    html = styled.to_html()
    assert "25.00%" in html and ">1<" in html
    assert "Int64" in html and "not_computed" in html and "no_target" in html
    assert "True" in html and "background-color:" in html
    assert_frame_equal(styled.data, frame)
    styled._compute()
    assert {column for _, column in styled.ctx} == {4, 5}


@pytest.mark.parametrize("backend", ["pandas", "polars"])
@pytest.mark.parametrize("metric", ["schema", "unseen"])
def test_explicit_comparison_sort_list_remains_supported(backend: str, metric: str) -> None:
    """合法多列排序与明确错误沿用公共查询契约。"""
    report = _comparison_report(backend, grouped=True)
    styled = report.show_trend(metric, sort_by=["feature"], sort_ascending=True)
    assert styled.data["feature"].tolist() == sorted(styled.data["feature"].tolist())
    assert "<table" in styled.to_html()
    with pytest.raises(ValueError, match="Unknown columns"):
        report.show_trend(metric, sort_by=["feature", "not_a_column"])


def test_existing_profile_trends_keep_default_total_order_and_percentage_format() -> None:
    """普通数值趋势保留按 total 降序与比例格式。"""
    report = MarsProfileReport(
        pd.DataFrame({"feature": ["income", "debt"], "missing_rate": [0.1, 0.4]}),
        dq_tables={
            "missing": pd.DataFrame(
                {"feature": ["income", "debt"], "total": [0.1, 0.4], "period": [0.2, 0.3]}
            )
        },
        stats_tables={},
    )
    styled = report.show_trend("missing")
    assert styled.data["feature"].tolist() == ["debt", "income"]
    assert "40.00%" in styled.to_html() and "10.00%" in styled.to_html()
    assert "background-color:" in report.show_overview().to_html()


@pytest.mark.parametrize("backend", ["pandas", "polars"])
def test_profile_excel_uses_semantic_display_formats_without_changing_values(
    backend: str, tmp_path: Path,
) -> None:
    """Excel 比例、PSI、KS、AUC、整数与加权小数按真实语义显示。"""
    from openpyxl import load_workbook

    overview = pd.DataFrame({
        "feature": ["first", "second"], "missing_rate": [0, 1],
        "count": [1234, 5], "weight_sum": [1.25, 2.5], "auc": [0.75, 0.8],
        "ks": [25.0, 40.0], "mode_value": ["10", "2"],
    })
    trend = pd.DataFrame({"feature": ["first", "second"], "period": [0.125, 0.0], "total": [0.25, 1.0]})
    report = MarsProfileReport(
        pl.from_pandas(overview) if backend == "polars" else overview,
        dq_tables={"missing": pl.from_pandas(trend) if backend == "polars" else trend},
        stats_tables={"psi": pl.from_pandas(trend) if backend == "polars" else trend},
        feature_metadata={"first": {"display_name": "<申请收入>", "data_source": "申请 & 外部"}},
    )
    path = tmp_path / "formats.xlsx"
    report.write_excel(str(path))
    workbook = load_workbook(path)
    sheet = workbook["Overview"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    rows = {sheet.cell(row, headers["feature"]).value: row for row in range(2, sheet.max_row + 1)}
    expected = {
        "missing_rate": (0, "0.00%"), "count": (1234, "#,##0"),
        "weight_sum": (1.25, "#,##0.00"), "auc": (0.75, "#,##0.00"),
        "ks": (25, "#,##0.00"),
    }
    for column, (value, number_format) in expected.items():
        cell = sheet.cell(rows["first"], headers[column])
        assert cell.value == value and cell.number_format == number_format, column
    assert sheet.cell(rows["first"], headers["mode_value"]).number_format == "General"
    for sheet_name, number_format in (("DQ_Missing", "0.00%"), ("Trend_Psi", "#,##0.00")):
        trend_sheet = workbook[sheet_name]
        trend_headers = {cell.value: cell.column for cell in trend_sheet[1]}
        assert trend_sheet.cell(2, trend_headers["period"]).number_format == number_format
        assert trend_sheet.cell(2, trend_headers["display_name"]).number_format == "General"
        assert trend_sheet.cell(2, trend_headers["data_source"]).number_format == "General"
    notebook_html = report.show_overview().to_html()
    assert "&lt;申请收入&gt;" in notebook_html and "申请 &amp; 外部" in notebook_html
    assert "<申请收入>" not in notebook_html
    workbook.close()
