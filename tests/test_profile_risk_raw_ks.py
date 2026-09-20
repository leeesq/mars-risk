"""原始值 KS 的计算与画像契约。"""

from __future__ import annotations

from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
import polars as pl
import pytest
from openpyxl import load_workbook
from polars.testing import assert_frame_equal
from scipy.stats import ks_2samp

from mars.analysis import profile_risk
from mars.compute._raw_ks import _calculate_raw_ks
from mars.reporting.plotter import MarsPlotter


def _compute(df: pl.DataFrame, **kwargs: Any) -> dict[str, Any]:
    return _calculate_raw_ks(
        df.with_columns(pl.lit("Total").alias("g")),
        features_by_target={"y": ["x"]},
        group_col="g",
        weights_col=kwargs.get("weights_col"),
        missing_values=kwargs.get("missing_values"),
        special_values=kwargs.get("special_values"),
    ).row(0, named=True)


def test_raw_ks_replaces_summary_and_preserves_default() -> None:
    df = pl.DataFrame({"x": list(range(12)), "y": [0, 0, 0, 0, 0, 1, 0, 1, 1, 1, 1, 1]})
    default = profile_risk(df, target="y", features=["x"], n_bins=2).report
    explicit = profile_risk(
        df, target="y", features=["x"], n_bins=2, ks_method="binned"
    ).report
    raw = profile_risk(df, target="y", features=["x"], n_bins=2, ks_method="raw").report
    expected = (
        ks_2samp(
            df.filter(pl.col("y") == 0)["x"], df.filter(pl.col("y") == 1)["x"]
        ).statistic
        * 100
    )
    assert raw.summary_table["ks"][0] == pytest.approx(expected)
    assert raw.detail_table["ks"].unique().to_list() == pytest.approx([expected])
    assert raw.trend_tables["ks"]["Total"][0] == pytest.approx(expected)
    assert_frame_equal(default.summary_table, explicit.summary_table)
    assert_frame_equal(default.detail_table, explicit.detail_table)
    assert "ks" not in default.detail_table.columns
    assert_frame_equal(default.summary_table.drop("ks"), raw.summary_table.drop("ks"))


@pytest.mark.parametrize(
    "x,y",
    [
        ([1, 1, 2, 2, 3, 3], [0, 1, 0, 1, 0, 1]),
        ([1, 2, 3, 4], [0, 0, 1, 1]),
        ([1, 2, 3, 4], [1, 1, 0, 0]),
        ([4, 4, 4, 4], [0, 0, 1, 1]),
        ([2**60, 2**60 + 1, 2**60 + 2, 2**60 + 3], [0, 0, 1, 1]),
    ],
)
def test_raw_empirical_ks_matches_scipy(x, y) -> None:
    df = pl.DataFrame({"x": x, "y": y})
    expected = (
        ks_2samp(
            df.filter(pl.col("y") == 0)["x"], df.filter(pl.col("y") == 1)["x"]
        ).statistic
        * 100
    )
    assert _compute(df)["ks"] == pytest.approx(expected)


def test_weighted_ks_matches_independent_threshold_reference() -> None:
    rng = np.random.default_rng(4)
    x = rng.integers(0, 30, 200)
    y = rng.integers(0, 2, 200)
    w = rng.uniform(0, 3, 200)
    w[::4] = 0
    expected = (
        max(
            abs(
                w[(x <= threshold) & (y == 0)].sum() / w[y == 0].sum()
                - w[(x <= threshold) & (y == 1)].sum() / w[y == 1].sum()
            )
            for threshold in np.unique(x)
        )
        * 100
    )
    df = pl.DataFrame({"x": x, "y": y, "w": w})
    assert _compute(df, weights_col="w")["ks"] == pytest.approx(expected)
    # 同倍缩放不改变分布，即使直接累加原始权重可能溢出。
    huge = df.with_columns(pl.col("w") * 1e307)
    assert _compute(huge, weights_col="w")["ks"] == pytest.approx(expected)
    report = profile_risk(
        df, target="y", features=["x"], weights_col="w", ks_method="raw"
    ).report
    assert report.summary_table["ks"][0] == pytest.approx(expected)


def test_raw_excludes_unobserved_missing_and_special_values() -> None:
    df = pl.DataFrame(
        {
            "x": [
                None,
                float("nan"),
                float("inf"),
                -float("inf"),
                -999.0,
                -888.0,
                1.0,
                2.0,
                3.0,
            ],
            "y": [0, 1, 0, 1, 0, 1, 0, 1, None],
        }
    )
    result = _compute(
        df, missing_values=[-999, "missing"], special_values=[-888, "special"]
    )
    assert result["ks"] == 100
    assert result["valid_count"] == 2
    report = profile_risk(
        df,
        target="y",
        features=["x"],
        ks_method="raw",
        missing_values=[-999],
        special_values=[-888],
        benchmark_df=df.filter(pl.col("x").is_finite()),
    ).report
    assert report.summary_table["ks"][0] == 100


@pytest.mark.parametrize(
    "weights",
    [
        [-1.0, 1.0, 1.0, 1.0],
        [None, 1.0, 1.0, 1.0],
        [float("inf"), 1.0, 1.0, 1.0],
        [float("nan"), 1.0, 1.0, 1.0],
    ],
)
def test_raw_rejects_invalid_participating_weights(weights) -> None:
    df = pl.DataFrame({"x": [1, 2, 3, 4], "y": [0, 1, 0, 1], "w": weights})
    with pytest.raises(
        ValueError, match="finite non-negative weights.*feature='x'.*target='y'"
    ):
        profile_risk(df, target="y", features=["x"], weights_col="w", ks_method="raw")


def test_weights_outside_raw_sample_are_not_validated() -> None:
    df = pl.DataFrame(
        {
            "x": [None, 1, 2, 3],
            "y": [0, 0, 1, None],
            "w": [-1.0, 1.0, 1.0, float("nan")],
        }
    )
    assert _compute(df, weights_col="w")["ks"] == 100


@pytest.mark.parametrize("target", [None, "absent", "y"])
def test_raw_label_free_is_null_with_diagnostic(target) -> None:
    df = pl.DataFrame(
        {"x": [1, 2, 3, 4], "y": [None] * 4, "dummy_target": [0, 0, 1, 1]}
    )
    report = profile_risk(df, target=target, features=["x"], ks_method="raw").report
    assert report.summary_table["ks"][0] is None
    assert report.detail_table["ks"].null_count() == report.detail_table.height
    assert report.report_meta["raw_ks_diagnostics"][0]["reason"] == "no_valid_samples"


def test_grouped_multi_target_and_categories_share_final_ks() -> None:
    df = pl.DataFrame(
        {
            "x": [1, 2, 3, 4, 1, 2, 3, 4, 1, 2, 3, 4],
            "cat": ["a", "b"] * 6,
            "y": [0, 1, 0, 1] * 3,
            "z": [0, 0, 1, 1] * 3,
            "g": ["a"] * 4 + ["b"] * 4 + [None] * 4,
        }
    )
    args = dict(target=["y", "z"], features=["x", "cat"], group_col="g", n_bins=2)
    default = profile_risk(df, **args).report
    report = profile_risk(df, **args, ks_method="raw").report
    assert_frame_equal(
        report.summary_table.filter(pl.col("feature") == "cat"),
        default.summary_table.filter(pl.col("feature") == "cat"),
    )
    assert report.report_meta["ks_source_by_feature"] == {"x": "raw", "cat": "binned"}
    numeric = report.summary_table.filter(pl.col("feature") == "x")
    assert numeric.filter(pl.col("target") == "y")["ks"][0] == 50
    assert numeric.filter(pl.col("target") == "z")["ks"][0] == 100
    trend = report.trend_tables["ks"].filter(pl.col("feature") == "x")
    assert trend.select("a", "b", "null", "Total").row(0) == (50, 50, 50, 50)
    for panel in report.detail_table.partition_by(["feature", "y", "mars_group"]):
        expected = 50 if panel["y"][0] == "y" else 100
        if panel["feature"][0] == "x":
            assert (
                MarsPlotter._summarize_binning_metrics(panel.to_pandas())[1] == expected
            )


def test_raw_single_class_after_filtering_does_not_fall_back() -> None:
    df = pl.DataFrame({"x": [1, 2, -999, -999], "y": [0, 0, 1, 1]})
    report = profile_risk(
        df, target="y", features=["x"], missing_values=[-999], ks_method="raw"
    ).report
    assert report.summary_table["ks"][0] is None
    assert np.isnan(
        MarsPlotter._summarize_binning_metrics(report.detail_table.to_pandas())[1]
    )
    assert (
        report.report_meta["raw_ks_diagnostics"][0]["reason"]
        == "insufficient_class_weight"
    )
    zero = _compute(df.with_columns(pl.lit(0.0).alias("w")), weights_col="w")
    assert zero["ks"] is None
    with pytest.raises(ValueError, match="at least 2 observed classes"):
        profile_risk(
            df.with_columns(pl.lit(0).alias("y")),
            target="y",
            features=["x"],
            ks_method="raw",
        )


def test_limit_precedes_binning_and_categories_do_not_count(monkeypatch) -> None:
    from mars.analysis.evaluator import MarsBinEvaluator

    df = pl.DataFrame(
        {
            **{f"x{i}": [1, 2, 3, 4] for i in range(51)},
            "y": [0, 0, 1, 1],
            "cat": ["a", "b"] * 2,
        }
    )
    actual = MarsBinEvaluator.evaluate

    def forbidden(*args, **kwargs):
        pytest.fail("limit must be checked before evaluation")

    monkeypatch.setattr(MarsBinEvaluator, "evaluate", forbidden)
    with pytest.raises(ValueError, match="51 numeric features.*max_raw_ks_features=50"):
        profile_risk(df, target="y", ks_method="raw")
    monkeypatch.setattr(MarsBinEvaluator, "evaluate", actual)
    profile_risk(df.drop("x50"), target="y", ks_method="raw")
    profile_risk(df, target="y", ks_method="raw", max_raw_ks_features=51)
    profile_risk(
        df, target="y", features=["x0", "cat"], ks_method="raw", max_raw_ks_features=1
    )
    profile_risk(df, target="y", max_raw_ks_features=1)


@pytest.mark.parametrize(
    "options",
    [
        {"ks_method": "invalid"},
        {"max_raw_ks_features": 0},
        {"max_raw_ks_features": True},
        {"max_raw_ks_features": 1.5},
    ],
)
def test_invalid_ks_options(options) -> None:
    with pytest.raises(ValueError):
        profile_risk(pl.DataFrame({"x": [1, 2], "y": [0, 1]}), target="y", **options)


def test_pandas_benchmark_and_binning_independence() -> None:
    df = pl.DataFrame({"x": list(range(20)), "y": [0, 1] * 10})
    baseline = df.with_columns(-pl.col("x"))
    expected = _compute(df)["ks"]
    for method, bins in [
        ("quantile", 2),
        ("quantile", 10),
        ("uniform", 3),
        ("cart", 3),
    ]:
        report = profile_risk(
            df.to_pandas(),
            target="y",
            features=["x"],
            ks_method="raw",
            method=method,
            n_bins=bins,
            benchmark_df=baseline.to_pandas(),
        ).report
        assert isinstance(report.summary_table, pd.DataFrame)
        assert isinstance(report.detail_table, pd.DataFrame)
        assert isinstance(report.trend_tables["ks"], pd.DataFrame)
        assert report.summary_table["ks"].iloc[0] == pytest.approx(expected)


def test_extreme_class_weight_scales_preserve_distributions() -> None:
    df = pl.DataFrame(
        {"x": [1, 2, 3, 4], "y": [0, 0, 1, 1], "w": [1e-300, 1e-300, 1e300, 1e300]}
    )
    assert _compute(df, weights_col="w")["ks"] == 100


def test_group_with_one_valid_class_has_null_ks() -> None:
    df = pl.DataFrame({"x": [1, 2, 3, 4], "y": [0, 0, 1, 1], "g": ["a", "a", "b", "b"]})
    report = profile_risk(
        df, target="y", features=["x"], group_col="g", ks_method="raw"
    ).report
    assert report.trend_tables["ks"].select("a", "b", "Total").row(0) == (
        None,
        None,
        100,
    )


def test_category_only_raw_mode_keeps_binned_values() -> None:
    df = pl.DataFrame({"cat": ["a", "b"] * 4, "y": [0, 1, 0, 0, 1, 1, 0, 1]})
    binned = profile_risk(df, target="y", features=["cat"]).report
    raw = profile_risk(df, target="y", features=["cat"], ks_method="raw").report
    assert_frame_equal(binned.summary_table, raw.summary_table)
    assert raw.report_meta["ks_source_by_feature"] == {"cat": "binned"}
    assert raw.report_meta["raw_ks_diagnostics"] == []


def test_inferred_feature_limit_excludes_weights_and_amount() -> None:
    df = pl.DataFrame(
        {
            "x": [1, 2, 3, 4],
            "y": [0, 0, 1, 1],
            "w": [1.0, 2.0, 3.0, 4.0],
            "amt": [10, 20, 30, 40],
        }
    )
    report = profile_risk(
        df,
        target="y",
        weights_col="w",
        amount_col="amt",
        ks_method="raw",
        max_raw_ks_features=1,
    ).report
    assert report.summary_table["feature"].to_list() == ["x"]


def test_raw_excel_html_and_figure_contain_final_values(tmp_path: Path) -> None:
    from mars.reporting._matplotlib import require_pyplot

    df = pl.DataFrame(
        {
            "x": list(range(12)),
            "y": [0, 1] * 6,
            "dt": ["2024-01-01"] * 6 + ["2024-02-01"] * 6,
        }
    )
    report = profile_risk(
        df,
        target="y",
        features=["x"],
        time_col="dt",
        ks_method="raw",
        n_bins=2,
    ).report
    excel = tmp_path / "raw.xlsx"
    report.write_excel(str(excel))
    workbook = load_workbook(excel, data_only=True)
    try:
        for name, expected in [
            ("综合指标", report.summary_table),
            ("分组明细", report.detail_table),
            ("KS", report.trend_tables["ks"]),
        ]:
            sheet = workbook[name]
            rows = list(sheet.values)
            assert list(rows[0]) == expected.columns
            actual = pd.DataFrame(rows[1:], columns=rows[0])
            cols = ["ks"] if "ks" in actual else ["Total"]
            np.testing.assert_allclose(
                actual[cols].to_numpy(dtype=float),
                expected.select(cols).to_numpy(),
                equal_nan=True,
            )
        assert workbook["KS计算诊断"].max_row > 1
    finally:
        workbook.close()
    html = tmp_path / "raw.html"
    report.write_html(str(html), max_plots=0)
    assert "16.67" in html.read_text(encoding="utf-8")
    figures = report.build_risk_trend_figures(features=["x"], show_risk="count")
    try:
        texts = [text.get_text() for figure in figures for text in figure.texts]
        assert any("KS: 16.7" in text for text in texts)
    finally:
        for figure in figures:
            require_pyplot(feature_name="test").close(figure)


def test_batch_plot_sort_uses_final_ks(monkeypatch) -> None:
    frame = pd.DataFrame(
        {
            "feature": ["a", "b", "c"],
            "g": ["Total"] * 3,
            "ks": [10.0, 90.0, float("nan")],
            "ks_bin": [90.0, 10.0, 100.0],
        }
    )
    order = []
    monkeypatch.setattr(
        MarsPlotter,
        "plot_feature_binning_risk_trend",
        lambda df_detail, feature, **kwargs: order.append(feature),
    )
    MarsPlotter.plot_feature_binning_risk_trend_batch(
        frame,
        ["a", "b", "c"],
        group_col="g",
        sort_by="ks",
        ascending=False,
    )
    assert order == ["b", "a", "c"]


def test_grouped_weighted_ks_is_independent_of_row_order() -> None:
    rng = np.random.default_rng(91)
    df = pl.DataFrame(
        {
            "x": rng.integers(0, 30, 250),
            "y": rng.integers(0, 2, 250),
            "g": [str(i % 5) for i in range(250)],
            "w": rng.uniform(0, 10, 250),
        }
    ).sample(fraction=1, shuffle=True, seed=10)
    values = _calculate_raw_ks(
        df,
        features_by_target={"y": ["x"]},
        group_col="g",
        weights_col="w",
        missing_values=None,
        special_values=None,
    )
    for row in values.iter_rows(named=True):
        part = df if row["group"] == "Total" else df.filter(pl.col("g") == row["group"])
        x, y, w = (part[c].to_numpy() for c in ["x", "y", "w"])
        expected = (
            max(
                abs(
                    w[(x <= t) & (y == 0)].sum() / w[y == 0].sum()
                    - w[(x <= t) & (y == 1)].sum() / w[y == 1].sum()
                )
                for t in np.unique(x)
            )
            * 100
        )
        assert row["ks"] == pytest.approx(expected)


def test_raw_parameters_are_only_on_profile_risk() -> None:
    from inspect import signature

    from mars.analysis import MarsBinEvaluator

    for name in ("ks_method", "max_raw_ks_features"):
        assert name in signature(profile_risk).parameters
        assert name not in signature(MarsBinEvaluator).parameters
        assert name not in signature(MarsBinEvaluator.evaluate).parameters
